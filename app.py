import io
import os
import datetime
import functools
import hashlib
import hmac
import json
import re
import secrets

import openpyxl
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import time
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image as PillowImage, UnidentifiedImageError
from sqlalchemy import text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError, IntegrityError

# 1. CONFIGURAÇÃO INICIAL
st.set_page_config(page_title="Pesagem e Balanceamento", layout="wide", initial_sidebar_state="expanded")

if 'usuario_logado' not in st.session_state:
    st.session_state['usuario_logado'] = False
    st.session_state['nivel_acesso'] = 0 
    st.session_state['nome_usuario'] = ""
if 'pagina_atual' not in st.session_state:
    st.session_state['pagina_atual'] = 'consulta'

st.markdown(
    """
    <style>
    [data-testid="stDecoration"], [data-testid="stAppDeployButton"],
    [data-testid="stMainMenu"] { display: none !important; }
    [data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"] {
        display: flex !important; visibility: visible !important;
    }
    header[data-testid="stHeader"] { background: transparent; }
    .block-container { padding-top: 2.2rem; max-width: 1280px; }
    section[data-testid="stSidebar"] {
        background: #FFFFFF; border-right: 1px solid #E6E8EC;
    }
    section[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap: .35rem; }
    section[data-testid="stSidebar"] hr { margin: .9rem 0; }
    input:disabled, textarea:disabled {
        color: #1D2433 !important; -webkit-text-fill-color: #1D2433 !important;
        opacity: 1 !important; cursor: default;
    }
    [data-testid="stWidgetLabel"] p { color: #1D2433 !important; opacity: 1 !important; }
    [data-baseweb="input"]:has(input:disabled),
    [data-baseweb="input"]:has(input:disabled) > div,
    input:disabled { background: #F1F3F6 !important; }
    section[data-testid="stSidebar"] .stButton button {
        justify-content: flex-start; border: none; box-shadow: none;
        padding: 0.4rem 0.75rem; min-height: 2.3rem; font-weight: 500;
    }
    section[data-testid="stSidebar"] .stButton button > div {
        justify-content: flex-start; width: 100%;
    }
    section[data-testid="stSidebar"] .stButton button[kind="secondary"] {
        background: transparent; color: #3B4352;
    }
    section[data-testid="stSidebar"] .stButton button[kind="secondary"]:hover {
        background: #FFF1E8; color: #FF6A13;
    }
    .marca { display: flex; align-items: center; gap: .6rem; margin-bottom: .4rem; }
    .marca-logo {
        width: 38px; height: 38px; border-radius: 10px; background: #FF6A13;
        color: #fff; display: flex; align-items: center; justify-content: center;
        font-weight: 700; font-size: .8rem; flex-shrink: 0;
    }
    .marca-nome { font-weight: 700; font-size: 1.02rem; line-height: 1.1; }
    .marca-sub { color: #6B7280; font-size: .78rem; }
    .usuario {
        display: flex; align-items: center; gap: .65rem; padding: .2rem .4rem;
        font-size: .85rem; color: #1F2430; line-height: 1.25; margin-bottom: .4rem;
    }
    .usuario span { color: #6B7280; font-size: .78rem; }
    .usuario-avatar {
        width: 34px; height: 34px; border-radius: 50%; background: #FFF1E8;
        color: #FF6A13; font-weight: 700; font-size: .8rem; flex-shrink: 0;
        display: flex; align-items: center; justify-content: center;
    }
    .grupo-menu {
        color: #9AA1AD; font-size: .7rem; font-weight: 600; letter-spacing: .06em;
        text-transform: uppercase; margin: 1.1rem 0 .1rem .75rem;
    }
    /* O Streamlit puxa o texto seguinte 1rem para cima; nos rótulos do menu isso
       fazia o botão ativo cobrir o nome do grupo. */
    section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"]:has(> .grupo-menu, > .usuario) {
        margin-bottom: 0;
    }
    .cabecalho h1 { font-size: 1.75rem; font-weight: 700; margin: 0; padding: 0; }
    .cabecalho p { color: #6B7280; margin: .25rem 0 1.2rem; }
    div[data-testid="stMetric"] {
        background: #FFFFFF; border: 1px solid #E6E8EC; border-radius: 12px;
        padding: .85rem 1rem;
    }
    div[data-testid="stMetricLabel"] { color: #6B7280; }
    .stTabs [data-baseweb="tab-list"] { gap: .25rem; }
    .stTabs [data-baseweb="tab"] { padding: .4rem .9rem; }
    </style>
    """,
    unsafe_allow_html=True,
)


NIVEL_CONSULTA, NIVEL_OPERACAO, NIVEL_ADMIN = 2, 1, 3
# Engenharia de Operações só vê, na consulta, a ficha mais recente de cada aeronave.
NIVEL_ENG_OPERACOES = 4
# Especialistas GOL abrem as telas de emissão e aprovação, mas não gravam nada.
NIVEL_ESPECIALISTA = 5
PERFIS = {
    NIVEL_CONSULTA: "Consulta",
    NIVEL_ENG_OPERACOES: "Engenharia de Operações",
    NIVEL_ESPECIALISTA: "Especialista GOL (só visualiza emissão)",
    NIVEL_OPERACAO: "Gera e aprova",
    NIVEL_ADMIN: "Administrador",
}


def nivel_atual():
    return st.session_state.get('nivel_acesso', 0)


# O Brasil não tem horário de verão desde 2019: Brasília é sempre UTC-3.
FUSO_BRASILIA = datetime.timezone(datetime.timedelta(hours=-3), "BRT")


def formatar_data_hora_brasilia(valor):
    try:
        momento = datetime.datetime.fromisoformat(safe_str(valor))
    except ValueError:
        return safe_str(valor) or "—"
    if momento.tzinfo is None:
        momento = momento.replace(tzinfo=datetime.timezone.utc)
    return momento.astimezone(FUSO_BRASILIA).strftime("%d/%m/%Y %H:%M")


def exigir_nivel(*niveis):
    if nivel_atual() in niveis:
        return True
    st.error("Seu perfil não tem acesso a esta tela.")
    return False


def pode_emitir():
    return nivel_atual() in (NIVEL_OPERACAO, NIVEL_ADMIN)


def garantir_pode_emitir():
    # Barreira no servidor: nenhuma gravação de ficha passa sem perfil de emissão,
    # mesmo que um botão desabilitado seja acionado por fora da interface.
    if not pode_emitir():
        raise PermissionError("Seu perfil não pode gerar, alterar ou aprovar fichas.")


def aviso_somente_visualizacao():
    if not pode_emitir():
        st.info(
            "Modo visualização: seu perfil pode ver esta tela, mas não pode "
            "gerar, alterar, devolver ou aprovar fichas."
        )


ICONES_LUCIDE = {
    'menu_consulta': 'layout-dashboard', 'menu_graficos': 'chart-column',
    'menu_nova_ficha': 'file-plus',
    'menu_edicao': 'file-pen-line', 'menu_aprovar': 'circle-check-big',
    'menu_cadastrar_aeronave': 'plane', 'menu_criar_login': 'users',
    'menu_assinaturas': 'signature', 'menu_excluir': 'trash-2',
    'menu_gerou_aprovou': 'user-check',
    'menu_sair': 'log-out', 'atualizar_fichas': 'refresh-cw',
    'atualizar_graficos': 'refresh-cw',
}


@st.cache_resource
def css_icones_lucide(icones):
    """Ícones Lucide (SVG) aplicados aos botões pela chave de cada um."""
    # O dicionário entra como argumento para o cache mudar quando um ícone é incluído.
    import base64
    regras = []
    for chave, nome in dict(icones).items():
        caminho = os.path.join("assets", "icons", f"{nome}.svg")
        if not os.path.exists(caminho):
            continue
        with open(caminho, "rb") as arquivo:
            dados = base64.b64encode(arquivo.read()).decode()
        url = f"url('data:image/svg+xml;base64,{dados}')"
        regras.append(
            f".st-key-{chave} button p::before {{ content: ''; display: inline-block;"
            f" width: 1.05rem; height: 1.05rem; margin-right: .6rem; vertical-align: -0.2rem;"
            f" background-color: currentColor; -webkit-mask: {url} no-repeat center / contain;"
            f" mask: {url} no-repeat center / contain; }}"
        )
    return "<style>" + "\n".join(regras) + "</style>"


st.markdown(css_icones_lucide(tuple(ICONES_LUCIDE.items())), unsafe_allow_html=True)


def cabecalho(titulo, subtitulo=""):
    st.markdown(
        f'<div class="cabecalho"><h1>{titulo}</h1>'
        + (f"<p>{subtitulo}</p>" if subtitulo else "<p></p>")
        + "</div>",
        unsafe_allow_html=True,
    )

def safe_float(val):
    if pd.isna(val):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s = str(val).strip()
        if not s: return 0.0
        if '.' in s and ',' in s: s = s.replace('.', '').replace(',', '.')
        elif ',' in s: s = s.replace(',', '.')
        return float(s)
    except:
        return 0.0

def safe_str(val):
    return "" if pd.isna(val) else str(val).strip()


def safe_identifier(val):
    if pd.isna(val):
        return ""
    if pd.api.types.is_number(val) and float(val).is_integer():
        return str(int(val))
    return str(val).strip()


def normalizar_chave_ficha(valor):
    valor = safe_str(valor)
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return valor
    return str(int(numero)) if numero.is_integer() else format(numero, "g")


def momento_flaps_do_registro(registro):
    descricoes_flaps = [
        safe_str(registro.get(f"Additions Description {indice}", ""))
        for indice in range(1, 17)
    ]
    if not any("flap" in descricao.casefold() for descricao in descricoes_flaps):
        return None

    momento = safe_float(registro.get("Adction LAP", 0))
    if momento > 0:
        return momento

    for indice, descricao in enumerate(descricoes_flaps, start=1):
        if "flap" in descricao.casefold():
            momento = safe_float(registro.get(f"Additions arm {indice}", 0))
            if momento > 0:
                return momento
    return None


@st.cache_data
def buscar_momento_flaps_banco(prefixo, tipo_aeronave):
    df = carregar_fichas_para_momento_flaps(prefixo, tipo_aeronave)
    if df.empty or "Tipo_aeronave" not in df.columns:
        return 0.0

    df = df.copy()
    coluna_revisao = "Revisao" if "Revisao" in df.columns else "revisao"
    df["_pesagem_num"] = pd.to_numeric(df.get("Pesagem"), errors="coerce").fillna(-1)
    df["_revisao_num"] = pd.to_numeric(df.get(coluna_revisao), errors="coerce").fillna(-1)

    if prefixo and "Prefixo" in df.columns:
        historico_prefixo = df.loc[
            df["Prefixo"].astype(str).str.strip().str.casefold()
            == safe_str(prefixo).casefold()
        ].sort_values(["_pesagem_num", "_revisao_num"], ascending=False)
        for _, registro in historico_prefixo.iterrows():
            momento = momento_flaps_do_registro(registro)
            if momento is not None:
                return momento

    modelo = safe_str(tipo_aeronave).casefold()
    historico_modelo = df.loc[
        df["Tipo_aeronave"].astype(str).str.strip().str.casefold() == modelo
    ]
    momentos = [
        momento
        for _, registro in historico_modelo.iterrows()
        if (momento := momento_flaps_do_registro(registro)) is not None
    ]
    if not momentos:
        return 0.0

    frequencias = pd.Series(momentos).value_counts()
    maior_frequencia = frequencias.max()
    return float(max(frequencias[frequencias == maior_frequencia].index))


def converter_unidade_canela(unidade_key, unidade_anterior_key, chaves_valores):
    unidade_anterior = st.session_state.get(unidade_anterior_key, "in")
    unidade_atual = st.session_state.get(unidade_key, unidade_anterior)
    if unidade_atual != unidade_anterior:
        fator = 25.4 if unidade_anterior == "in" else 1 / 25.4
        for chave in chaves_valores:
            if chave in st.session_state:
                st.session_state[chave] *= fator
        st.session_state[unidade_anterior_key] = unidade_atual


def valor_canela_em_polegadas(valor, unidade):
    return valor / 25.4 if unidade == "mm" else valor


CAMPO_ULTIMA_PESAGEM_INFORMADA = 'Ultima pesagem informada'
CAMPO_ULTIMA_REVISAO_INFORMADA = 'Ultima revisao informada'


def ultima_ficha_gravada(registro):
    """Última pesagem/revisão digitadas à mão, para quando a ficha anterior
    não está no sistema ("" quando não há). As colunas antigas 'Pesagem
    anterior'/'Revisão Anterior' não são usadas: vieram do banco antigo sem
    padrão."""
    pesagem = normalizar_chave_ficha(registro.get(CAMPO_ULTIMA_PESAGEM_INFORMADA))
    if safe_float(pesagem) <= 0:
        return "", ""
    revisao = normalizar_chave_ficha(registro.get(CAMPO_ULTIMA_REVISAO_INFORMADA))
    return pesagem, revisao or "0"


def formatar_data_ficha(valor):
    """Datas da ficha no formato 20/10/2026 (o banco guarda 2026-10-20)."""
    texto = safe_str(valor)
    if re.match(r"^\d{4}-\d{2}-\d{2}", texto):
        try:
            return datetime.date.fromisoformat(texto[:10]).strftime("%d/%m/%Y")
        except ValueError:
            return texto
    return texto


def separar_campos_lopa(registro):
    valores = [
        safe_str(registro.get(" LOPA", registro.get("LOPA", ""))),
        safe_str(registro.get("Configuração LOPA ", registro.get("Configuração LOPA", ""))),
    ]
    valores = [
        valor for valor in valores
        if valor and valor.casefold() not in {"lopa", "configuração lopa"}
    ]
    codigo_lopa = next(
        (valor for valor in valores if valor.upper().startswith("GLP-")), ""
    )
    configuracao = next(
        (valor for valor in valores if not valor.upper().startswith("GLP-")), ""
    )
    return codigo_lopa, configuracao


LEVEL_CORRECTION_VALUES = {
    "-2": 3.0,
    "-1 7/8": 2.8,
    "-1 3/4": 2.7,
    "-1 5/8": 2.5,
    "-1 1/2": 2.3,
    "-1 3/8": 2.1,
    "-1 1/4": 1.9,
    "-1 1/8": 1.7,
    "-1": 1.5,
    "-7/8": 1.3,
    "-3/4": 1.1,
    "-5/8": 0.9,
    "-1/2": 0.8,
    "-3/8": 0.6,
    "-1/4": 0.4,
    "-1/8": 0.2,
    "0": 0.0,
    "1/8": -0.2,
    "1/4": -0.4,
    "3/8": -0.6,
    "1/2": -0.8,
    "5/8": -0.9,
    "3/4": -1.1,
    "7/8": -1.3,
    "1": -1.5,
    "1 1/8": -1.7,
    "1 1/4": -1.9,
    "1 3/8": -2.1,
    "1 1/2": -2.3,
    "1 5/8": -2.5,
    "1 3/4": -2.7,
    "1 7/8": -2.8,
    "2": -3.0,
}
LEVEL_CORRECTION_LABEL = "Level Correction (Value added to CG)"


def normalizar_angulo_level_correction(angulo):
    angulo = safe_str(angulo).replace("\xa0", " ")
    angulo = " ".join(angulo.split())
    if angulo.endswith(".0") and angulo[:-2] in LEVEL_CORRECTION_VALUES:
        angulo = angulo[:-2]
    return angulo if angulo in LEVEL_CORRECTION_VALUES else ""


def calcular_level_correction(angulo, peso_base):
    angulo = normalizar_angulo_level_correction(angulo)
    if not angulo:
        return 0.0, 0.0, None

    valor_correspondente = LEVEL_CORRECTION_VALUES[angulo]
    if angulo.startswith("-"):
        return valor_correspondente, valor_correspondente * peso_base, "additions"
    return valor_correspondente, -valor_correspondente * peso_base, "deductions"

def protegido(tela):
    """Mostra uma mensagem clara em vez de derrubar o app se algo falhar."""
    @functools.wraps(tela)
    def executar(*args, **kwargs):
        try:
            return tela(*args, **kwargs)
        except Exception as erro:
            st.cache_data.clear()
            st.error(
                "Algo deu errado nesta tela e nada foi perdido do que já estava "
                "salvo. Tente novamente; se o erro continuar, envie o detalhe "
                "abaixo para a Engenharia."
            )
            with st.expander("Detalhe do erro"):
                st.exception(erro)
    return executar


# 2. CONEXÃO E CARREGAMENTO DOS BANCOS DE DADOS
@st.cache_resource
def obter_conexao_postgresql():
    configuracao = st.secrets.get("connections", {}).get("postgresql", {})
    url = configuracao.get("url") if configuracao else None
    if url:
        try:
            make_url(url)
        except ValueError:
            raise ValueError(
                "A URL em [connections.postgresql].url é inválida. "
                "Confira se está no formato esperado e se a porta após o host "
                "é numérica, como 5432 ou 6543."
            ) from None
    # prepare_threshold=None evita erros de prepared statement no pooler.
    # Conexões derrubadas pelo Supabase são refeitas em ConexaoPostgreSQL.execute,
    # sem o custo de um "ping" antes de cada consulta.
    return st.connection(
        "postgresql",
        type="sql",
        pool_recycle=240,
        pool_size=5,
        max_overflow=5,
        connect_args={"prepare_threshold": None},
    )


class LinhaBanco:
    def __init__(self, linha):
        self._valores = tuple(linha)
        self._mapeamento = dict(linha._mapping)

    def __getitem__(self, chave):
        if isinstance(chave, int):
            return self._valores[chave]
        return self._mapeamento[chave]

    def __iter__(self):
        return iter(self._valores)

    def keys(self):
        return self._mapeamento.keys()


class ResultadoBanco:
    def __init__(self, resultado):
        self._resultado = resultado
        self.rowcount = resultado.rowcount

    def fetchone(self):
        linha = self._resultado.fetchone()
        return LinhaBanco(linha) if linha is not None else None

    def fetchall(self):
        return [LinhaBanco(linha) for linha in self._resultado.fetchall()]

    def __iter__(self):
        return iter(self.fetchall())


class ConexaoPostgreSQL:
    def __init__(self):
        self._sessao = obter_conexao_postgresql().session
        self._primeira_consulta = True

    @property
    def row_factory(self):
        return None

    @row_factory.setter
    def row_factory(self, _valor):
        pass

    @staticmethod
    def _preparar(sql, parametros):
        if parametros is None or isinstance(parametros, dict):
            return sql, parametros or {}
        partes = sql.split("?")
        if len(partes) - 1 != len(parametros):
            raise ValueError("A quantidade de parâmetros SQL não corresponde aos placeholders.")
        consulta = partes[0]
        valores = {}
        for indice, valor in enumerate(parametros):
            nome = f"param_{indice}"
            consulta += f":{nome}{partes[indice + 1]}"
            valores[nome] = valor
        return consulta, valores

    def execute(self, sql, parametros=None):
        consulta, valores = self._preparar(sql, parametros)
        primeira = self._primeira_consulta
        self._primeira_consulta = False
        try:
            return ResultadoBanco(self._sessao.execute(text(consulta), valores))
        except DBAPIError as erro:
            # Conexão encerrada pelo servidor: refaz uma vez se nada foi
            # executado ainda nesta transação.
            if not (primeira and erro.connection_invalidated):
                raise
            self._sessao.rollback()
            self._sessao.close()
            self._sessao = obter_conexao_postgresql().session
            return ResultadoBanco(self._sessao.execute(text(consulta), valores))

    def executemany(self, sql, sequencias):
        if not sequencias:
            return None
        consulta, _ = self._preparar(sql, sequencias[0])
        valores = [self._preparar(sql, linha)[1] for linha in sequencias]
        return ResultadoBanco(self._sessao.execute(text(consulta), valores))

    def __enter__(self):
        return self

    def __exit__(self, tipo_erro, _erro, _rastreio):
        try:
            if tipo_erro is None:
                self._sessao.commit()
            else:
                self._sessao.rollback()
        finally:
            self._sessao.close()

    def close(self):
        self._sessao.commit()
        self._sessao.close()


def conectar_banco():
    return ConexaoPostgreSQL()


def normalizar_valor_json(valor):
    if valor is None:
        return None
    if isinstance(valor, dict):
        return {
            str(chave): normalizar_valor_json(item)
            for chave, item in valor.items()
        }
    if isinstance(valor, (list, tuple)):
        return [normalizar_valor_json(item) for item in valor]
    if hasattr(valor, "item"):
        return normalizar_valor_json(valor.item())
    if pd.isna(valor):
        return None
    if isinstance(valor, (datetime.date, datetime.datetime)):
        return valor.isoformat()
    return valor


def preparar_registro_json(registro):
    dados = dict(registro)
    revisao = dados.pop("revisao", None)
    if safe_str(dados.get("Revisao", "")) == "":
        dados["Revisao"] = revisao
    for chave in ("Pesagem", "Revisao"):
        dados[chave] = normalizar_chave_ficha(dados.get(chave, ""))
    dados["Prefixo"] = safe_str(dados.get("Prefixo", ""))
    dados.pop("Pesagem_num", None)
    dados.pop("Revisao_num", None)
    return json.dumps(
        normalizar_valor_json(dados),
        ensure_ascii=False,
        allow_nan=False,
    )


def buscar_id_ficha(conn, prefixo, pesagem, revisao):
    registros = conn.execute(
        """SELECT id, dados->>'Pesagem' AS pesagem,
                  COALESCE(dados->>'Revisao', dados->>'revisao') AS revisao
           FROM pesagens
           WHERE lower(btrim(dados->>'Prefixo')) = lower(btrim(?))
           ORDER BY id DESC""",
        (safe_str(prefixo),),
    ).fetchall()
    return [
        item["id"]
        for item in registros
        if normalizar_chave_ficha(item["pesagem"]) == normalizar_chave_ficha(pesagem)
        and normalizar_chave_ficha(item["revisao"]) == normalizar_chave_ficha(revisao)
    ]


def bloquear_prefixo(conn, prefixo):
    # Serializa gravações da mesma aeronave para evitar fichas duplicadas
    # quando o botão é clicado duas vezes ou duas pessoas salvam juntas.
    conn.execute(
        "SELECT pg_advisory_xact_lock(hashtext(lower(btrim(?))))",
        (safe_str(prefixo),),
    )


def inserir_ficha(registro):
    garantir_pode_emitir()
    with conectar_banco() as conn:
        bloquear_prefixo(conn, registro.get("Prefixo", ""))
        if buscar_id_ficha(
            conn,
            registro.get("Prefixo", ""),
            registro.get("Pesagem", ""),
            registro.get("Revisao", registro.get("revisao", "")),
        ):
            return False
        conn.execute(
            "INSERT INTO pesagens (dados) VALUES (CAST(? AS JSONB))",
            (preparar_registro_json(registro),),
        )
    invalidar_cache_fichas()
    return True


CAMPOS_FORA_DO_HISTORICO = {
    "Prefixo", "Pesagem", "Revisao", "Pesagem_num", "Revisao_num", "revisao",
}


def valor_para_comparar(valor):
    # Campo vazio e zero contam como iguais: o formulário grava 0 nos campos
    # numéricos que estavam em branco, e isso não é uma alteração de dado.
    if isinstance(valor, (list, dict)):
        return json.dumps(valor, ensure_ascii=False, sort_keys=True)
    texto = safe_str(valor)
    if texto == "":
        return "0"
    # A mesma data pode estar gravada como 29/07/2025 ou 2025-07-29.
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            pass
    try:
        numero = float(texto.replace(",", "."))
    except ValueError:
        return texto
    return format(numero, "g") if numero == numero else texto


def valor_para_exibir(valor):
    if isinstance(valor, (list, dict)):
        return json.dumps(valor, ensure_ascii=False)
    texto = safe_str(valor)
    if texto == "":
        return "—"
    try:
        return format(float(texto.replace(",", ".")), "g")
    except ValueError:
        return texto


def comparar_registros(anterior, novo):
    alteracoes = []
    for campo in list(anterior) + [c for c in novo if c not in anterior]:
        if campo in CAMPOS_FORA_DO_HISTORICO:
            continue
        antes, depois = anterior.get(campo), novo.get(campo)
        if valor_para_comparar(antes) != valor_para_comparar(depois):
            alteracoes.append({
                "campo": campo,
                "antes": valor_para_exibir(antes),
                "depois": valor_para_exibir(depois),
            })
    return alteracoes


def atualizar_ficha(prefixo, pesagem, revisao, registro):
    """Grava a ficha e devolve a lista de campos alterados (None se não achou)."""
    garantir_pode_emitir()
    dados_novos = preparar_registro_json(registro)
    with conectar_banco() as conn:
        bloquear_prefixo(conn, prefixo)
        ids = buscar_id_ficha(conn, prefixo, pesagem, revisao)
        if not ids:
            return None

        anterior = conn.execute(
            "SELECT dados FROM pesagens WHERE id = ?", (ids[0],)
        ).fetchone()
        resultado = conn.execute(
            "UPDATE pesagens SET dados = CAST(? AS JSONB) WHERE id = ?",
            (dados_novos, ids[0]),
        )
    if resultado.rowcount != 1:
        return None
    invalidar_cache_fichas()
    return comparar_registros(
        anterior["dados"] if anterior else {}, json.loads(dados_novos)
    )


@st.cache_data
def carregar_dados_banco():
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT dados FROM pesagens
               WHERE NULLIF(btrim(dados->>'Prefixo'), '') IS NOT NULL
               ORDER BY id"""
        ).fetchall()
    return dataframe_fichas(registros)


def dataframe_fichas(registros):
    df = (
        pd.DataFrame([registro[0] for registro in registros])
        if registros
        else pd.DataFrame(columns=["Prefixo", "Pesagem", "Revisao"])
    )
    if 'revisao' in df.columns:
        if 'Revisao' in df.columns:
            df['Revisao'] = df['Revisao'].where(
                df['Revisao'].map(safe_str) != "", df['revisao']
            )
            df = df.drop(columns=['revisao'])
        else:
            df = df.rename(columns={'revisao': 'Revisao'})
    for coluna in ('Pesagem', 'Revisao'):
        if coluna in df.columns:
            df[coluna] = df[coluna].map(normalizar_chave_ficha)
    if 'Prefixo' in df.columns:
        df = df.loc[df['Prefixo'].map(safe_str) != ""].copy()
        df['Prefixo'] = df['Prefixo'].map(safe_str)
    return df.reset_index(drop=True)


@st.cache_data
def carregar_prefixos_fichas():
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT DISTINCT dados->>'Prefixo' AS prefixo
               FROM pesagens
               WHERE NULLIF(btrim(dados->>'Prefixo'), '') IS NOT NULL
               ORDER BY prefixo"""
        ).fetchall()
    return [safe_str(registro["prefixo"]) for registro in registros]


@st.cache_data
def carregar_fichas_prefixo(prefixo):
    return dataframe_fichas(consultar_registros_prefixo(prefixo))


def consultar_registros_prefixo(prefixo):
    with conectar_banco() as conn:
        return conn.execute(
            """SELECT dados FROM pesagens
               WHERE lower(btrim(dados->>'Prefixo')) = lower(btrim(?))
               ORDER BY id""",
            (safe_str(prefixo),),
        ).fetchall()


@st.cache_data
def carregar_linha_ficha(prefixo, pesagem, revisao):
    df = carregar_fichas_prefixo(prefixo)
    if not {"Pesagem", "Revisao"}.issubset(df.columns):
        return None
    linhas = df.loc[
        (df["Pesagem"].map(normalizar_chave_ficha)
         == normalizar_chave_ficha(pesagem))
        & (df["Revisao"].map(normalizar_chave_ficha)
           == normalizar_chave_ficha(revisao))
    ]
    return linhas.iloc[0] if not linhas.empty else None


@st.cache_data
def carregar_fichas_para_momento_flaps(prefixo, tipo_aeronave):
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT dados FROM pesagens
               WHERE lower(btrim(dados->>'Prefixo')) = lower(btrim(?))
                  OR lower(btrim(dados->>'Tipo_aeronave')) = lower(btrim(?))
               ORDER BY id DESC""",
            (safe_str(prefixo), safe_str(tipo_aeronave)),
        ).fetchall()
    return dataframe_fichas(registros)


def invalidar_cache_fichas():
    carregar_dados_banco.clear()
    carregar_linha_ficha.clear()
    listar_fichas_pendentes.clear()
    carregar_presets_lopa.clear()
    carregar_prefixos_fichas.clear()
    carregar_fichas_prefixo.clear()
    carregar_fichas_para_momento_flaps.clear()
    buscar_momento_flaps_banco.clear()
    buscar_fluxo_ficha.clear()
    montar_resumo_frota.clear()
    montar_historico_frota.clear()
    carregar_campos_com_erro.clear()
    carregar_historico_ficha.clear()
    carregar_assinatura_usuario.clear()
    listar_fichas_devolvidas_usuario.clear()
    listar_fichas_gerou_aprovou.clear()


@st.cache_data
def carregar_tipos_aeronave():
    excel_path = 'Cópia de Ficha_Pesagem_v2.xlsm'
    if not os.path.exists(excel_path):
        st.warning(f"Planilha não encontrada: '{excel_path}'.")
        return {}
    try:
        df_tipos = pd.read_excel(excel_path, sheet_name='Sheet1')
        
        col_modelo = 'MODEL' if 'MODEL' in df_tipos.columns else df_tipos.columns[0]
        col_prefixo = 'VRG REG.' if 'VRG REG.' in df_tipos.columns else df_tipos.columns[1]
        col_armnose = 'ARMNOSE' if 'ARMNOSE' in df_tipos.columns else None
        
        col_vrbl = next((c for c in df_tipos.columns if 'VRBL' in str(c).upper() or 'VRBL NUMBER' in str(c).upper()), None)
        nomes_colunas = {
            coluna: " ".join(str(coluna).replace("\n", " ").split()).upper()
            for coluna in df_tipos.columns
        }
        col_serial = next(
            (
                coluna for coluna, nome in nomes_colunas.items()
                if nome in {"S/N", "SERIAL", "SERIAL NUMBER"}
            ),
            None,
        )
        if col_serial is None:
            col_serial = next(
                (
                    coluna for coluna, nome in nomes_colunas.items()
                    if "SERIAL" in nome and "ENGINE" not in nome
                ),
                None,
            )
        col_line = next((c for c in df_tipos.columns if 'LINE' in str(c).upper()), None)
        
        dict_aero = {}
        for _, row in df_tipos.iterrows():
            prefixo = str(row[col_prefixo]).strip()
            if prefixo and prefixo != 'nan':
                modelo = str(row[col_modelo]).strip()
                arm_nose = 93.0000
                if col_armnose and pd.notna(row[col_armnose]):
                    try: arm_nose = float(row[col_armnose])
                    except: pass
                
                dict_aero[prefixo] = {
                    'modelo': modelo, 
                    'armnose': arm_nose,
                    'vrbl': safe_identifier(row[col_vrbl]) if col_vrbl else "",
                    'serial': safe_identifier(row[col_serial]) if col_serial else "",
                    'line': safe_identifier(row[col_line]) if col_line else ""
                }
        return dict_aero
    except Exception as e:
        return {}

dict_tipos_aeronave = carregar_tipos_aeronave()


def normalizar_prefixo(prefixo):
    return "".join(safe_str(prefixo).upper().split())


@st.cache_data
def carregar_aeronaves_cadastradas():
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT prefixo, modelo, serial, vrbl, line, armnose,
                      cadastrado_por, cadastrado_em
               FROM aeronaves ORDER BY prefixo"""
        ).fetchall()
    return {
        registro['prefixo']: {
            'modelo': registro['modelo'],
            'armnose': (
                float(registro['armnose'])
                if registro['armnose'] is not None
                else armnose_padrao_modelo(registro['modelo'])
            ),
            'vrbl': safe_str(registro['vrbl']),
            'serial': safe_str(registro['serial']),
            'line': safe_str(registro['line']),
            'cadastrado_por': safe_str(registro['cadastrado_por']),
            'cadastrado_em': safe_str(registro['cadastrado_em']),
        }
        for registro in registros
    }


def armnose_padrao_modelo(modelo):
    valores = [
        dados['armnose']
        for dados in dict_tipos_aeronave.values()
        if dados.get('modelo') == modelo
    ]
    if not valores:
        return 93.0
    return float(pd.Series(valores).mode().max())


def cadastrar_aeronave(prefixo, modelo, serial, vrbl, line, armnose, usuario):
    with conectar_banco() as conn:
        conn.execute(
            """INSERT INTO aeronaves
               (prefixo, modelo, serial, vrbl, line, armnose,
                cadastrado_por, cadastrado_em)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (prefixo, modelo, serial, vrbl, line, float(armnose), usuario,
             datetime.datetime.now(datetime.timezone.utc).isoformat()),
        )
    carregar_aeronaves_cadastradas.clear()


def remover_aeronave_cadastrada(prefixo):
    with conectar_banco() as conn:
        conn.execute('DELETE FROM aeronaves WHERE prefixo = ?', (prefixo,))
    carregar_aeronaves_cadastradas.clear()


def obter_senha_admin_inicial():
    senha = os.environ.get("INITIAL_ADMIN_PASSWORD")
    if senha:
        return senha
    try:
        return st.secrets.get("INITIAL_ADMIN_PASSWORD")
    except Exception:
        return None


@st.cache_resource
def inicializar_controle_acesso():
    with conectar_banco() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS pesagens (
                id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                dados JSONB NOT NULL
            )
        ''')
        conn.execute('''
            CREATE INDEX IF NOT EXISTS pesagens_prefixo_idx
            ON pesagens (lower(btrim(dados->>'Prefixo')))
        ''')
        conn.execute('''
            CREATE INDEX IF NOT EXISTS pesagens_tipo_aeronave_idx
            ON pesagens (lower(btrim(dados->>'Tipo_aeronave')))
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS usuarios (
                usuario TEXT PRIMARY KEY,
                nome TEXT NOT NULL,
                senha_hash TEXT NOT NULL,
                salt TEXT NOT NULL,
                nivel_acesso INTEGER NOT NULL,
                criado_em TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE UNIQUE INDEX IF NOT EXISTS usuarios_usuario_casefold_uidx
            ON usuarios (lower(usuario))
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS fluxo_fichas (
                prefixo TEXT NOT NULL,
                pesagem TEXT NOT NULL,
                revisao TEXT NOT NULL,
                gerado_por TEXT,
                aprovado_por TEXT,
                criado_em TEXT NOT NULL,
                aprovado_em TEXT,
                PRIMARY KEY (prefixo, pesagem, revisao)
            )
        ''')
        # Fichas migradas que estavam pendentes quando a Aprovação foi limpa:
        # ficam só em "Gerou e aprovou" até alguém gerar ou editar de novo.
        conn.execute('''
            CREATE TABLE IF NOT EXISTS fichas_fora_aprovacao (
                prefixo TEXT NOT NULL,
                pesagem TEXT NOT NULL,
                revisao TEXT NOT NULL,
                separado_em TEXT NOT NULL,
                PRIMARY KEY (prefixo, pesagem, revisao)
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS ajustes_aplicados (
                nome TEXT PRIMARY KEY,
                aplicado_em TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS campos_com_erro (
                prefixo TEXT NOT NULL,
                pesagem TEXT NOT NULL,
                revisao TEXT NOT NULL,
                campo TEXT NOT NULL,
                descricao TEXT NOT NULL,
                marcado_por TEXT,
                marcado_em TEXT NOT NULL,
                PRIMARY KEY (prefixo, pesagem, revisao, campo)
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS historico_fichas (
                id BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
                prefixo TEXT NOT NULL,
                pesagem TEXT NOT NULL,
                revisao TEXT NOT NULL,
                acao TEXT NOT NULL,
                usuario TEXT,
                nome TEXT,
                registrado_em TEXT NOT NULL
            )
        ''')
        conn.execute(
            "ALTER TABLE historico_fichas ADD COLUMN IF NOT EXISTS alteracoes JSONB"
        )
        conn.execute('''
            CREATE INDEX IF NOT EXISTS historico_fichas_chave_idx
            ON historico_fichas (prefixo, pesagem, revisao)
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS opcoes_cadastradas (
                categoria TEXT NOT NULL,
                valor TEXT NOT NULL,
                cadastrado_por TEXT,
                cadastrado_em TEXT NOT NULL,
                PRIMARY KEY (categoria, valor)
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS aeronaves (
                prefixo TEXT PRIMARY KEY,
                modelo TEXT NOT NULL,
                serial TEXT,
                vrbl TEXT,
                line TEXT,
                armnose DOUBLE PRECISION NOT NULL,
                cadastrado_por TEXT,
                cadastrado_em TEXT NOT NULL
            )
        ''')
        # Bancos que já tinham a tabela da versão anterior (só prefixo, modelo,
        # serial, vrbl e line) recebem as colunas novas.
        conn.execute('''
            ALTER TABLE aeronaves
                ADD COLUMN IF NOT EXISTS serial TEXT,
                ADD COLUMN IF NOT EXISTS vrbl TEXT,
                ADD COLUMN IF NOT EXISTS line TEXT,
                ADD COLUMN IF NOT EXISTS armnose DOUBLE PRECISION,
                ADD COLUMN IF NOT EXISTS cadastrado_por TEXT,
                ADD COLUMN IF NOT EXISTS cadastrado_em TEXT
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS assinaturas_usuarios (
                usuario TEXT PRIMARY KEY,
                imagem_png BYTEA NOT NULL,
                atualizado_em TEXT NOT NULL,
                FOREIGN KEY (usuario) REFERENCES usuarios(usuario) ON DELETE CASCADE
            )
        ''')
        # A conta inicial da Engenharia passa a ser Administrador (nível 3).
        conn.execute(
            "UPDATE usuarios SET nivel_acesso = 3 WHERE usuario = 'engenharia' AND nivel_acesso = 1"
        )
        usuario_existente = conn.execute('SELECT 1 FROM usuarios LIMIT 1').fetchone()
        if not usuario_existente:
            senha_inicial = obter_senha_admin_inicial()
            if not senha_inicial:
                st.error(
                    "Configure INITIAL_ADMIN_PASSWORD em Streamlit Secrets "
                    "para criar o primeiro usuário de Engenharia."
                )
            else:
                salt = secrets.token_hex(16)
                senha_hash = hashlib.pbkdf2_hmac(
                    'sha256', senha_inicial.encode('utf-8'), bytes.fromhex(salt), 600000
                ).hex()
                conn.execute(
                    '''INSERT INTO usuarios
                       (usuario, nome, senha_hash, salt, nivel_acesso, criado_em)
                       VALUES (?, ?, ?, ?, ?, ?)''',
                    ('engenharia', 'Engenharia GOL', senha_hash, salt, 3,
                     datetime.datetime.now(datetime.timezone.utc).isoformat())
                )
def autenticar_usuario(usuario, senha):
    with conectar_banco() as conn:
        registro = conn.execute(
            'SELECT * FROM usuarios WHERE lower(usuario) = lower(?)',
            (usuario.strip(),),
        ).fetchone()
    if not registro:
        return None
    senha_hash = hashlib.pbkdf2_hmac(
        'sha256', senha.encode('utf-8'), bytes.fromhex(registro['salt']), 600000
    ).hex()
    if not hmac.compare_digest(senha_hash, registro['senha_hash']):
        senha_inicial = obter_senha_admin_inicial()
        if (
            registro['usuario'].casefold() != 'engenharia'
            or not senha_inicial
            or not hmac.compare_digest(senha, senha_inicial)
        ):
            return None

        salt = secrets.token_hex(16)
        senha_hash = hashlib.pbkdf2_hmac(
            'sha256', senha_inicial.encode('utf-8'), bytes.fromhex(salt), 600000
        ).hex()
        with conectar_banco() as conn:
            conn.execute(
                'UPDATE usuarios SET senha_hash = ?, salt = ? WHERE usuario = ?',
                (senha_hash, salt, registro['usuario']),
            )
        invalidar_cache_fichas()
        registro = dict(registro)
        registro['senha_hash'] = senha_hash
        registro['salt'] = salt
    return dict(registro)

def criar_usuario(nome, usuario, senha, nivel_acesso):
    salt = secrets.token_hex(16)
    senha_hash = hashlib.pbkdf2_hmac(
        'sha256', senha.encode('utf-8'), bytes.fromhex(salt), 600000
    ).hex()
    with conectar_banco() as conn:
        conn.execute(
            '''INSERT INTO usuarios
               (usuario, nome, senha_hash, salt, nivel_acesso, criado_em)
               VALUES (?, ?, ?, ?, ?, ?)''',
            (usuario.strip(), nome.strip(), senha_hash, salt, nivel_acesso,
             datetime.datetime.now(datetime.timezone.utc).isoformat())
        )
    invalidar_cache_fichas()


def redefinir_senha(usuario, senha):
    salt = secrets.token_hex(16)
    senha_hash = hashlib.pbkdf2_hmac(
        'sha256', senha.encode('utf-8'), bytes.fromhex(salt), 600000
    ).hex()
    with conectar_banco() as conn:
        conn.execute(
            'UPDATE usuarios SET senha_hash = ?, salt = ? WHERE usuario = ?',
            (senha_hash, salt, usuario),
        )


def listar_usuarios_assinaturas():
    with conectar_banco() as conn:
        registros = conn.execute(
            '''SELECT usuarios.usuario, usuarios.nome, usuarios.nivel_acesso,
                      assinaturas_usuarios.atualizado_em
               FROM usuarios
               LEFT JOIN assinaturas_usuarios
                   ON assinaturas_usuarios.usuario = usuarios.usuario
               ORDER BY lower(usuarios.nome)'''
        ).fetchall()
    return [dict(registro) for registro in registros]


@st.cache_data
def carregar_assinatura_usuario(usuario):
    if not usuario:
        return None
    with conectar_banco() as conn:
        registro = conn.execute(
            'SELECT imagem_png FROM assinaturas_usuarios WHERE usuario = ?',
            (usuario,),
        ).fetchone()
    return bytes(registro[0]) if registro else None


def salvar_assinatura_usuario(usuario, imagem_png):
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with conectar_banco() as conn:
        conn.execute(
            '''INSERT INTO assinaturas_usuarios (usuario, imagem_png, atualizado_em)
               VALUES (?, ?, ?)
               ON CONFLICT(usuario) DO UPDATE SET
                   imagem_png = excluded.imagem_png,
                   atualizado_em = excluded.atualizado_em''',
            (usuario, imagem_png, agora),
        )
    invalidar_cache_fichas()


def remover_assinatura_usuario(usuario):
    with conectar_banco() as conn:
        conn.execute(
            'DELETE FROM assinaturas_usuarios WHERE usuario = ?',
            (usuario,),
        )
    invalidar_cache_fichas()


def normalizar_imagem_assinatura(imagem):
    if len(imagem) > 5 * 1024 * 1024:
        raise ValueError("A imagem deve ter no máximo 5 MB.")

    try:
        with PillowImage.open(io.BytesIO(imagem)) as arquivo:
            if arquivo.format not in {"PNG", "JPEG"}:
                raise ValueError("Use uma imagem PNG ou JPEG.")
            if arquivo.width * arquivo.height > 20_000_000:
                raise ValueError("A imagem excede o limite de resolução permitido.")
            imagem_convertida = arquivo.convert("RGBA")
            imagem_convertida.thumbnail((1200, 400))
            saida = io.BytesIO()
            imagem_convertida.save(saida, format="PNG")
    except UnidentifiedImageError as erro:
        raise ValueError("O arquivo enviado não é uma imagem válida.") from erro
    except OSError as erro:
        raise ValueError("Não foi possível ler a imagem enviada.") from erro

    return saida.getvalue()


def inserir_historico(
    conn, prefixo, pesagem, revisao, acao, usuario, agora, alteracoes=None
):
    # Guarda também o nome, para o histórico continuar legível se o usuário
    # for renomeado ou removido depois.
    conn.execute(
        '''INSERT INTO historico_fichas
           (prefixo, pesagem, revisao, acao, usuario, nome, registrado_em,
            alteracoes)
           VALUES (?, ?, ?, ?, ?,
                   (SELECT nome FROM usuarios WHERE usuario = ?), ?,
                   CAST(? AS JSONB))''',
        (
            safe_str(prefixo),
            normalizar_chave_ficha(pesagem),
            normalizar_chave_ficha(revisao),
            acao,
            usuario,
            usuario,
            agora,
            json.dumps(alteracoes, ensure_ascii=False) if alteracoes else None,
        ),
    )


@st.cache_data
def carregar_historico_ficha(prefixo, pesagem, revisao):
    with conectar_banco() as conn:
        registros = conn.execute(
            '''SELECT historico_fichas.acao, historico_fichas.usuario,
                      COALESCE(historico_fichas.nome, usuarios.nome) AS nome,
                      historico_fichas.registrado_em,
                      historico_fichas.alteracoes
               FROM historico_fichas
               LEFT JOIN usuarios ON usuarios.usuario = historico_fichas.usuario
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?
               ORDER BY historico_fichas.id DESC''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
            ),
        ).fetchall()
    # Uma linha por campo alterado; ações sem campos (gerar, aprovar) ficam
    # em uma linha só.
    linhas = []
    for registro in registros:
        base = {
            "Data e hora (Brasília)": formatar_data_hora_brasilia(registro["registrado_em"]),
            "Ação": registro["acao"],
            "Por": registro["nome"] or registro["usuario"] or "—",
        }
        alteracoes = registro["alteracoes"] or []
        if registro["acao"] in ("Alterada", "Corrigida") and not alteracoes:
            alteracoes = [{"campo": "Nenhum dado alterado", "antes": "", "depois": ""}]
        for alteracao in alteracoes or [{"campo": "", "antes": "", "depois": ""}]:
            linhas.append({
                **base,
                "Campo": alteracao["campo"],
                "Antes": alteracao["antes"],
                "Depois": alteracao["depois"],
            })
    return linhas


def registrar_ficha(prefixo, pesagem, revisao, usuario, acao="Gerada", alteracoes=None):
    # Quem gera ou edita a ficha passa a ser o emissor (e a assinatura dele vai
    # na ficha); a aprovação anterior é desfeita e ela volta para a Aprovação.
    garantir_pode_emitir()
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with conectar_banco() as conn:
        conn.execute(
            '''DELETE FROM fichas_fora_aprovacao
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
            ),
        )
        inserir_historico(
            conn, prefixo, pesagem, revisao, acao, usuario, agora, alteracoes
        )
        conn.execute(
            '''INSERT INTO fluxo_fichas
               (prefixo, pesagem, revisao, gerado_por, criado_em)
               VALUES (?, ?, ?, ?, ?)
               ON CONFLICT(prefixo, pesagem, revisao) DO UPDATE SET
                   gerado_por = excluded.gerado_por,
                   aprovado_por = NULL,
                   criado_em = excluded.criado_em,
                   aprovado_em = NULL''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
                usuario,
                agora,
            )
        )
    invalidar_cache_fichas()

@st.cache_data
def buscar_fluxo_ficha(prefixo, pesagem, revisao):
    with conectar_banco() as conn:
        registro = conn.execute(
            '''SELECT fluxo_fichas.gerado_por AS gerador_usuario,
                      gerador.nome AS gerador_nome,
                      fluxo_fichas.aprovado_por AS aprovador_usuario,
                      aprovador.nome AS aprovador_nome,
                      fluxo_fichas.aprovado_em
               FROM fluxo_fichas
               LEFT JOIN usuarios AS gerador
                   ON gerador.usuario = fluxo_fichas.gerado_por
               LEFT JOIN usuarios AS aprovador
                   ON aprovador.usuario = fluxo_fichas.aprovado_por
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
            )
        ).fetchone()
    if not registro:
        return {
            'gerador_usuario': None,
            'gerador_nome': 'Não registrado',
            'aprovador_usuario': None,
            'aprovador_nome': 'Pendente',
            'aprovado_em': None,
        }
    return {
        'gerador_usuario': registro['gerador_usuario'],
        'gerador_nome': registro['gerador_nome'] or 'Não registrado',
        'aprovador_usuario': registro['aprovador_usuario'],
        'aprovador_nome': registro['aprovador_nome'] or 'Pendente',
        'aprovado_em': registro['aprovado_em'],
    }

def carregar_fluxos_fichas():
    with conectar_banco() as conn:
        registros = conn.execute(
            '''SELECT fluxo_fichas.prefixo, fluxo_fichas.pesagem,
                      fluxo_fichas.revisao,
                      fluxo_fichas.gerado_por AS gerador_usuario,
                      gerador.nome AS gerador_nome,
                      fluxo_fichas.aprovado_por AS aprovador_usuario,
                      aprovador.nome AS aprovador_nome
               FROM fluxo_fichas
               LEFT JOIN usuarios AS gerador
                   ON gerador.usuario = fluxo_fichas.gerado_por
               LEFT JOIN usuarios AS aprovador
                   ON aprovador.usuario = fluxo_fichas.aprovado_por'''
        ).fetchall()
    return {
        (
            safe_str(registro['prefixo']),
            normalizar_chave_ficha(registro['pesagem']),
            normalizar_chave_ficha(registro['revisao']),
        ): {
            'gerador_usuario': registro['gerador_usuario'],
            'gerador_nome': registro['gerador_nome'] or 'Não registrado',
            'aprovador_usuario': registro['aprovador_usuario'],
            'aprovador_nome': registro['aprovador_nome'] or 'Pendente',
        }
        for registro in registros
    }


@st.cache_data
def carregar_campos_com_erro(prefixo, pesagem, revisao):
    with conectar_banco() as conn:
        registros = conn.execute(
            '''SELECT campos_com_erro.campo, campos_com_erro.descricao,
                      campos_com_erro.marcado_por, usuarios.nome AS marcador_nome
               FROM campos_com_erro
               LEFT JOIN usuarios ON usuarios.usuario = campos_com_erro.marcado_por
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?
               ORDER BY campo''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
            ),
        ).fetchall()
    return [dict(registro) for registro in registros]


def carregar_chaves_fichas():
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT dados->>'Prefixo' AS prefixo,
                      dados->>'Pesagem' AS pesagem,
                      COALESCE(dados->>'Revisao', dados->>'revisao') AS revisao
               FROM pesagens
               WHERE NULLIF(btrim(dados->>'Prefixo'), '') IS NOT NULL"""
        ).fetchall()
    return [dict(registro) for registro in registros]


@st.cache_data
def listar_fichas_devolvidas_usuario(usuario):
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT campos_com_erro.prefixo, campos_com_erro.pesagem,
                      campos_com_erro.revisao, campos_com_erro.campo,
                      campos_com_erro.descricao,
                      fluxo_fichas.gerado_por AS gerador_usuario,
                      usuarios.nome AS gerador_nome
               FROM campos_com_erro
               JOIN fluxo_fichas
                 ON fluxo_fichas.prefixo = campos_com_erro.prefixo
                AND fluxo_fichas.pesagem = campos_com_erro.pesagem
                AND fluxo_fichas.revisao = campos_com_erro.revisao
               LEFT JOIN usuarios
                 ON usuarios.usuario = fluxo_fichas.gerado_por
               WHERE fluxo_fichas.gerado_por = ?
                 AND fluxo_fichas.aprovado_por IS NULL
               ORDER BY campos_com_erro.prefixo,
                        campos_com_erro.pesagem,
                        campos_com_erro.revisao,
                        campos_com_erro.campo""",
            (usuario,),
        ).fetchall()

    fichas = {}
    for registro in registros:
        chave = (
            safe_str(registro["prefixo"]),
            normalizar_chave_ficha(registro["pesagem"]),
            normalizar_chave_ficha(registro["revisao"]),
        )
        ficha = fichas.setdefault(
            chave,
            {
                "chave": chave,
                "pesagem_num": safe_float(chave[1]),
                "revisao_num": safe_float(chave[2]),
                "fluxo": {
                    "gerador_usuario": registro["gerador_usuario"],
                    "gerador_nome": registro["gerador_nome"] or "Não registrado",
                },
                "erros": [],
            },
        )
        ficha["erros"].append(
            {"campo": registro["campo"], "descricao": registro["descricao"]}
        )
    return list(fichas.values())


@st.cache_data
def listar_fichas_pendentes():
    fluxos = carregar_fluxos_fichas()
    fora_aprovacao = carregar_fichas_fora_aprovacao()
    with conectar_banco() as conn:
        registros_erros = conn.execute(
            '''SELECT prefixo, pesagem, revisao, campo, descricao
               FROM campos_com_erro
               ORDER BY prefixo, pesagem, revisao, campo'''
        ).fetchall()

    erros_por_ficha = {}
    for prefixo, pesagem, revisao, campo, descricao in registros_erros:
        chave = (
            safe_str(prefixo),
            normalizar_chave_ficha(pesagem),
            normalizar_chave_ficha(revisao),
        )
        erros_por_ficha.setdefault(chave, []).append(
            {"campo": campo, "descricao": descricao}
        )

    fichas = []
    ultima_revisao = {}
    for linha in carregar_chaves_fichas():
        prefixo = safe_str(linha.get("prefixo", ""))
        pesagem = normalizar_chave_ficha(linha.get("pesagem", ""))
        revisao = normalizar_chave_ficha(linha.get("revisao", ""))
        chave = (prefixo, pesagem, revisao)
        pesagem_num = safe_float(pesagem)
        revisao_num = safe_float(revisao)
        chave_pesagem = (prefixo, pesagem)
        ultima_revisao[chave_pesagem] = max(
            revisao_num,
            ultima_revisao.get(chave_pesagem, float("-inf")),
        )
        fichas.append(
            {
                "chave": chave,
                "chave_pesagem": chave_pesagem,
                "pesagem_num": pesagem_num,
                "revisao_num": revisao_num,
                "erros": erros_por_ficha.get(chave, []),
                "fluxo": fluxos.get(
                    chave,
                    {
                        "gerador_usuario": None,
                        "gerador_nome": "Não registrado",
                        "aprovador_usuario": None,
                        "aprovador_nome": "Pendente",
                    },
                ),
            }
        )

    pendentes_aprovacao = []
    pendentes_correcao = []
    for ficha in fichas:
        if ficha["erros"]:
            if not ficha["fluxo"]["aprovador_usuario"]:
                pendentes_correcao.append(ficha)
        elif (
            not ficha["fluxo"]["aprovador_usuario"]
            and ficha["revisao_num"]
            == ultima_revisao[ficha["chave_pesagem"]]
            and ficha["chave"] not in fora_aprovacao
        ):
            pendentes_aprovacao.append(ficha)

    ordenar = lambda ficha: (
        ficha["chave"][0],
        ficha["pesagem_num"],
        ficha["revisao_num"],
    )
    return sorted(pendentes_aprovacao, key=ordenar), sorted(
        pendentes_correcao, key=ordenar
    )


def carregar_fichas_fora_aprovacao():
    with conectar_banco() as conn:
        registros = conn.execute(
            'SELECT prefixo, pesagem, revisao FROM fichas_fora_aprovacao'
        ).fetchall()
    return {
        (
            safe_str(registro['prefixo']),
            normalizar_chave_ficha(registro['pesagem']),
            normalizar_chave_ficha(registro['revisao']),
        )
        for registro in registros
    }


AJUSTE_PENDENTES_MIGRADAS = "pendentes_migradas_fora_aprovacao"


@st.cache_resource
def separar_pendentes_migradas():
    """Roda uma vez só: as fichas que aguardavam aprovação em 21/10/2026 (as
    migradas sem quem gerou/aprovou) saem da Aprovação e ficam só em
    "Gerou e aprovou". Fichas geradas depois seguem o fluxo normal."""
    with conectar_banco() as conn:
        if conn.execute(
            'SELECT 1 FROM ajustes_aplicados WHERE nome = ?',
            (AJUSTE_PENDENTES_MIGRADAS,),
        ).fetchone():
            return
    listar_fichas_pendentes.clear()
    pendentes_aprovacao, _ = listar_fichas_pendentes()
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with conectar_banco() as conn:
        if pendentes_aprovacao:
            conn.executemany(
                '''INSERT INTO fichas_fora_aprovacao
                   (prefixo, pesagem, revisao, separado_em)
                   VALUES (?, ?, ?, ?)
                   ON CONFLICT DO NOTHING''',
                [(*ficha["chave"], agora) for ficha in pendentes_aprovacao],
            )
        conn.execute(
            '''INSERT INTO ajustes_aplicados (nome, aplicado_em) VALUES (?, ?)
               ON CONFLICT DO NOTHING''',
            (AJUSTE_PENDENTES_MIGRADAS, agora),
        )
    listar_fichas_pendentes.clear()


def salvar_campos_com_erro(prefixo, pesagem, revisao, campos, usuario):
    garantir_pode_emitir()
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    chave = (
        safe_str(prefixo),
        normalizar_chave_ficha(pesagem),
        normalizar_chave_ficha(revisao),
    )
    with conectar_banco() as conn:
        conn.execute(
            '''DELETE FROM campos_com_erro
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
            chave,
        )
        conn.executemany(
            '''INSERT INTO campos_com_erro
               (prefixo, pesagem, revisao, campo, descricao, marcado_por, marcado_em)
               VALUES (?, ?, ?, ?, ?, ?, ?)''',
            [
                (*chave, campo, descricao, usuario, agora)
                for campo, descricao in campos.items()
            ],
        )
        inserir_historico(
            conn, *chave,
            "Devolvida para correção" if campos else "Sinalizações removidas",
            usuario, agora,
        )
    invalidar_cache_fichas()


def limpar_campos_com_erro(prefixo, pesagem, revisao):
    garantir_pode_emitir()
    with conectar_banco() as conn:
        conn.execute(
            '''DELETE FROM campos_com_erro
               WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
            ),
        )
    invalidar_cache_fichas()

def aprovar_ficha(prefixo, pesagem, revisao, usuario):
    garantir_pode_emitir()
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with conectar_banco() as conn:
        inserir_historico(conn, prefixo, pesagem, revisao, "Aprovada", usuario, agora)
        conn.execute(
            '''INSERT INTO fluxo_fichas
               (prefixo, pesagem, revisao, aprovado_por, criado_em, aprovado_em)
               VALUES (?, ?, ?, ?, ?, ?)
               ON CONFLICT(prefixo, pesagem, revisao) DO UPDATE SET
                   aprovado_por = excluded.aprovado_por,
                   aprovado_em = excluded.aprovado_em''',
            (
                safe_str(prefixo),
                normalizar_chave_ficha(pesagem),
                normalizar_chave_ficha(revisao),
                usuario,
                agora,
                agora,
            )
        )
    invalidar_cache_fichas()


# Só a conta da Engenharia preenche, à mão, quem gerou e quem aprovou as fichas
# que vieram do banco antigo sem esse registro.
USUARIO_GEROU_APROVOU = "engenharia"


def pode_preencher_gerou_aprovou():
    return (
        safe_str(st.session_state.get('usuario_id')).casefold()
        == USUARIO_GEROU_APROVOU
    )


@st.cache_data
def listar_fichas_gerou_aprovou():
    fluxos = carregar_fluxos_fichas()
    with conectar_banco() as conn:
        registros = conn.execute(
            """SELECT dados->>'Prefixo' AS prefixo,
                      dados->>'Pesagem' AS pesagem,
                      COALESCE(dados->>'Revisao', dados->>'revisao') AS revisao,
                      dados->>'Data_da_Pesagem' AS data
               FROM pesagens
               WHERE NULLIF(btrim(dados->>'Prefixo'), '') IS NOT NULL"""
        ).fetchall()
    linhas = []
    for registro in registros:
        chave = (
            safe_str(registro["prefixo"]),
            normalizar_chave_ficha(registro["pesagem"]),
            normalizar_chave_ficha(registro["revisao"]),
        )
        fluxo = fluxos.get(chave, {})
        linhas.append({
            "chave": chave,
            "Aeronave": chave[0],
            "Pesagem": chave[1],
            "Revisão": chave[2],
            "Data da pesagem": data_pesagem(registro["data"]),
            "gerado_por": fluxo.get("gerador_usuario") or None,
            "aprovado_por": fluxo.get("aprovador_usuario") or None,
            "pesagem_num": safe_float(chave[1]),
            "revisao_num": safe_float(chave[2]),
        })
    colunas = [
        "chave", "Aeronave", "Pesagem", "Revisão", "Data da pesagem",
        "gerado_por", "aprovado_por", "pesagem_num", "revisao_num",
    ]
    df = pd.DataFrame(linhas, columns=colunas)
    return (
        df.sort_values(["Aeronave", "pesagem_num", "revisao_num"])
        .drop_duplicates("chave")
        .reset_index(drop=True)
    )


def definir_gerou_aprovou(mudancas, usuario):
    """Grava quem gerou/aprovou sem a regra da edição (não troca o emissor
    nem desfaz a aprovação de quem não foi alterado)."""
    if not pode_preencher_gerou_aprovou():
        raise PermissionError("Só o usuário Engenharia pode informar quem gerou e aprovou.")
    agora = datetime.datetime.now(datetime.timezone.utc).isoformat()
    with conectar_banco() as conn:
        for chave, mudanca in mudancas.items():
            conn.execute(
                '''INSERT INTO fluxo_fichas
                   (prefixo, pesagem, revisao, gerado_por, aprovado_por,
                    criado_em, aprovado_em)
                   VALUES (?, ?, ?, ?, ?, ?, ?)
                   ON CONFLICT(prefixo, pesagem, revisao) DO UPDATE SET
                       gerado_por = excluded.gerado_por,
                       aprovado_por = excluded.aprovado_por,
                       aprovado_em = CASE
                           WHEN excluded.aprovado_por IS NULL THEN NULL
                           ELSE COALESCE(fluxo_fichas.aprovado_em, excluded.aprovado_em)
                       END''',
                (
                    *chave,
                    mudanca["gerado_por"],
                    mudanca["aprovado_por"],
                    agora,
                    agora if mudanca["aprovado_por"] else None,
                ),
            )
            inserir_historico(
                conn, *chave, "Emissor/aprovador informado", usuario, agora,
                mudanca["alteracoes"],
            )
    invalidar_cache_fichas()

@st.cache_resource
def normalizar_chaves_fluxo():
    """Converte chaves antigas como '15.0' para '15' no fluxo de aprovação.

    O app procura as fichas por '15'; registros migrados com '15.0' ficavam
    invisíveis (gerador "Não registrado", devoluções que não apareciam).
    Linhas que colidiriam com uma já normalizada são mantidas como estão.
    """
    normalizar = r"regexp_replace({0}, '^(-?[0-9]+)[.]0+$', '\1')"
    for tabela, chave_extra in (("fluxo_fichas", ""), ("campos_com_erro", " AND g.campo = f.campo")):
        try:
            with conectar_banco() as conn:
                conn.execute(
                    f"""UPDATE {tabela} AS f
                        SET pesagem = {normalizar.format('f.pesagem')},
                            revisao = {normalizar.format('f.revisao')}
                        WHERE (f.pesagem ~ '[.]0+$' OR f.revisao ~ '[.]0+$')
                          AND NOT EXISTS (
                              SELECT 1 FROM {tabela} AS g
                              WHERE g.prefixo = f.prefixo
                                AND g.pesagem = {normalizar.format('f.pesagem')}
                                AND g.revisao = {normalizar.format('f.revisao')}
                                {chave_extra}
                          )"""
                )
        except Exception:
            pass


def avisar_banco_indisponivel(erro):
    st.error(
        "Não foi possível conectar ao banco de dados. Nada do que já estava "
        "salvo foi perdido."
    )
    st.markdown(
        "- **Supabase pausado:** projetos gratuitos são pausados depois de "
        "7 dias sem uso. Entre em supabase.com/dashboard, abra o projeto e "
        "clique em **Restore project**; depois de alguns minutos, tente de novo.\n"
        "- **URL do banco:** confira em Settings → Secrets se "
        "[connections.postgresql].url ainda é a do pooler do Supabase "
        "(Project Settings → Database → Connection string)."
    )
    if st.button("Tentar novamente", type="primary"):
        obter_conexao_postgresql.clear()
        st.rerun()
    with st.expander("Detalhe do erro"):
        st.code(f"{type(getattr(erro, 'orig', erro)).__name__}: "
                f"{getattr(erro, 'orig', erro)}")
    st.stop()


try:
    inicializar_controle_acesso()
except DBAPIError as erro:
    avisar_banco_indisponivel(erro)
normalizar_chaves_fluxo()
separar_pendentes_migradas()
dict_tipos_aeronave.update(carregar_aeronaves_cadastradas())
df_historico = carregar_dados_banco()

# 3. FUNÇÃO DE EXPORTAÇÃO PARA EXCEL
def gerar_excel_por_template(dados, caminho_template="exemplo_ficha.xlsx"):
    try:
        wb = openpyxl.load_workbook(caminho_template)
    except FileNotFoundError as erro:
        raise FileNotFoundError(
            f"O arquivo de template '{caminho_template}' não foi encontrado na pasta."
        ) from erro

    ws = wb['Ficha de pesagem - Template']
    
    # CABEÇALHO
    ws['G8'] = dados.get('prefixo', '')
    ws['V8'] = dados.get('pesado_por', '')
    ws['G9'] = dados.get('modelo', '')
    ws['V9'] = dados.get('local', '')
    ws['V10'] = formatar_data_ficha(dados.get('data', ''))
    ws['V12'] = dados.get('razao', '')
    ws['G13'] = dados.get('config_lopa', '')
    ws['G14'] = dados.get('lopa', '')
    ws['S22'] = dados.get('arm_a', '')
    ws['S24'] = dados.get('arm_b_lh', '')
    ws['S25'] = dados.get('arm_b_rh', '')
    
    ws['P43'] = formatar_data_ficha(dados.get('issue_date', ''))
    ws['O65'] = dados.get('ultima_pesagem', '') 
    ws['W65'] = dados.get('ultima_revisao', '') 
    
    ws['G10'] = dados.get('vrbl_number', '') 
    ws['G11'] = dados.get('serial_number', '') 
    ws['G12'] = dados.get('line_number', '') 
    
    ws['X43'] = dados.get('numero_pesagem', '') 
    ws['AC43'] = dados.get('revisao', '') 
    
    # REACTIONS & BASIC WEIGHT
    reacoes = ['LH', 'RH', 'NOSE', 'TAIL', 'TOTAL REGISTERED', 'DEDUCTIONS', 'ADDITIONS', 'AIRCRAFT BASIC WEIGHT']
    linha_reacao = 31
    for r in reacoes:
        rd = dados['reactions'].get(r, {})
        ws[f'N{linha_reacao}'] = rd.get('weight', 0.0)
        ws[f'S{linha_reacao}'] = rd.get('arm', 0.0)
        ws[f'X{linha_reacao}'] = rd.get('moment', 0.0)
        linha_reacao += 1

    ws['J44'] = dados.get('cg_mac', 0.0)

    def set_cell_value(sheet, row, col_letter, value):
        cell = sheet[f"{col_letter}{row}"]
        if isinstance(cell, openpyxl.cell.cell.MergedCell):
            for range_merged in sheet.merged_cells.ranges:
                if f"{col_letter}{row}" in range_merged:
                    min_col, min_row, _, _ = range_merged.bounds
                    cell = sheet.cell(row=min_row, column=min_col)
                    break
        cell.value = value

    deducoes = dados.get('deductions', [])
    if len(deducoes) > 22:
        raise ValueError("O template suporta no máximo 22 itens de deduções.")
    for idx in range(70, 92):
        if (idx - 70) < len(deducoes):
            d = deducoes[idx - 70]
            desc = d.get('desc', '')
            w_val = d.get('w', 0.0)
            a_val = d.get('a', 0.0)
            m_val = d.get('m', 0.0)
        else:
            desc, w_val, a_val, m_val = "", 0.0, 0.0, 0.0
            
        set_cell_value(ws, idx, 'B', desc)
        set_cell_value(ws, idx, 'R', w_val)
        set_cell_value(ws, idx, 'V', a_val)
        set_cell_value(ws, idx, 'Z', m_val)

    adicoes = dados.get('additions', [])
    if len(adicoes) > 32:
        raise ValueError("O template suporta no máximo 32 itens de adições.")
    for idx in range(95, 127):
        if (idx - 95) < len(adicoes):
            a = adicoes[idx - 95]
            desc = a.get('desc', '')
            w_val = a.get('w', 0.0)
            a_val = a.get('a', 0.0)
            m_val = a.get('m', 0.0)
        else:
            desc, w_val, a_val, m_val = "", 0.0, 0.0, 0.0
            
        set_cell_value(ws, idx, 'B', desc)
        set_cell_value(ws, idx, 'R', w_val)
        set_cell_value(ws, idx, 'V', a_val)
        set_cell_value(ws, idx, 'Z', m_val)

    for celula, chave in (
        ("I50", "assinatura_emissor"),
        ("U50", "assinatura_aprovador"),
    ):
        assinatura = dados.get(chave)
        if assinatura:
            imagem = ExcelImage(io.BytesIO(assinatura))
            escala = min(180 / imagem.width, 54 / imagem.height, 1.5)
            imagem.width = int(imagem.width * escala * 1.5)
            imagem.height = int(imagem.height * escala)
            imagem.anchor = celula
            ws.add_image(imagem)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def caixa_total(container, valor, espaco_px):
    """Caixa de total somente leitura, alinhada ao lado dos campos somados."""
    texto = f"{safe_float(valor):,.0f}".replace(",", ".")
    container.markdown(
        f"<div style='height:{espaco_px}px'></div>"
        "<div style='font-size:0.875rem;margin-bottom:6px'>Total</div>"
        "<div style='border:1px solid #C9CED6;border-radius:8px;padding:8px 12px;"
        f"background:#FFF4EC;font-weight:700;min-height:40px'>{texto}</div>",
        unsafe_allow_html=True,
    )


def grade_celulas(render):
    """Monta as células como na ficha manual: Nariz (LH/RH) com o total NLG ao lado,
    RH MLG 1/2 e LH MLG 1/2 empilhados, cada lado com o seu total.
    render(container, indice) desenha o campo e devolve o valor; índices:
    0 Nariz LH, 1 Nariz RH, 2 RH MLG 1, 3 RH MLG 2, 4 LH MLG 1, 5 LH MLG 2."""
    c1, c2, c3, c4 = st.columns(4)
    c1.markdown("<div style='height:75px'></div>", unsafe_allow_html=True)
    v = [None] * 6
    v[0] = render(c1, 0)
    v[1] = render(c1, 1)
    for indice in (2, 3, 4, 5):
        v[indice] = render(c3, indice)
    caixa_total(c2, safe_float(v[0]) + safe_float(v[1]), 112)
    caixa_total(c4, safe_float(v[2]) + safe_float(v[3]), 40)
    caixa_total(c4, safe_float(v[4]) + safe_float(v[5]), 80)
    return v


def mostrar_relatorio_pesagem(tabela, peso, braco, cg_mac):
    r1, r2, r3 = st.columns(3)
    r1.metric("Aircraft Basic Weight", f"{peso:,.2f} kgf")
    r2.metric("Aircraft Basic Arm", f"{braco:,.4f} in")
    r3.metric("C.G. (% MAC)", f"{cg_mac:.2f} %")
    st.dataframe(
        pd.DataFrame(tabela),
        use_container_width=True,
        hide_index=True,
        column_config={
            "Weight (kgf)": st.column_config.NumberColumn(format="%.2f"),
            "Arm (inch)": st.column_config.NumberColumn(format="%.4f"),
            "Moment (kgf x inch)": st.column_config.NumberColumn(format="%,.2f"),
        },
    )


# 4. TELAS E FUNCIONALIDADES

@protegido
def tela_consulta():
    cabecalho("Consultar fichas", "Visão geral e consulta do histórico de pesagens.")
    if st.button("Atualizar", key="atualizar_fichas"):
        invalidar_cache_fichas()
        st.rerun()

    if df_historico.empty:
        st.warning(
            "Nenhuma ficha foi carregada do banco conectado ao aplicativo. "
            "No Streamlit Community Cloud, confira em Settings → Secrets se "
            "[connections.postgresql].url aponta para o mesmo banco Supabase "
            "usado na importação e reinicie o app."
        )
        return

    somente_ultimas = nivel_atual() == NIVEL_ENG_OPERACOES
    df_fichas = ultimas_fichas(df_historico) if somente_ultimas else df_historico
    prefixos_unicos = sorted(df_fichas['Prefixo'].dropna().unique().tolist()) if 'Prefixo' in df_fichas.columns else []
    if somente_ultimas:
        st.caption("Seu perfil mostra apenas a ficha mais recente de cada aeronave.")
        st.metric("Aeronaves com ficha", len(prefixos_unicos))
    else:
        pendentes_aprovacao, pendentes_correcao = listar_fichas_pendentes()
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Aeronaves com ficha", len(prefixos_unicos))
        k2.metric("Fichas emitidas", len(df_fichas))
        k3.metric("Aguardando aprovação", len(pendentes_aprovacao))
        k4.metric("Devolvidas para correção", len(pendentes_correcao))

    ordenar_numero = lambda valores: sorted(valores, key=lambda v: (safe_float(v), v))
    with st.container(border=True):
        st.markdown("**Consultar ficha**")
        c1, c2, c3 = st.columns([2, 1, 1])
        prefixo = c1.selectbox(
            "Aeronave", prefixos_unicos, index=None, placeholder="Digite ou escolha o prefixo"
        )
        df_filtrado = (
            df_fichas[df_fichas['Prefixo'] == prefixo] if prefixo else df_fichas.iloc[0:0]
        )
        pesagens = ordenar_numero(df_filtrado['Pesagem'].dropna().unique().tolist())
        pesagem = c2.selectbox(
            "Pesagem", pesagens, index=len(pesagens) - 1 if pesagens else None,
            disabled=not prefixo, placeholder="—",
        )
        revisoes = ordenar_numero(
            df_filtrado[df_filtrado['Pesagem'] == pesagem]['Revisao'].dropna().unique().tolist()
        ) if pesagem else []
        revisao = c3.selectbox(
            "Revisão", revisoes, index=len(revisoes) - 1 if revisoes else None,
            disabled=not pesagem, placeholder="—",
        )

    if prefixo and pesagem:
        if revisao:
                st.divider()
                linha = df_filtrado[(df_filtrado['Pesagem'] == pesagem) & (df_filtrado['Revisao'] == revisao)].iloc[0]
                renderizar_ficha_visualizacao(prefixo, pesagem, revisao, linha)

def ultimas_fichas(df):
    """Mantém só a maior pesagem e, dentro dela, a maior revisão de cada aeronave."""
    if df.empty or not {'Prefixo', 'Pesagem', 'Revisao'}.issubset(df.columns):
        return df.iloc[0:0]
    ordem = df.assign(
        _pesagem=df['Pesagem'].map(safe_float),
        _revisao=df['Revisao'].map(safe_float),
    ).sort_values(['Prefixo', '_pesagem', '_revisao'])
    return (
        ordem.groupby('Prefixo', sort=False).tail(1)
        .drop(columns=['_pesagem', '_revisao'])
        .reset_index(drop=True)
    )


CORES_MODELOS = {
    "737-700": "#3B6BC7",
    "737-800": "#E8590C",
    "737 MAX 8": "#2A9D8F",
    "737-800 BCF": "#EDA100",
}


@st.cache_data
def carregar_variantes_aeronave():
    """Segunda coluna MODEL da planilha da frota (ex.: 8EH, MAX, BCF), por prefixo."""
    try:
        df_tipos = pd.read_excel('Cópia de Ficha_Pesagem_v2.xlsm', sheet_name='Sheet1')
    except Exception:
        return {}
    if 'VRG REG.' not in df_tipos.columns or 'MODEL.1' not in df_tipos.columns:
        return {}
    return {
        safe_str(prefixo): safe_str(variante)
        for prefixo, variante in zip(df_tipos['VRG REG.'], df_tipos['MODEL.1'])
        if safe_str(prefixo)
    }


def nome_modelo_grafico(modelo, prefixo=""):
    texto = safe_str(modelo).upper()
    # Cargueiros convertidos (BCF) têm peso e C.G. bem diferentes dos de passageiros.
    # Todos os PS-GF* da frota são 737-800 BCF.
    if (normalizar_prefixo(prefixo).startswith("PS-GF")
            or "BCF" in carregar_variantes_aeronave().get(safe_str(prefixo), "").upper()):
        return "737-800 BCF"
    if "MAX" in texto:
        return "737 MAX 8"
    if "800" in texto:
        return "737-800"
    if "700" in texto:
        return "737-700"
    return safe_str(modelo) or "Não informado"


def calcular_resumo_pesagem(prefixo, row):
    """Peso básico e C.G. da ficha, com as mesmas contas do Weighing Report."""
    info_aero = dict_tipos_aeronave.get(prefixo, {})
    lh = (
        safe_float(row.get('Peso MLG LH 1', 0)) + safe_float(row.get('Peso MLG LH2', 0))
        + safe_float(row.get('Peso MLG LH 1 pesagem 2', 0))
        + safe_float(row.get('Peso MLG LH2 pesagem 2', 0))
    ) / 2
    rh = (
        safe_float(row.get('Peso MLG RH 1 ', 0)) + safe_float(row.get('Peso MLG RH 2', 0))
        + safe_float(row.get('Peso MLG RH 1  pesagem 2', row.get('Peso MLG RH 1 pesagem 2', 0)))
        + safe_float(row.get('Peso MLG RH 2 pesagem 2', 0))
    ) / 2
    nose = (
        (safe_float(row.get('Peso nariz LH', 0)) + safe_float(row.get('Peso nariz LH pesagem 2', 0))) / 2
        + (safe_float(row.get('Peso Nariz RH', 0)) + safe_float(row.get('Peso Nariz RH pesagem 2', 0))) / 2
    )
    tail = safe_float(row.get('Tail', 0))
    arm_b_lh = (safe_float(row.get('Canela MLG LH', 0)) * 0.1304) + 706.17
    arm_b_rh = (safe_float(row.get('Canela MLG RH', 0)) * 0.1304) + 706.17
    arm_a = info_aero.get('armnose', 93.0000)
    arm_c = 1298.9000

    tot_reg_weight = lh + rh + nose + tail
    tot_reg_moment = lh * arm_b_lh + rh * arm_b_rh + nose * arm_a + tail * arm_c

    _, level_correction_moment, level_correction_side = calcular_level_correction(
        normalizar_angulo_level_correction(row.get('Graus correção do cg', '')),
        tot_reg_weight,
    )

    def itens(campo_desc, campo_peso, campo_braco, limite):
        lista = []
        for i in range(1, limite + 1):
            desc = row.get(campo_desc.format(i), None)
            if pd.notna(desc) and str(desc).strip() != '':
                lista.append((
                    str(desc),
                    safe_float(row.get(campo_peso.format(i), None)),
                    safe_float(row.get(campo_braco.format(i), None)),
                ))
        return lista

    deducoes = itens('Deductions description {}', 'Deductions Weigth {}', 'Deductions arm {}', MAX_DEDUCTIONS)
    adicoes = itens('Additions Description {}', 'Additions weigth {}', 'Additions arm {}', MAX_ADDITIONS)

    ded_weight = sum(peso for _, peso, _ in deducoes)
    ded_moment = sum(peso * braco for _, peso, braco in deducoes)
    if level_correction_side == "deductions":
        ded_moment += level_correction_moment

    momento_extra_flaps = 0.0
    if any("flap" in desc.casefold() for desc, _, _ in adicoes):
        momento_extra_flaps = (
            momento_flaps_do_registro(row)
            or buscar_momento_flaps_banco(prefixo, info_aero.get('modelo', ""))
        )
    add_weight = sum(peso for _, peso, _ in adicoes)
    add_moment = sum(peso * braco for _, peso, braco in adicoes) + momento_extra_flaps
    if level_correction_side == "additions":
        add_moment += level_correction_moment

    basic_weight = tot_reg_weight + add_weight - ded_weight
    basic_moment = tot_reg_moment + add_moment - ded_moment
    basic_arm = basic_moment / basic_weight if basic_weight > 0 else 0.0
    cg_mac = ((basic_arm - 627.1) / 1.558) if basic_arm > 0 else 0.0
    return {"peso": basic_weight, "braco": basic_arm, "cg_mac": cg_mac}


def data_pesagem(valor):
    texto = safe_str(valor)
    if not texto:
        return pd.NaT
    iso = len(texto) >= 10 and texto[4] == "-"
    data = pd.to_datetime(texto, dayfirst=not iso, errors="coerce")
    # Datas no futuro são erro de digitação na ficha; ficam sem data.
    if pd.notna(data) and data > pd.Timestamp.now():
        return pd.NaT
    return data


@st.cache_data
def montar_resumo_frota():
    linhas = []
    for _, row in ultimas_fichas(carregar_dados_banco()).iterrows():
        prefixo = safe_str(row.get('Prefixo'))
        info_aero = dict_tipos_aeronave.get(prefixo, {})
        resumo = calcular_resumo_pesagem(prefixo, row)
        linhas.append({
            "Aeronave": prefixo,
            "Modelo": nome_modelo_grafico(info_aero.get('modelo') or row.get('Tipo_aeronave'), prefixo),
            "Pesagem": safe_str(row.get('Pesagem')),
            "Revisão": safe_str(row.get('Revisao')),
            "Data da pesagem": data_pesagem(row.get('Data_da_Pesagem')),
            "Peso básico (kgf)": round(resumo["peso"], 1),
            "C.G. (% MAC)": round(resumo["cg_mac"], 2),
        })
    return pd.DataFrame(
        linhas,
        columns=["Aeronave", "Modelo", "Pesagem", "Revisão", "Data da pesagem",
                 "Peso básico (kgf)", "C.G. (% MAC)"],
    )


def formatar_kg(valor):
    return f"{valor:,.0f} kgf".replace(",", ".")


def formatar_pct(valor):
    return f"{valor:.2f} %".replace(".", ",")


LOCALE_GRAFICOS = {
    "number": {"decimal": ",", "thousands": ".", "grouping": [3], "currency": ["R$ ", ""]},
}


def mostrar_grafico(grafico, **kwargs):
    st.altair_chart(grafico.configure(locale=LOCALE_GRAFICOS), **kwargs)


def motivos_fora_do_padrao(df):
    """Aponta fichas cujo peso ou C.G. foge muito do que é normal para o modelo."""
    mediana_peso = df.groupby("Modelo")["Peso básico (kgf)"].transform("median")
    mediana_cg = df.groupby("Modelo")["C.G. (% MAC)"].transform("median")
    peso_fora = (df["Peso básico (kgf)"] - mediana_peso).abs() > 0.04 * mediana_peso
    cg_fora = (df["C.G. (% MAC)"] - mediana_cg).abs() > 4
    motivos = pd.Series("", index=df.index)
    motivos[peso_fora] = "Peso fora do padrão"
    motivos[cg_fora] = "C.G. fora do padrão"
    motivos[peso_fora & cg_fora] = "Peso e C.G. fora do padrão"
    return motivos


# Peso máximo de pouso de cada modelo (kgf), base dos limites da FAA AC 120-27F.
# Conferir no AFM da frota.
PESO_MAXIMO_POUSO = {
    "737-700": 58060,
    "737-800": 65320,
    "737 MAX 8": 69310,
    "737-800 BCF": 65320,
}
# AC 120-27F §2.1.3.1: restabelecer o peso básico quando a soma das mudanças
# passar de 0,5 % do peso máximo de pouso ou de 0,5 % MAC.
LIMITE_VARIACAO_PESO = 0.005
LIMITE_VARIACAO_CG = 0.5


def tolerancia_peso_modelo(modelo):
    mlw = PESO_MAXIMO_POUSO.get(modelo)
    return mlw * LIMITE_VARIACAO_PESO if mlw else None


@st.cache_data
def montar_historico_frota():
    """Peso e C.G. de todas as pesagens e revisões de cada aeronave."""
    df = carregar_dados_banco()
    if df.empty or not {'Prefixo', 'Pesagem', 'Revisao'}.issubset(df.columns):
        return pd.DataFrame()
    # Fichas repetidas no banco: vale a última gravada.
    df = df.drop_duplicates(subset=['Prefixo', 'Pesagem', 'Revisao'], keep='last')
    linhas = []
    for _, row in df.iterrows():
        prefixo = safe_str(row.get('Prefixo'))
        info_aero = dict_tipos_aeronave.get(prefixo, {})
        resumo = calcular_resumo_pesagem(prefixo, row)
        if resumo["peso"] <= 0 or resumo["cg_mac"] <= 0:
            continue
        pesagem = safe_float(row.get('Pesagem'))
        revisao = safe_float(row.get('Revisao'))
        linhas.append({
            "aeronave": prefixo,
            "modelo": nome_modelo_grafico(info_aero.get('modelo') or row.get('Tipo_aeronave'), prefixo),
            "pesagem_num": pesagem,
            "revisao_num": revisao,
            "ficha": f"P{pesagem:g} R{revisao:g}",
            "data": data_pesagem(row.get('Data_da_Pesagem')),
            "peso": round(resumo["peso"], 1),
            "cg": round(resumo["cg_mac"], 2),
        })
    historico = pd.DataFrame(linhas)
    if historico.empty:
        return historico
    return historico.sort_values(["aeronave", "pesagem_num", "revisao_num"]).reset_index(drop=True)


def variacao_desde_pesagem_anterior(historico):
    """Última ficha de cada aeronave contra a última revisão da pesagem anterior."""
    linhas = []
    for aeronave, fichas in historico.groupby("aeronave", sort=True):
        atual = fichas.iloc[-1]
        anteriores = fichas[fichas["pesagem_num"] < atual["pesagem_num"]]
        if anteriores.empty:
            continue
        anterior = anteriores.iloc[-1]
        tol_peso = tolerancia_peso_modelo(atual["modelo"])
        delta_peso = atual["peso"] - anterior["peso"]
        delta_cg = atual["cg"] - anterior["cg"]
        fora = []
        if tol_peso and abs(delta_peso) > tol_peso:
            fora.append("peso")
        if abs(delta_cg) > LIMITE_VARIACAO_CG:
            fora.append("C.G.")
        linhas.append({
            "aeronave": aeronave,
            "modelo": atual["modelo"],
            "de": anterior["ficha"],
            "para": atual["ficha"],
            "delta_peso": round(delta_peso, 1),
            "delta_cg": round(delta_cg, 2),
            "tol_peso": tol_peso,
            "situacao": ("Passou do limite de " + " e ".join(fora)) if fora else "Dentro do limite",
        })
    return pd.DataFrame(linhas)


def aba_variacao(alt, historico, modelos, escala_cor, cor):
    st.caption(
        "Última ficha de cada aeronave comparada com a última revisão da pesagem "
        "anterior. A caixa tracejada é o limite da FAA AC 120-27F (§2.1.3.1): "
        "0,5 % do peso máximo de pouso e 0,5 % MAC. Fora dela, o peso básico "
        "deve ser restabelecido."
    )
    variacao = variacao_desde_pesagem_anterior(historico)
    if not variacao.empty:
        variacao = variacao[variacao["modelo"].isin(modelos)]
    if variacao.empty:
        st.info("Nenhuma aeronave tem mais de uma pesagem para comparar.")
        return
    fora = variacao[variacao["situacao"] != "Dentro do limite"]
    k1, k2, k3 = st.columns(3)
    k1.metric("Aeronaves comparadas", len(variacao))
    k2.metric("Passaram do limite", len(fora))
    k3.metric(
        "Variação média de peso",
        f"{variacao['delta_peso'].mean():+,.0f} kgf".replace(",", "."),
    )

    tol_maior = max(
        (t for t in variacao["tol_peso"].dropna()), default=0
    )
    caixas = pd.DataFrame([
        {"modelo": m, "x": -LIMITE_VARIACAO_CG, "x2": LIMITE_VARIACAO_CG,
         "y": -tolerancia_peso_modelo(m), "y2": tolerancia_peso_modelo(m)}
        for m in modelos if tolerancia_peso_modelo(m)
    ])
    # Corta o gráfico em volta da caixa para os valores absurdos não esmagarem o resto.
    lim_x = max(LIMITE_VARIACAO_CG * 4, float(variacao["delta_cg"].abs().quantile(0.9)) * 1.2)
    lim_y = max(tol_maior * 4, float(variacao["delta_peso"].abs().quantile(0.9)) * 1.2)
    dentro_do_grafico = variacao[
        (variacao["delta_cg"].abs() <= lim_x) & (variacao["delta_peso"].abs() <= lim_y)
    ]
    dica = [
        alt.Tooltip("aeronave:N", title="Aeronave"),
        alt.Tooltip("modelo:N", title="Modelo"),
        alt.Tooltip("de:N", title="De"),
        alt.Tooltip("para:N", title="Para"),
        alt.Tooltip("delta_peso:Q", title="Δ peso (kgf)", format="+,.1f"),
        alt.Tooltip("delta_cg:Q", title="Δ C.G. (% MAC)", format="+.2f"),
        alt.Tooltip("situacao:N", title="Situação"),
    ]
    with st.container(border=True):
        st.markdown("**Quanto cada aeronave mudou desde a pesagem anterior**")
        escala_x = alt.Scale(domain=[-lim_x, lim_x])
        escala_y = alt.Scale(domain=[-lim_y, lim_y])
        camadas = []
        if not caixas.empty:
            camadas.append(
                alt.Chart(caixas).mark_rect(fill=None, strokeDash=[5, 4], strokeWidth=1.5).encode(
                    x=alt.X("x:Q", scale=escala_x), x2="x2:Q",
                    y=alt.Y("y:Q", scale=escala_y), y2="y2:Q",
                    stroke=alt.Color("modelo:N", scale=escala_cor, legend=None),
                )
            )
        camadas.append(
            alt.Chart(dentro_do_grafico).mark_circle(
                size=90, opacity=0.85, stroke="white", strokeWidth=1.5
            ).encode(
                x=alt.X("delta_cg:Q", title="Δ C.G. (% MAC)", scale=escala_x),
                y=alt.Y("delta_peso:Q", title="Δ peso básico (kgf)", scale=escala_y),
                color=cor,
                tooltip=dica,
            )
        )
        mostrar_grafico(alt.layer(*camadas).properties(height=420), use_container_width=True)
        escondidas = len(variacao) - len(dentro_do_grafico)
        if escondidas:
            st.caption(f"{escondidas} aeronave(s) com variação muito grande ficaram fora da área do gráfico e aparecem na tabela.")

    if not fora.empty:
        with st.container(border=True):
            st.markdown(f"**Aeronaves que passaram do limite ({len(fora)})**")
            st.dataframe(
                fora.assign(abs_peso=fora["delta_peso"].abs())
                .sort_values("abs_peso", ascending=False)
                .drop(columns=["abs_peso", "tol_peso"])
                .rename(columns={
                    "aeronave": "Aeronave", "modelo": "Modelo", "de": "De", "para": "Para",
                    "delta_peso": "Δ peso (kgf)", "delta_cg": "Δ C.G. (% MAC)",
                    "situacao": "Situação",
                }),
                hide_index=True, use_container_width=True,
                column_config={
                    "Δ peso (kgf)": st.column_config.NumberColumn(format="%+.1f"),
                    "Δ C.G. (% MAC)": st.column_config.NumberColumn(format="%+.2f"),
                },
            )


def aba_historico(alt, historico, modelos, escala_cor):
    aeronaves = sorted(historico.loc[historico["modelo"].isin(modelos), "aeronave"].unique())
    if not aeronaves:
        st.info("Nenhuma aeronave com fichas para mostrar.")
        return
    contagem = historico[historico["aeronave"].isin(aeronaves)]["aeronave"].value_counts()
    aeronave = st.selectbox(
        "Aeronave", aeronaves, index=aeronaves.index(contagem.idxmax()),
        key="graficos_historico_aeronave",
    )
    fichas = historico[historico["aeronave"] == aeronave].copy()
    modelo = fichas["modelo"].iloc[-1]
    tol_peso = tolerancia_peso_modelo(modelo)
    # Faixa de cada pesagem: o valor da primeira revisão ± o limite da AC.
    primeira = fichas.groupby("pesagem_num")[["peso", "cg"]].transform("first")
    fichas["peso_min"] = primeira["peso"] - (tol_peso or 0)
    fichas["peso_max"] = primeira["peso"] + (tol_peso or 0)
    fichas["cg_min"] = primeira["cg"] - LIMITE_VARIACAO_CG
    fichas["cg_max"] = primeira["cg"] + LIMITE_VARIACAO_CG
    fichas["pesagem"] = fichas["pesagem_num"].map(lambda v: f"Pesagem {v:g}")
    ordem = fichas["ficha"].tolist()
    cor_modelo = CORES_MODELOS.get(modelo, "#8A8F98")
    st.caption(
        f"{modelo} · {len(fichas)} ficha(s). A faixa clara em cada pesagem vai do valor "
        "da primeira revisão até o limite da FAA AC 120-27F (0,5 % do peso máximo de "
        "pouso e 0,5 % MAC)."
    )
    dica = [
        alt.Tooltip("ficha:N", title="Ficha"),
        alt.Tooltip("data:T", title="Data da pesagem", format="%d/%m/%Y"),
        alt.Tooltip("peso:Q", title="Peso básico (kgf)", format=",.1f"),
        alt.Tooltip("cg:Q", title="C.G. (% MAC)", format=".2f"),
    ]

    def evolucao(campo, minimo, maximo, titulo, formato, mostrar_faixa=True):
        base = alt.Chart(fichas).encode(
            x=alt.X("ficha:N", sort=ordem, title=None, axis=alt.Axis(labelAngle=0)),
        )
        camadas = []
        if mostrar_faixa:
            camadas.append(base.mark_area(opacity=0.15, color=cor_modelo).encode(
                y=alt.Y(f"{minimo}:Q", title=titulo, scale=alt.Scale(zero=False)),
                y2=f"{maximo}:Q",
                detail="pesagem:N",
            ))
        camadas.append(base.mark_line(color=cor_modelo, strokeWidth=2).encode(
            y=alt.Y(f"{campo}:Q", title=titulo, scale=alt.Scale(zero=False)),
        ))
        camadas.append(base.mark_circle(size=80, color=cor_modelo, opacity=1).encode(
            y=f"{campo}:Q", tooltip=dica,
        ))
        camadas.append(base.mark_text(dy=-12, color="#1F2430").encode(
            y=f"{campo}:Q", text=alt.Text(f"{campo}:Q", format=formato),
        ))
        return alt.layer(*camadas).properties(height=260)

    with st.container(border=True):
        st.markdown("**Peso básico em cada ficha**")
        mostrar_grafico(
            evolucao("peso", "peso_min", "peso_max", "Peso básico (kgf)", ",.0f", bool(tol_peso)),
            use_container_width=True,
        )
    with st.container(border=True):
        st.markdown("**C.G. em cada ficha**")
        mostrar_grafico(
            evolucao("cg", "cg_min", "cg_max", "C.G. (% MAC)", ".2f"),
            use_container_width=True,
        )


def aba_carta_controle(alt, selecionadas, modelos):
    modelo = st.radio(
        "Modelo", modelos, horizontal=True, key="graficos_carta_modelo",
    )
    dados = selecionadas[selecionadas["Modelo"] == modelo].rename(columns={
        "Aeronave": "aeronave", "Peso básico (kgf)": "peso", "C.G. (% MAC)": "cg",
    })
    if len(dados) < 3:
        st.info("Esse modelo tem poucas aeronaves para montar a carta de controle.")
        return
    st.caption(
        "Carta de controle de valores individuais (Shewhart), com a última ficha de "
        "cada aeronave. A linha do meio é a mediana do modelo; as tracejadas ficam a 3 "
        "desvios-padrão, estimados pela mediana da amplitude móvel. Ponto fora das linhas é "
        "aeronave que foge do padrão da frota."
    )
    cor_modelo = CORES_MODELOS.get(modelo, "#8A8F98")

    def carta(campo, titulo, formato):
        # Mediana no centro: uma ficha com valor absurdo não puxa a linha do meio.
        media = dados[campo].median()
        # Desvio estimado pela mediana das amplitudes móveis (MR̃/0,954), que não
        # se deixa inflar por uma ficha com valor absurdo como a média das amplitudes.
        ordem = dados.sort_values("aeronave")[campo]
        desvio = ordem.diff().abs().median() / 0.954
        lsc, lic = media + 3 * desvio, media - 3 * desvio
        tabela = dados.assign(
            fora=((dados[campo] > lsc) | (dados[campo] < lic)).map(
                {True: "Fora dos limites", False: "Dentro dos limites"}
            )
        ).sort_values("aeronave")
        fora = tabela[tabela["fora"] == "Fora dos limites"]
        linhas = pd.DataFrame({
            "valor": [lsc, media, lic],
            "nome": ["Limite superior", "Mediana", "Limite inferior"],
            "traco": ["t", "m", "t"],
        })
        eixo_x = alt.X("aeronave:N", sort=tabela["aeronave"].tolist(), title=None,
                       axis=alt.Axis(labelAngle=-90, labelFontSize=9))
        regras = alt.Chart(linhas).mark_rule(color="#6B7280").encode(
            y="valor:Q",
            tooltip=[alt.Tooltip("nome:N", title="Linha"),
                     alt.Tooltip("valor:Q", title=titulo, format=f",.{formato}f")],
            strokeDash=alt.StrokeDash("traco:N", scale=alt.Scale(
                domain=["m", "t"], range=[[1, 0], [5, 4]]), legend=None),
        )
        pontos = alt.Chart(tabela).mark_point(filled=True, size=70, opacity=1).encode(
            x=eixo_x,
            y=alt.Y(f"{campo}:Q", title=titulo, scale=alt.Scale(zero=False)),
            color=alt.Color("fora:N", scale=alt.Scale(
                domain=["Dentro dos limites", "Fora dos limites"],
                range=[cor_modelo, "#1F2430"]), legend=alt.Legend(orient="top", title=None)),
            shape=alt.Shape("fora:N", scale=alt.Scale(
                domain=["Dentro dos limites", "Fora dos limites"],
                range=["circle", "diamond"]), legend=None),
            tooltip=[alt.Tooltip("aeronave:N", title="Aeronave"),
                     alt.Tooltip(f"{campo}:Q", title=titulo, format=f",.{formato}f"),
                     alt.Tooltip("fora:N", title="Situação")],
        )
        rotulo_fora = alt.Chart(fora).mark_text(dy=-12, color="#1F2430", fontWeight="bold").encode(
            x=eixo_x, y=f"{campo}:Q", text="aeronave:N",
        )
        # Com muitos pontos fora, os nomes se sobrepõem; a lista vai embaixo do gráfico.
        camadas = [regras, pontos] + ([rotulo_fora] if len(fora) <= 6 else [])
        grafico = alt.layer(*camadas).properties(height=320)
        return grafico, fora, (lic, media, lsc)

    for campo, titulo, formato in (("peso", "Peso básico (kgf)", 0), ("cg", "C.G. (% MAC)", 2)):
        grafico, fora, (lic, media, lsc) = carta(campo, titulo, formato)
        with st.container(border=True):
            st.markdown(f"**Carta de controle · {titulo}**")
            st.caption(
                (f"Mediana {media:,.{formato}f} · limites {lic:,.{formato}f} a {lsc:,.{formato}f}")
                .replace(",", "X").replace(".", ",").replace("X", ".")
            )
            mostrar_grafico(grafico, use_container_width=True)
            if not fora.empty:
                st.caption("Fora dos limites: " + ", ".join(fora["aeronave"]))


@protegido
def tela_graficos():
    import altair as alt

    cabecalho(
        "Gráficos da frota",
        "Peso básico e C.G. de cada aeronave pela última pesagem e revisão.",
    )
    if st.button("Atualizar", key="atualizar_graficos"):
        invalidar_cache_fichas()
        st.rerun()

    frota = montar_resumo_frota()
    validas = frota[(frota["Peso básico (kgf)"] > 0) & (frota["C.G. (% MAC)"] > 0)]
    if validas.empty:
        st.info("Ainda não há fichas com dados de célula de carga para montar os gráficos.")
        return

    modelos_disponiveis = [m for m in CORES_MODELOS if m in set(validas["Modelo"])]
    modelos_disponiveis += sorted(set(validas["Modelo"]) - set(modelos_disponiveis))
    modelos = st.pills(
        "Modelos", modelos_disponiveis, selection_mode="multi",
        default=modelos_disponiveis, key="graficos_modelos",
    ) or modelos_disponiveis
    selecionadas = validas[validas["Modelo"].isin(modelos)].copy()
    selecionadas["Conferir"] = motivos_fora_do_padrao(selecionadas)
    para_conferir = selecionadas[selecionadas["Conferir"] != ""]
    ignorar = st.toggle(
        "Deixar de fora as fichas com valores fora do padrão do modelo",
        value=True, key="graficos_ignorar_fora_padrao",
        disabled=para_conferir.empty,
    )
    dados = selecionadas[selecionadas["Conferir"] == ""] if ignorar else selecionadas
    dados = dados.drop(columns=["Conferir"])

    sem_celulas = len(frota) - len(validas)
    st.caption(
        f"{len(dados)} aeronaves nos gráficos, cada uma com a sua ficha mais recente "
        "(aprovada ou não)."
        + (f" {sem_celulas} ficha(s) sem pesos de célula de carga ficaram de fora." if sem_celulas else "")
        + (f" {len(para_conferir)} ficha(s) para conferir, listadas no fim da aba Frota hoje."
           if not para_conferir.empty else "")
    )

    escala_cor = alt.Scale(
        domain=modelos_disponiveis,
        range=[CORES_MODELOS.get(m, "#8A8F98") for m in modelos_disponiveis],
    )
    cor = alt.Color("modelo:N", scale=escala_cor, legend=alt.Legend(orient="top", title=None))
    aba_hoje, aba_var, aba_hist, aba_carta = st.tabs(
        ["Frota hoje", "Variação desde a pesagem anterior", "Histórico por aeronave", "Carta de controle"]
    )
    historico = montar_historico_frota()
    with aba_var:
        aba_variacao(alt, historico, modelos, escala_cor, cor)
    with aba_hist:
        aba_historico(alt, historico, modelos, escala_cor)
    with aba_carta:
        aba_carta_controle(alt, selecionadas, modelos)
    with aba_hoje:
        mais_pesada = dados.loc[dados["Peso básico (kgf)"].idxmax()]
        mais_leve = dados.loc[dados["Peso básico (kgf)"].idxmin()]
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Peso básico médio", formatar_kg(dados["Peso básico (kgf)"].mean()))
        k2.metric("C.G. médio", formatar_pct(dados["C.G. (% MAC)"].mean()))
        k3.metric(
            "Mais pesada", mais_pesada["Aeronave"],
            formatar_kg(mais_pesada["Peso básico (kgf)"]), delta_color="off", delta_arrow="off",
        )
        k4.metric(
            "Mais leve", mais_leve["Aeronave"],
            formatar_kg(mais_leve["Peso básico (kgf)"]), delta_color="off", delta_arrow="off",
        )

        # O Altair lê "." no nome do campo como campo aninhado; os gráficos usam nomes simples.
        grafico = dados.rename(columns={
            "Aeronave": "aeronave", "Modelo": "modelo", "Pesagem": "pesagem",
            "Revisão": "revisao", "Data da pesagem": "data",
            "Peso básico (kgf)": "peso", "C.G. (% MAC)": "cg",
        })

        dica = [
            alt.Tooltip("aeronave:N", title="Aeronave"),
            alt.Tooltip("modelo:N", title="Modelo"),
            alt.Tooltip("peso:Q", title="Peso básico (kgf)", format=",.0f"),
            alt.Tooltip("cg:Q", title="C.G. (% MAC)", format=".2f"),
            alt.Tooltip("pesagem:N", title="Pesagem"),
            alt.Tooltip("revisao:N", title="Revisão"),
            alt.Tooltip("data:T", title="Data da pesagem", format="%d/%m/%Y"),
        ]

        por_modelo = (
            grafico.groupby("modelo", as_index=False)
            .agg(aeronaves=("aeronave", "count"), peso=("peso", "mean"), cg=("cg", "mean"))
        )

        def barras_por_modelo(campo, titulo, formato):
            base = alt.Chart(por_modelo).encode(
                y=alt.Y("modelo:N", sort=modelos_disponiveis, title=None),
                x=alt.X(f"{campo}:Q", title=titulo),
                tooltip=[
                    alt.Tooltip("modelo:N", title="Modelo"),
                    alt.Tooltip("aeronaves:Q", title="Aeronaves"),
                    alt.Tooltip(f"{campo}:Q", title=titulo, format=formato),
                ],
            )
            barras = base.mark_bar(cornerRadiusEnd=4, height=22).encode(
                color=alt.Color("modelo:N", scale=escala_cor, legend=None)
            )
            rotulos = base.mark_text(align="left", dx=6, color="#1F2430").encode(
                text=alt.Text(f"{campo}:Q", format=formato)
            )
            return (barras + rotulos).properties(height=60 * len(por_modelo) + 30)

        c1, c2 = st.columns(2)
        with c1.container(border=True):
            st.markdown("**Peso básico médio por modelo**")
            mostrar_grafico(
                barras_por_modelo("peso", "kgf", ",.0f"), use_container_width=True
            )
        with c2.container(border=True):
            st.markdown("**C.G. médio por modelo**")
            mostrar_grafico(
                barras_por_modelo("cg", "% MAC", ".2f"), use_container_width=True
            )

        with st.container(border=True):
            st.markdown("**Peso × C.G. de cada aeronave**")
            st.caption("As linhas tracejadas marcam a média da seleção. Passe o mouse para ver a aeronave.")
            pontos = alt.Chart(grafico).mark_circle(size=90, opacity=0.85, stroke="white", strokeWidth=1.5).encode(
                x=alt.X("cg:Q", title="C.G. (% MAC)", scale=alt.Scale(zero=False)),
                y=alt.Y("peso:Q", title="Peso básico (kgf)", scale=alt.Scale(zero=False)),
                color=cor,
                tooltip=dica,
            )
            media_x = alt.Chart(pd.DataFrame({"v": [dados["C.G. (% MAC)"].mean()]})).mark_rule(
                strokeDash=[4, 4], color="#6B7280"
            ).encode(x="v:Q")
            media_y = alt.Chart(pd.DataFrame({"v": [dados["Peso básico (kgf)"].mean()]})).mark_rule(
                strokeDash=[4, 4], color="#6B7280"
            ).encode(y="v:Q")
            mostrar_grafico((media_x + media_y + pontos).properties(height=380), use_container_width=True)

        c3, c4 = st.columns(2)
        with c3.container(border=True):
            st.markdown("**10 aeronaves mais pesadas**")
            top = grafico.nlargest(10, "peso")
            # Pontos em vez de barras: as dez ficam perto umas das outras e a barra
            # partindo do zero esconderia a diferença.
            base = alt.Chart(top).encode(
                y=alt.Y("aeronave:N", title=None,
                        sort=alt.EncodingSortField("peso", order="descending")),
                x=alt.X("peso:Q", title="Peso básico (kgf)",
                        scale=alt.Scale(zero=False, nice=True, padding=30)),
                tooltip=dica,
            )
            mostrar_grafico(
                (
                    base.mark_circle(size=120, opacity=1).encode(color=cor)
                    + base.mark_text(align="left", dx=10, color="#1F2430").encode(
                        text=alt.Text("peso:Q", format=",.0f")
                    )
                ).properties(height=330),
                use_container_width=True,
            )
        with c4.container(border=True):
            st.markdown("**Distribuição do C.G.**")
            mostrar_grafico(
                alt.Chart(grafico).mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4).encode(
                    x=alt.X("cg:Q", bin=alt.Bin(step=1), title="C.G. (% MAC), faixas de 1 %"),
                    y=alt.Y("count():Q", title="Aeronaves", stack=True),
                    color=cor,
                    tooltip=[alt.Tooltip("modelo:N", title="Modelo"), alt.Tooltip("count():Q", title="Aeronaves")],
                ).properties(height=330),
                use_container_width=True,
            )

        com_data = grafico.dropna(subset=["data"]).copy()
        if not com_data.empty:
            with st.container(border=True):
                st.markdown("**Ano da última pesagem**")
                sem_data = len(dados) - len(com_data)
                st.caption(
                    "Quantas aeronaves tiveram a última pesagem em cada ano."
                    + (f" {sem_data} sem data válida na ficha." if sem_data else "")
                )
                com_data["ano"] = com_data["data"].dt.year.astype(str)
                mostrar_grafico(
                    alt.Chart(com_data).mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4).encode(
                        x=alt.X("ano:O", title=None, axis=alt.Axis(labelAngle=0)),
                        y=alt.Y("count():Q", title="Aeronaves"),
                        color=cor,
                        tooltip=[alt.Tooltip("ano:O", title="Ano"), alt.Tooltip("modelo:N", title="Modelo"),
                                 alt.Tooltip("count():Q", title="Aeronaves")],
                    ).properties(height=260),
                    use_container_width=True,
                )

        if not para_conferir.empty:
            with st.expander(f"Fichas para conferir ({len(para_conferir)})"):
                st.caption(
                    "Peso mais de 4 % longe da mediana do modelo ou C.G. mais de 4 % MAC "
                    "longe da mediana do modelo. Costuma ser célula de carga, canela ou "
                    "item digitado errado na ficha."
                )
                st.dataframe(
                    para_conferir.sort_values(["Modelo", "Aeronave"]),
                    hide_index=True, use_container_width=True,
                    column_config={
                        "Data da pesagem": st.column_config.DateColumn(format="DD/MM/YYYY"),
                        "Peso básico (kgf)": st.column_config.NumberColumn(format="%.1f"),
                        "C.G. (% MAC)": st.column_config.NumberColumn(format="%.2f"),
                    },
                )

        with st.expander("Ver tabela das aeronaves"):
            tabela = dados.sort_values("Peso básico (kgf)", ascending=False)
            st.dataframe(
                tabela, hide_index=True, use_container_width=True,
                column_config={
                    "Data da pesagem": st.column_config.DateColumn(format="DD/MM/YYYY"),
                    "Peso básico (kgf)": st.column_config.NumberColumn(format="%.1f"),
                    "C.G. (% MAC)": st.column_config.NumberColumn(format="%.2f"),
                },
            )
            st.download_button(
                "Baixar CSV",
                tabela.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
                file_name="frota_peso_cg.csv",
                mime="text/csv",
            )


def renderizar_ficha_visualizacao(
    prefixo,
    pesagem,
    revisao,
    row,
    modo_aprovacao=False,
    destacar_erros=False,
):
    st.subheader(f"{prefixo} · Pesagem {pesagem} · Revisão {revisao}")
    fluxo = buscar_fluxo_ficha(prefixo, pesagem, revisao)
    st.caption(
        f"Emissor (responsável, último a gerar ou editar): {fluxo['gerador_nome']}"
        f" | Aprovada por: {fluxo['aprovador_nome']}"
    )
    if nivel_atual() != NIVEL_ENG_OPERACOES:
        historico = carregar_historico_ficha(prefixo, pesagem, revisao)
        with st.expander("Histórico de alterações"):
            if historico:
                st.dataframe(
                    pd.DataFrame(historico), hide_index=True, use_container_width=True
                )
            else:
                st.caption(
                    "Nenhuma alteração registrada. O histórico começa a contar "
                    "a partir de 14/10/2026."
                )
    campos_pendentes = carregar_campos_com_erro(prefixo, pesagem, revisao)
    if campos_pendentes and not modo_aprovacao:
        st.warning(
            f"Ficha devolvida ao emissor ({fluxo['gerador_nome']}) para correção: "
            + "; ".join(campo['descricao'] for campo in campos_pendentes)
        )
    sinalizacoes_atuais = {
        campo['campo']: campo
        for campo in campos_pendentes
    }
    marcacoes_campos = {}
    rotulos_campos = {}

    def mostrar_campo(campo, rotulo, valor, container=st, permitir_erro=True):
        if not modo_aprovacao or not permitir_erro:
            if destacar_erros and campo in sinalizacoes_atuais:
                container.markdown(f":red[🔴 {rotulo} precisa de correção]")
            container.text_input(
                rotulo,
                value=safe_str(valor),
                disabled=True,
                key=(
                    f"visual_{'aprovacao' if modo_aprovacao else 'consulta'}_"
                    f"{prefixo}_{pesagem}_{revisao}_{campo}_"
                    f"{'erros' if destacar_erros else 'normal'}"
                ),
            )
            return

        coluna_valor, coluna_erro = container.columns([5, 1])
        coluna_valor.text_input(
            rotulo,
            value=safe_str(valor),
            disabled=True,
            key=f"visual_{prefixo}_{pesagem}_{revisao}_{campo}",
        )
        marcacoes_campos[campo] = coluna_erro.checkbox(
            "Erro",
            value=campo in sinalizacoes_atuais,
            key=f"aprovacao_erro_{prefixo}_{pesagem}_{revisao}_{campo}",
            help="Marque para solicitar ao emissor a correção deste campo.",
        )
        if marcacoes_campos[campo]:
            coluna_valor.markdown(f":red[🔴 {rotulo} marcado para correção]")
        rotulos_campos[campo] = f"{rotulo}: {safe_str(valor)}"

    def tabela_itens_aprovacao(tipo, campo_desc, campo_peso, campo_braco, limite):
        itens = []
        for indice in range(1, limite + 1):
            descricao = row.get(f'{tipo} {campo_desc} {indice}', None)
            if pd.isna(descricao) or not str(descricao).strip():
                continue
            peso = safe_float(row.get(f'{tipo} {campo_peso} {indice}', 0))
            braco = safe_float(row.get(f'{tipo} {campo_braco} {indice}', 0))
            itens.append((indice, safe_str(descricao), peso, braco))
        if not itens:
            st.caption("Nenhum item.")
            return
        larguras = [0.5, 4, 1.2, 1.2, 1.5, 0.8]
        cab = st.columns(larguras, vertical_alignment="center")
        for coluna, titulo in zip(cab, ["#", "Descrição", "Peso (kgf)", "Braço (in)", "Momento", "Erro"]):
            coluna.markdown(f"<span style='color:#6B7280;font-size:.82rem;font-weight:600'>{titulo}</span>", unsafe_allow_html=True)
        for indice, descricao, peso, braco in itens:
            campo = f"{tipo} {campo_desc} {indice}"
            campos_item = {f"{tipo} {c} {indice}" for c in (campo_desc, campo_peso, campo_braco)}
            marcado_antes = bool(campos_item & set(sinalizacoes_atuais))
            st.markdown("<hr style='margin:.1rem 0;border:none;border-top:1px solid #EEF0F3'>", unsafe_allow_html=True)
            c = st.columns(larguras, vertical_alignment="center")
            c[0].write(str(indice))
            c[1].write(descricao)
            c[2].write(f"{peso:,.2f}")
            c[3].write(f"{braco:,.2f}")
            c[4].write(f"{peso * braco:,.2f}")
            marcado = c[5].checkbox(
                "Erro", value=marcado_antes, label_visibility="collapsed",
                key=f"aprovacao_erro_{prefixo}_{pesagem}_{revisao}_{campo}",
            )
            marcacoes_campos[campo] = marcado
            rotulos_campos[campo] = f"{tipo} item {indice}: {descricao}"
            if marcado:
                c[1].markdown(":red[🔴 marcado para correção]")

    aba1, aba2, aba3, aba4, aba5 = st.tabs(["Dados Gerais", "Células de Carga", "Deductions", "Additions", "Weighing Report"])
    info_aero = dict_tipos_aeronave.get(prefixo, {})
    lopa, config_lopa = separar_campos_lopa(row)

    with aba1:
        c_p, c_r = st.columns(2)
        mostrar_campo(
            "Pesagem", "Número da Pesagem:", row.get('Pesagem', ''),
            c_p, permitir_erro=False,
        )
        mostrar_campo(
            "Revisao", "Número da Revisão:", row.get('Revisao', ''),
            c_r, permitir_erro=False,
        )

        c1, c2, c3 = st.columns(3)
        mostrar_campo("Data da ficha", "Data de emissão:", row.get('Data da ficha', row.get('Data_da_Pesagem', '')), c1)
        mostrar_campo("Pesado Por", "Pesado por:", row.get('Pesado Por', row.get('WEIGHED BY', '')), c2)
        mostrar_campo("Local da pesagem", "Local da pesagem:", row.get('Local da pesagem', ''), c3)
        
        c4, c5, c6 = st.columns(3)
        mostrar_campo("Data_da_Pesagem", "Data da pesagem:", row.get('Data_da_Pesagem', ''), c4)
        mostrar_campo("lopa", "LOPA:", lopa, c5)
        mostrar_campo("config_lopa", "Configuração LOPA:", config_lopa, c6)
        
        c7, c8, c9 = st.columns(3)
        mostrar_campo("VRBL", "VRBL. NUMBER:", safe_str(row.get('VRBL')) or info_aero.get('vrbl', ''), c7)
        mostrar_campo("SERIAL", "SERIAL NUMBER:", safe_str(row.get('SERIAL')) or info_aero.get('serial', ''), c8)
        mostrar_campo("LINE", "LINE NUMBER:", safe_str(row.get('LINE')) or info_aero.get('line', ''), c9)
        mostrar_campo("Motivo", "Razão para emissão:", row.get('Motivo', ''))

    with aba2:
        campos_pesagem_1 = [
            ("Peso nariz LH", "Nariz LH"),
            ("Peso Nariz RH", "Nariz RH"),
            ("Peso MLG RH 1 ", "RH MLG 1"),
            ("Peso MLG LH 1", "LH MLG 1"),
            ("Peso MLG RH 2", "RH MLG 2"),
            ("Peso MLG LH2", "LH MLG 2"),
        ]
        campos_pesagem_2 = [
            ("Peso nariz LH pesagem 2", "Nariz LH"),
            ("Peso Nariz RH pesagem 2", "Nariz RH"),
            ("Peso MLG RH 1  pesagem 2", "RH MLG 1"),
            ("Peso MLG LH 1 pesagem 2", "LH MLG 1"),
            ("Peso MLG RH 2 pesagem 2", "RH MLG 2"),
            ("Peso MLG LH2 pesagem 2", "LH MLG 2"),
        ]
        alternativas_totais = {
            "Peso MLG RH 1 ": ["Peso MLG RH 1 ", "Peso MLG RH 1"],
            "Peso MLG LH 1": ["Peso MLG LH 1", "Peso MLG LH 1 "],
            "Peso MLG RH 2": ["Peso MLG RH 2", "Peso MLG RH 2 "],
            "Peso MLG LH2": ["Peso MLG LH2", "Peso MLG LH 2"],
            "Peso MLG RH 1  pesagem 2": ["Peso MLG RH 1  pesagem 2", "Peso MLG RH 1 pesagem 2"],
        }
        for titulo, campos in (
            ("Pesagem 01", campos_pesagem_1),
            ("Pesagem 02", campos_pesagem_2),
        ):
            st.markdown(f"**{titulo}**")
            ordem = [campos[0], campos[1], campos[2], campos[4], campos[3], campos[5]]

            def render_consulta(container, indice, ordem=ordem):
                campo, rotulo = ordem[indice]
                campo = get_real_col(alternativas_totais.get(campo, [campo]))
                valor = row.get(campo, "")
                mostrar_campo(campo, rotulo, valor, container)
                return valor

            grade_celulas(render_consulta)

        st.markdown("**Medidas das canelas**")
        canela_lh, canela_rh = st.columns(2)
        mostrar_campo(
            "Canela MLG LH", "Canela MLG LH (in)",
            row.get("Canela MLG LH", ""), canela_lh,
        )
        mostrar_campo(
            "Canela MLG RH", "Canela MLG RH (in)",
            row.get("Canela MLG RH", ""), canela_rh,
        )

    with aba3:
        deducoes_lista = []
        for i in range(1, MAX_DEDUCTIONS + 1):
            desc = row.get(f'Deductions description {i}', None)
            peso = row.get(f'Deductions Weigth {i}', None)
            arm = row.get(f'Deductions arm {i}', None)
            if pd.notna(desc) and str(desc).strip() != '':
                deducoes_lista.append({"Descrição": desc, "Peso [kgf]": safe_float(peso), "Arm [pol]": safe_float(arm)})
        if modo_aprovacao:
            tabela_itens_aprovacao("Deductions", "description", "Weigth", "arm", MAX_DEDUCTIONS)
        elif deducoes_lista:
            st.dataframe(pd.DataFrame(deducoes_lista), use_container_width=True, hide_index=True)

    with aba4:
        adicoes_lista = []
        for i in range(1, MAX_ADDITIONS + 1):
            desc = row.get(f'Additions Description {i}', None)
            peso = row.get(f'Additions weigth {i}', None)
            arm = row.get(f'Additions arm {i}', None)
            if pd.notna(desc) and str(desc).strip() != '':
                adicoes_lista.append({"Descrição": desc, "Peso [kgf]": safe_float(peso), "Arm [pol]": safe_float(arm)})
        if modo_aprovacao:
            tabela_itens_aprovacao("Additions", "Description", "weigth", "arm", MAX_ADDITIONS)
        elif adicoes_lista:
            st.dataframe(pd.DataFrame(adicoes_lista), use_container_width=True, hide_index=True)

    with aba5:
        try:
            lh_1 = safe_float(row.get('Peso MLG LH 1', 0))
            lh_2 = safe_float(row.get('Peso MLG LH2', 0))
            lh_1_p2 = safe_float(row.get('Peso MLG LH 1 pesagem 2', 0))
            lh_2_p2 = safe_float(row.get('Peso MLG LH2 pesagem 2', 0))
            rh_1 = safe_float(row.get('Peso MLG RH 1 ', 0))
            rh_2 = safe_float(row.get('Peso MLG RH 2', 0))
            rh_1_p2 = safe_float(row.get('Peso MLG RH 1  pesagem 2', row.get('Peso MLG RH 1 pesagem 2', 0)))
            rh_2_p2 = safe_float(row.get('Peso MLG RH 2 pesagem 2', 0))
            nose_lh_1 = safe_float(row.get('Peso nariz LH', 0))
            nose_lh_2 = safe_float(row.get('Peso nariz LH pesagem 2', 0))
            nose_rh_1 = safe_float(row.get('Peso Nariz RH', 0))
            nose_rh_2 = safe_float(row.get('Peso Nariz RH pesagem 2', 0))
            
            nariz_lh = (nose_lh_1 + nose_lh_2)/2
            nariz_rh = (nose_rh_1 + nose_rh_2)/2
            nose = nariz_lh + nariz_rh
            
            tail = safe_float(row.get('Tail', 0))
            canela_mlg_lh = safe_float(row.get('Canela MLG LH', 0))
            canela_mlg_rh = safe_float(row.get('Canela MLG RH', 0))
        except:
            lh, rh, nose, tail = 0.0, 0.0, 0.0, 0.0
            canela_mlg_lh, canela_mlg_rh = 0.0, 0.0

        arm_b_lh = (canela_mlg_lh * 0.1304) + 706.17
        arm_b_rh = (canela_mlg_rh * 0.1304) + 706.17
        tipo_aeronave = info_aero.get('modelo', "")
        arm_a = info_aero.get('armnose', 93.0000)
        arm_c = 1298.9000

        rh = (rh_1 + rh_2 + rh_1_p2 + rh_2_p2) / 2
        lh = (lh_1 + lh_2 + lh_1_p2 + lh_2_p2) / 2

        m_lh = lh * arm_b_lh
        m_rh = rh * arm_b_rh
        m_nose = nose * arm_a
        m_tail = tail * arm_c

        tot_reg_weight = lh + rh + nose + tail
        tot_reg_moment = m_lh + m_rh + m_nose + m_tail
        tot_reg_arm = tot_reg_moment / tot_reg_weight if tot_reg_weight > 0 else 0.0

        level_correction_angle = normalizar_angulo_level_correction(
            row.get('Graus correção do cg', '')
        )
        level_correction_factor, level_correction_moment, level_correction_side = (
            calcular_level_correction(level_correction_angle, tot_reg_weight)
        )

        ded_weight = sum([item["Peso [kgf]"] for item in deducoes_lista])
        ded_moment = sum([item["Peso [kgf]"] * item["Arm [pol]"] for item in deducoes_lista])
        if level_correction_side == "deductions":
            ded_moment += level_correction_moment
        ded_arm = ded_moment / ded_weight if ded_weight > 0 else 0.0

        flaps_ativo = any("flap" in str(item.get("Descrição", "")).casefold() for item in adicoes_lista)
        momento_extra_flaps = 0.0
        if flaps_ativo:
            momento_extra_flaps = (
                momento_flaps_do_registro(row)
                or buscar_momento_flaps_banco(prefixo, tipo_aeronave)
            )

        add_weight = sum([item["Peso [kgf]"] for item in adicoes_lista])
        add_moment = sum([item["Peso [kgf]"] * item["Arm [pol]"] for item in adicoes_lista]) + momento_extra_flaps
        if level_correction_side == "additions":
            add_moment += level_correction_moment
        add_arm = add_moment / add_weight if add_weight > 0 else 0.0

        basic_weight = tot_reg_weight + add_weight - ded_weight
        basic_moment = tot_reg_moment + add_moment - ded_moment
        basic_arm = basic_moment / basic_weight if basic_weight > 0 else 0.0
        cg_mac = ((basic_arm - 627.1) / 1.558) if basic_arm > 0 else 0.0

        report_data = {
            "Reaction / Item": ["LH", "RH", "NOSE", "TAIL", "TOTAL REGISTERED", "LEVEL CORRECTION", "DEDUCTIONS", "ADDITIONS", "AIRCRAFT BASIC WEIGHT"],
            "Weight (kgf)": [lh, rh, nose, tail, tot_reg_weight, 0.0, ded_weight, add_weight, basic_weight],
            "Arm (inch)": [arm_b_lh, arm_b_rh, arm_a, arm_c, tot_reg_arm, level_correction_factor, ded_arm, add_arm, basic_arm],
            "Moment (kgf x inch)": [m_lh, m_rh, m_nose, m_tail, tot_reg_moment, level_correction_moment, ded_moment, add_moment, basic_moment]
        }
        mostrar_relatorio_pesagem(report_data, basic_weight, basic_arm, cg_mac)
            
        pesagem_atual = safe_float(row.get('Pesagem', 0))
        revisao_atual = safe_float(row.get('Revisao', 0))
        
        df_aero = carregar_fichas_prefixo(prefixo)
        df_aero['Pesagem_num'] = pd.to_numeric(df_aero['Pesagem'], errors='coerce').fillna(0)
        df_aero['Revisao_num'] = pd.to_numeric(df_aero['Revisao'], errors='coerce').fillna(0)
        df_prev = df_aero[(df_aero['Pesagem_num'] < pesagem_atual) | ((df_aero['Pesagem_num'] == pesagem_atual) & (df_aero['Revisao_num'] < revisao_atual))]
        df_prev = df_prev.sort_values(by=['Pesagem_num', 'Revisao_num'])
        
        ultima_p = str(int(df_prev.iloc[-1]['Pesagem_num'])) if not df_prev.empty else ""
        ultima_r = str(int(df_prev.iloc[-1]['Revisao_num'])) if not df_prev.empty else ""
        up_salva, ur_salva = ultima_ficha_gravada(row)
        if up_salva:
            ultima_p, ultima_r = up_salva, ur_salva
        lopa, config_lopa = separar_campos_lopa(row)

        deducoes_excel = [
            {'desc': d["Descrição"], 'w': d["Peso [kgf]"], 'a': d["Arm [pol]"], 'm': d["Peso [kgf]"] * d["Arm [pol]"]}
            for d in deducoes_lista
        ]
        adicoes_excel = [
            {'desc': a["Descrição"], 'w': a["Peso [kgf]"], 'a': a["Arm [pol]"], 'm': a["Peso [kgf]"] * a["Arm [pol]"] + (momento_extra_flaps if "flap" in a["Descrição"].casefold() else 0)}
            for a in adicoes_lista
        ]
        if level_correction_side == "deductions":
            deducoes_excel.append({
                'desc': LEVEL_CORRECTION_LABEL, 'w': 0.0,
                'a': level_correction_factor, 'm': level_correction_moment,
            })
        elif level_correction_side == "additions":
            adicoes_excel.append({
                'desc': LEVEL_CORRECTION_LABEL, 'w': 0.0,
                'a': level_correction_factor, 'm': level_correction_moment,
            })

        dados_excel = {
            "prefixo": prefixo, "modelo": tipo_aeronave, "pesado_por": safe_str(row.get('Pesado Por', '')),
            "local": safe_str(row.get('Local da pesagem', '')), "data": safe_str(row.get('Data_da_Pesagem', '')),
            "config_lopa": config_lopa, "lopa": lopa,
            "razao": safe_str(row.get('Motivo', '')), "cg_mac": cg_mac,
            "arm_a": arm_a, "arm_b_lh": arm_b_lh, "arm_b_rh": arm_b_rh,
            "issue_date": safe_str(row.get('Data da ficha', '')),
            "ultima_pesagem": ultima_p, "ultima_revisao": ultima_r, 
            "vrbl_number": safe_str(safe_str(row.get('VRBL')) or info_aero.get('vrbl', '')),
            "serial_number": safe_str(safe_str(row.get('SERIAL')) or info_aero.get('serial', '')),
            "line_number": safe_str(safe_str(row.get('LINE')) or info_aero.get('line', '')),
            "numero_pesagem": str(int(pesagem_atual)) if pesagem_atual else "",
            "revisao": str(int(revisao_atual)),
            "reactions": {
                'LH': {'weight': lh, 'arm': arm_b_lh, 'moment': m_lh},
                'RH': {'weight': rh, 'arm': arm_b_rh, 'moment': m_rh},
                'NOSE': {'weight': nose, 'arm': arm_a, 'moment': m_nose},
                'TAIL': {'weight': tail, 'arm': arm_c, 'moment': m_tail},
                'TOTAL REGISTERED': {'weight': tot_reg_weight, 'arm': tot_reg_arm, 'moment': tot_reg_moment},
                'DEDUCTIONS': {'weight': ded_weight, 'arm': ded_arm, 'moment': ded_moment},
                'ADDITIONS': {'weight': add_weight, 'arm': add_arm, 'moment': add_moment},
                'AIRCRAFT BASIC WEIGHT': {'weight': basic_weight, 'arm': basic_arm, 'moment': basic_moment}
            },
            "deductions": deducoes_excel,
            "additions": adicoes_excel
        }
        dados_excel["assinatura_emissor"] = carregar_assinatura_usuario(
            fluxo['gerador_usuario']
        )
        dados_excel["assinatura_aprovador"] = carregar_assinatura_usuario(
            fluxo['aprovador_usuario']
        )
        
        if modo_aprovacao:
            st.divider()
            st.subheader("Campos marcados para correção")
            campos_selecionados = {
                campo: rotulos_campos[campo]
                for campo, marcado in marcacoes_campos.items()
                if marcado
            }
            if campos_selecionados:
                for rotulo in campos_selecionados.values():
                    st.markdown(f"- {rotulo}")
                if st.button(
                    "Enviar para o emissor alterar",
                    type="secondary",
                    disabled=not pode_emitir(),
                    key=f"enviar_correcao_{prefixo}_{pesagem}_{revisao}",
                ):
                    salvar_campos_com_erro(
                        prefixo,
                        pesagem,
                        revisao,
                        campos_selecionados,
                        st.session_state['usuario_id'],
                    )
                    st.success(
                        "Os campos marcados foram enviados ao emissor para alteração."
                    )
                    st.rerun()
            elif campos_pendentes:
                st.info(
                    "Nenhum campo está marcado. Envie a lista vazia para remover "
                    "as sinalizações existentes."
                )
                if st.button(
                    "Limpar sinalizações",
                    disabled=not pode_emitir(),
                    key=f"limpar_correcao_{prefixo}_{pesagem}_{revisao}",
                ):
                    salvar_campos_com_erro(
                        prefixo,
                        pesagem,
                        revisao,
                        {},
                        st.session_state['usuario_id'],
                    )
                    st.success("As sinalizações foram removidas.")
                    st.rerun()
            else:
                st.caption("Marque “Erro” ao lado dos campos que precisam de correção.")

            pode_aprovar = (
                pode_emitir()
                and not campos_pendentes
                and not campos_selecionados
                and fluxo['gerador_usuario'] != st.session_state['usuario_id']
            )
            conferiu_cg = False
            if campos_pendentes:
                st.error(
                    "A aprovação está bloqueada até o emissor corrigir "
                    "os campos sinalizados."
                )
            elif fluxo['gerador_usuario'] == st.session_state['usuario_id']:
                st.warning("Quem emitiu a ficha não pode aprová-la.")
            elif not pode_emitir():
                st.caption("Seu perfil não pode aprovar fichas.")
            else:
                conferiu_cg = st.checkbox(
                    "Confirmei os dados e o valor do C.G. (% MAC).",
                    key=f"conferiu_cg_{prefixo}_{pesagem}_{revisao}",
                )

            coluna_excel, coluna_aprovacao = st.columns(2)
            botoes_exportacao(
                coluna_excel, lambda: dados_excel, prefixo, pesagem, revisao,
                f"aprovacao_{prefixo}_{pesagem}_{revisao}",
            )
            if coluna_aprovacao.button(
                "Aprovar ficha",
                type="primary",
                disabled=not pode_aprovar or not conferiu_cg,
                key=f"aprovar_{prefixo}_{pesagem}_{revisao}",
                use_container_width=True,
            ):
                aprovar_ficha(
                    prefixo,
                    pesagem,
                    revisao,
                    st.session_state['usuario_id'],
                )
                st.session_state["ficha_recem_aprovada"] = (prefixo, pesagem, revisao)
                st.rerun()
        else:
            botoes_exportacao(
                st, lambda: dados_excel, prefixo, pesagem, revisao,
                f"consulta_{prefixo}_{pesagem}_{revisao}_{destacar_erros}",
            )

def gerar_pdf_por_template(dados, caminho_template="exemplo_ficha.xlsx"):
    """Gera o Excel oficial e o converte em PDF com o LibreOffice."""
    import shutil
    import subprocess
    import tempfile

    executavel = shutil.which("soffice") or shutil.which("libreoffice")
    if not executavel:
        raise RuntimeError("LibreOffice não está instalado no servidor.")
    with tempfile.TemporaryDirectory() as pasta:
        caminho_xlsx = os.path.join(pasta, "ficha.xlsx")
        wb = openpyxl.load_workbook(
            io.BytesIO(gerar_excel_por_template(dados, caminho_template))
        )
        for planilha in wb.worksheets:
            # Ajusta a largura da ficha a uma página, como na impressão do Excel.
            planilha.sheet_properties.pageSetUpPr.fitToPage = True
            planilha.page_setup.fitToWidth = 1
            planilha.page_setup.fitToHeight = 0
            planilha.page_setup.paperSize = planilha.PAPERSIZE_A4
            # Só a ficha (colunas A:AD); à direita ficam tabelas auxiliares.
            planilha.print_area = f"A1:AD{min(planilha.max_row, 130)}"
        wb.save(caminho_xlsx)
        subprocess.run(
            [
                executavel, "--headless", "--norestore",
                f"-env:UserInstallation=file://{pasta}/perfil",
                "--convert-to", "pdf", "--outdir", pasta, caminho_xlsx,
            ],
            check=True, capture_output=True, timeout=120,
        )
        with open(os.path.join(pasta, "ficha.pdf"), "rb") as arquivo:
            return arquivo.read()


def nome_arquivo_ficha(prefixo, pesagem, revisao):
    return (
        f"{safe_str(prefixo)}_Pesagem_{normalizar_chave_ficha(pesagem)}"
        f"_Revisao_{normalizar_chave_ficha(revisao)}"
    )


def botoes_exportacao(container, obter_dados, prefixo, pesagem, revisao, chave):
    # obter_dados roda no clique do download, fora da execução da tela, onde
    # st.session_state não está disponível: passe os dados já capturados.
    nome = nome_arquivo_ficha(prefixo, pesagem, revisao)
    col_excel, col_pdf = container.columns(2)
    col_excel.download_button(
        "Exportar Excel",
        icon=":material/table_view:",
        data=lambda: gerar_excel_por_template(obter_dados(), "exemplo_ficha.xlsx"),
        file_name=f"{nome}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key=f"excel_{chave}",
    )
    col_pdf.download_button(
        "Exportar PDF",
        icon=":material/picture_as_pdf:",
        data=lambda: gerar_pdf_por_template(obter_dados(), "exemplo_ficha.xlsx"),
        file_name=f"{nome}.pdf",
        mime="application/pdf",
        use_container_width=True,
        key=f"pdf_{chave}",
    )


# Função auxiliar para mapear as colunas corretas do Banco de Dados
def get_real_col(possible_names):
    colunas = st.session_state.get("historico_colunas", df_historico.columns)
    if len(colunas) == 0:
        return possible_names[0]
    for n in possible_names:
        if n in colunas:
            return n
    return possible_names[0]

MAX_DEDUCTIONS = 15
MAX_ADDITIONS = 16

PRESETS = {
    "motivo": [
        "5 Years Check",
        "IFE Installation",
        "New Painting",
        "New Production Aircraft",
        "Scheduled",
        "Seats Reconfiguration",
    ],
    "pesado_por": ["GOL", "Boeing", "DIGEX", "Flightstar", "Aeroman", "Transavia"],
    "local": [
        "Lagoa Santa, MG - Brazil",
        "Renton, WA - EUA",
        "Seattle, WA - EUA",
        "Jacksonville, FL - EUA",
        "São José dos Campos, SP - Brazil",
        "San Salvador - El Salvador",
        "Amsterdam - Netherlands",
        "Everett, WA - EUA",
    ],
    "config_lopa": [],
    "lopa": [],
    "deductions": [
        "Down Lock, Nose Gear",
        "Down Lock, Main Gear",
        "Plumb Bob",
        "In-seat Power Removal",
        "Seats Removal (177 Pax)",
        "Seats Removal (138 Pax)",
        "Tail Jack Adaptors",
        "Wing Jack Adaptors",
        "Carbon Brake Retrofit Program (SB 737-32-1429)",
        "Radial Tires Replacement (SB 737-32-1535)",
        "Flaps 0 - 40° (Up when weighed)",
        "Wiring Diagrams",
        "Production Flight Emergency Equipment",
        "Protective Carpet Runner",
        "Plastic Seat Cover",
        "Loto Kit",
        "Leather Seats Cover Installation",
        "Seats Removal (168 Pax)",
    ],
    "additions": [
        "2Ku System Antenna Installation (SB 737-44-1016)",
        "2Ku System Cabin Equip Installation (SB 737-44-1017)",
        "Carbon Brake Retrofit Program (SB 737-32-1429)",
        "Container, Oxygen Mask Spares",
        "Disable Person Kit Aft (1 EA)",
        "Dual navAero EFB System Installation",
        "Emergency Care Kit (1 EA)",
        "Emergency Medical Kit (1 EA)",
        "First Aid Kit Aft (2 EA)",
        "First Aid Kit Fwd (1 EA)",
        "Flaps 0 - 40° (Up when weighed)",
        "Fuel, Drain, Unusable (5.6 Gal)",
        "Fuel, Trapped, Unusable (14.2 Gal)",
        "Fuel, Trapped, Usable (5.5 Gal)",
        "Fwd Cargo, Aft Bulkhead Panels",
        "In-Seat Power Supply System (ISPSS) Installation",
        "Seats Installation (138 Pax)",
        "Seats Installation (186 Pax)",
        "Seats Installation (177 Pax)",
        "Lavatory Dry Supplies",
        "Lavatory Dry Supplies(2EA)",
        "Lavatory Pre-charge Fluid (6.0 Gal)",
        "Life Vest Installation",
        "Waste Container",
        "Nitrogen Generation System Install. (SB 737-47-1003)",
        "Oil, Auxiliary Power Unit (2.3 Gal)",
        "Oil, Drain, Unusable Engine (6.6 Gal)",
        "Oil, Drain, Usable Engine (10.2 Gal)",
        "Oil, Integrated Drive Gen (4.4 Gal)",
        "Oil, Trapped, Unusable Engine (1.0 Gal)",
        "Water, Potable (62.6 Gal)",
        "Water, Potable (42.2 Gal)",
        "Water, Potable (52.2 Gal)",
        "Oven Tray (8 EA)",
        "Oxigen Mask, Disposable (2 EA)",
        "Oxigen Mask, Disposable (3 EA)",
        "Oxygen Access Panel",
        "Oxygen Bottle, Portable (2 EA)",
        "Oxygen Cilinder Assy",
        "Oxygen Cilinder Charge (18 EA)",
        "Sat LINK Communication Management System",
        "Jungle Survival Kit Aft (1 EA)",
        "Jungle Survival Kit Aft (2 EA)",
        "Jungle Survival Kit Fwd (2 EA)",
        "Securelink/ MQAR System",
        "Standard Container",
        "Toillet Chemicals",
        "Operations Manual",
        "Oven",
        "Oven Rack",
        "Wheelchair (1 EA)",
    ],
}

@st.cache_data
def carregar_presets_lopa():
    df = carregar_dados_banco()
    valores = set()
    for coluna in (" LOPA", "LOPA", "Configuração LOPA ", "Configuração LOPA"):
        if coluna in df.columns:
            valores.update(
                safe_str(valor)
                for valor in df[coluna].dropna().tolist()
                if safe_str(valor).casefold() not in {"lopa", "configuração lopa"}
            )
    lopa = sorted(
        {v for v in valores if v.upper().startswith("GLP-")} | {"GLP-MAX8-001-XMC"},
        key=str.casefold,
    )
    config = sorted(
        {v for v in valores if not v.upper().startswith("GLP-")} | {"186 Pax + 10 Flight Crew"},
        key=str.casefold,
    )
    return lopa, config


PRESETS["lopa"], PRESETS["config_lopa"] = carregar_presets_lopa()


@st.cache_data
def carregar_opcoes_cadastradas():
    with conectar_banco() as conn:
        registros = conn.execute(
            'SELECT categoria, valor FROM opcoes_cadastradas ORDER BY valor'
        ).fetchall()
    opcoes = {}
    for registro in registros:
        opcoes.setdefault(registro['categoria'], []).append(registro['valor'])
    return opcoes


def incluir_opcoes_preset(categoria, valores):
    existentes = {opcao.casefold() for opcao in PRESETS.setdefault(categoria, [])}
    for valor in valores:
        if valor.casefold() not in existentes:
            PRESETS[categoria].append(valor)
            existentes.add(valor.casefold())


for _categoria, _valores in carregar_opcoes_cadastradas().items():
    incluir_opcoes_preset(_categoria, _valores)


def salvar_opcao_digitada(categoria, chave_widget):
    """Grava no banco uma opção digitada que ainda não existe na lista."""
    valor = safe_str(st.session_state.get(chave_widget))
    if not pode_emitir() or not valor or any(
        valor.casefold() == opcao.casefold() for opcao in PRESETS.get(categoria, [])
    ):
        return
    try:
        with conectar_banco() as conn:
            conn.execute(
                '''INSERT INTO opcoes_cadastradas
                   (categoria, valor, cadastrado_por, cadastrado_em)
                   VALUES (?, ?, ?, ?)
                   ON CONFLICT (categoria, valor) DO NOTHING''',
                (
                    categoria,
                    valor,
                    st.session_state.get('usuario_id'),
                    datetime.datetime.now(datetime.timezone.utc).isoformat(),
                ),
            )
    except Exception:
        st.toast(f"Não foi possível salvar “{valor}” na lista de opções.")
        return
    carregar_opcoes_cadastradas.clear()
    incluir_opcoes_preset(categoria, [valor])
    st.toast(f"“{valor}” foi adicionado à lista de opções.")


def campo_com_preset(container, label, valor, opcoes, key, categoria=None):
    valor = safe_str(valor)
    escolhas = [""] + list(opcoes)
    if valor and valor not in escolhas:
        escolhas.insert(1, valor)
    return container.selectbox(
        label,
        escolhas,
        index=escolhas.index(valor) if valor else 0,
        key=key,
        accept_new_options=True,
        placeholder="Digite ou selecione",
        on_change=salvar_opcao_digitada if categoria else None,
        args=(categoria, key) if categoria else None,
    )


def renderizar_editor_itens(
    container,
    itens,
    opcoes,
    key,
    momento_flaps=0,
    campos_com_erro=None,
    prefixo_campos_erro="",
    categoria=None,
):
    campos_com_erro = {safe_str(campo).casefold() for campo in campos_com_erro or ()}
    colunas = ["Descrição", "Peso (kgf)", "Braço (in)", "Momento (kgf.in)"]
    quantidade_inicial = max(len(itens), 1)
    chave_linhas = f"{key}_linhas"
    chave_proximo_id = f"{key}_proximo_id"
    chave_base = f"{key}_base"
    # Se a ficha de origem mudou (outra ficha, ou a mesma já salva com outros
    # itens), as linhas guardadas da tela anterior não valem mais: sem isso a
    # edição mostrava linhas vazias no lugar dos itens gravados.
    base_itens = repr([
        (safe_str(item.get("Descrição")), item.get("Peso (kgf)"), item.get("Braço (in)"))
        for item in itens
    ])
    if st.session_state.get(chave_base) != base_itens:
        for chave_widget in [
            chave for chave in st.session_state
            if isinstance(chave, str) and chave.startswith(f"{key}_")
        ]:
            del st.session_state[chave_widget]
        st.session_state[chave_base] = base_itens
    if chave_linhas not in st.session_state:
        quantidade_anterior = max(
            quantidade_inicial,
            st.session_state.get(f"{key}_quantidade", quantidade_inicial),
        )
        st.session_state[chave_linhas] = [str(indice) for indice in range(quantidade_anterior)]
        st.session_state[chave_proximo_id] = quantidade_anterior

    ids_linhas = list(st.session_state[chave_linhas])
    cabecalho = container.columns([4, 1, 1, 1, 0.9])
    for coluna, titulo in zip(cabecalho, [*colunas, "Ação"]):
        coluna.caption(titulo)

    linhas = []
    for posicao, id_linha in enumerate(ids_linhas):
        indice_item = int(id_linha) if id_linha.isdigit() else len(itens)
        item = itens[indice_item] if indice_item < len(itens) else {}
        descricao_inicial = safe_str(item.get("Descrição", ""))
        escolhas = [""] + list(opcoes)
        if descricao_inicial and descricao_inicial not in escolhas:
            escolhas.append(descricao_inicial)

        descricao_col, peso_col, braco_col, momento_col, excluir_col = container.columns([4, 1, 1, 1, 0.9])
        if prefixo_campos_erro and item:
            slot = item.get("slot", indice_item + 1)
            campos_item = {
                f"{prefixo_campos_erro} {campo} {slot}".casefold()
                for campo in ("description", "weigth", "arm")
            }
            if campos_com_erro & campos_item:
                descricao_col.markdown(":red[🔴 Este item foi devolvido para correção]")
        descricao = descricao_col.selectbox(
            f"Descrição, linha {posicao + 1}",
            escolhas,
            index=escolhas.index(descricao_inicial),
            key=f"{key}_{id_linha}_descricao",
            accept_new_options=True,
            placeholder="Digite ou selecione",
            label_visibility="collapsed",
            on_change=salvar_opcao_digitada if categoria else None,
            args=(categoria, f"{key}_{id_linha}_descricao") if categoria else None,
        )
        # O flap entra só pelo momento extra; peso e braço ficam travados em 0.
        eh_flap = "flap" in safe_str(descricao).casefold()
        sufixo_flap = "_flap" if eh_flap else ""
        peso = peso_col.number_input(
            f"Peso (kgf), linha {posicao + 1}",
            value=0.0 if eh_flap else safe_float(item.get("Peso (kgf)", 0.0)),
            step=0.1,
            key=f"{key}_{id_linha}_peso{sufixo_flap}",
            label_visibility="collapsed",
            disabled=eh_flap,
        )
        braco = braco_col.number_input(
            f"Braço (in), linha {posicao + 1}",
            value=0.0 if eh_flap else safe_float(item.get("Braço (in)", 0.0)),
            step=0.1,
            key=f"{key}_{id_linha}_braco{sufixo_flap}",
            label_visibility="collapsed",
            disabled=eh_flap,
            help="O flap entra só pelo momento; peso e braço ficam em 0." if eh_flap else None,
        )

        descricao = safe_str(descricao)
        momento = peso * braco
        if momento_flaps and "flap" in descricao.casefold():
            momento += momento_flaps
        momento_col.write(f"{momento:,.2f}")

        if descricao:
            linhas.append({
                "Descrição": descricao,
                "Peso (kgf)": peso,
                "Braço (in)": braco,
                "Momento (kgf.in)": momento,
            })

        excluir_col.button(
            "",
            key=f"{key}_{id_linha}_excluir",
            help="Excluir esta linha",
            on_click=excluir_linha,
            args=(chave_linhas, key, id_linha),
            type="tertiary",
            icon=":material/delete:",
        )

    adicionar_col, _ = container.columns(2)
    adicionar_col.button(
        "Adicionar item",
        key=f"{key}_adicionar",
        on_click=adicionar_linha,
        args=(chave_linhas, chave_proximo_id),
    )
    return pd.DataFrame(linhas, columns=colunas)


def adicionar_linha(chave_linhas, chave_proximo_id):
    novo_id = st.session_state[chave_proximo_id]
    st.session_state[chave_linhas] = [
        *st.session_state[chave_linhas], f"extra_{novo_id}"
    ]
    st.session_state[chave_proximo_id] = novo_id + 1


def excluir_linha(chave_linhas, prefixo_widget, id_linha):
    st.session_state[chave_linhas] = [
        linha for linha in st.session_state[chave_linhas] if linha != id_linha
    ]
    for campo in ("descricao", "peso", "braco", "peso_flap", "braco_flap"):
        st.session_state.pop(f"{prefixo_widget}_{id_linha}_{campo}", None)

def formulario_pesagem(
    prefixo_selecionado,
    p_sugerida,
    r_sugerida,
    p_anterior,
    r_anterior,
    linha_existente=None,
    campos_com_erro=None,
    modo="nova",
    checar_celulas=False,
    celulas_base=None,
):
    """Desenha a ficha em etapas e devolve (registro, pesagem, revisão,
    na_ultima_etapa, celulas_pendentes)."""
    if linha_existente is None: linha_existente = {}
    campos_com_erro = set(campos_com_erro or ())

    def rotulo_form(campo, rotulo, container=st):
        if campo in campos_com_erro:
            container.markdown(f":red[🔴 Corrigir: {rotulo}]")
        return rotulo

    info_aero = dict_tipos_aeronave.get(prefixo_selecionado, {})
    tipo_a = info_aero.get('modelo', "")
    # O modo separa os campos de cada tela: sem ele, a Nova ficha e o Editar
    # ficha da mesma revisão dividiam as linhas de Additions/Deductions.
    form_key = f"{modo}_{prefixo_selecionado}_{p_sugerida}_{r_sugerida}"

    # Etapas no lugar das abas: todas são desenhadas (os campos não perdem o
    # valor), mas só a etapa atual aparece. "Próxima" leva à seguinte.
    chave_passo = f"{form_key}_passo"
    st.session_state.setdefault(chave_passo, PASSOS_FICHA[0])
    st.segmented_control(
        "Etapa",
        PASSOS_FICHA,
        key=chave_passo,
        required=True,
        format_func=lambda passo: f"{PASSOS_FICHA.index(passo) + 1}. {passo}",
        label_visibility="collapsed",
        width="stretch",
    )
    passo_atual = st.session_state[chave_passo]
    indice_passo = PASSOS_FICHA.index(passo_atual)
    st.html(
        "<style>"
        + "".join(
            f".st-key-passo_ficha_{i} {{display: none;}}"
            for i in range(len(PASSOS_FICHA)) if i != indice_passo
        )
        + "</style>"
    )
    aba1, aba2, aba4, aba3, aba5, aba6 = [
        st.container(key=f"passo_ficha_{i}") for i in range(len(PASSOS_FICHA))
    ]

    with aba1:
        st.markdown("### Controle de Identificação")
        cp1, cp2 = st.columns(2)
        id_pesagem_input = cp1.number_input("Número da Pesagem:", value=int(p_sugerida), disabled=True)
        rev_input = cp2.number_input("Número da Revisão:", value=int(r_sugerida), disabled=True)
        # A última ficha do rodapé do PDF vem do sistema. Só quando ela não
        # está no banco a pessoa abre esta opção e digita os números.
        # Na Nova ficha a base é a ficha anterior: o valor digitado nela não vale aqui.
        mesma_ficha = (
            normalizar_chave_ficha(linha_existente.get('Pesagem')) == str(int(p_sugerida))
            and normalizar_chave_ficha(linha_existente.get('Revisao')) == str(int(r_sugerida))
        )
        up_salva, ur_salva = (
            ultima_ficha_gravada(linha_existente) if mesma_ficha else ("", "")
        )
        ultima_sistema = (
            f"Pesagem {int(p_anterior)}, Revisão {int(r_anterior)}"
            if p_anterior > 0 else "nenhuma"
        )
        with st.expander(
            f"Última ficha no sistema: {ultima_sistema}. "
            "Ela não está no sistema? Clique para informar.",
            expanded=bool(up_salva),
        ):
            cu1, cu2 = st.columns(2)
            pesagem_informada = cu1.text_input(
                "Última pesagem:",
                value=up_salva,
                key=f"{form_key}_ultima_pesagem",
                placeholder="Deixe em branco para usar a do sistema",
            ).strip()
            revisao_informada = cu2.text_input(
                "Última revisão:",
                value=ur_salva,
                key=f"{form_key}_ultima_revisao",
                placeholder="Deixe em branco para usar a do sistema",
            ).strip()
        if pesagem_informada:
            ultima_pesagem, ultima_revisao = pesagem_informada, revisao_informada or "0"
        elif p_anterior > 0:
            ultima_pesagem, ultima_revisao = str(int(p_anterior)), str(int(r_anterior))
        else:
            ultima_pesagem = ultima_revisao = ""
        st.divider()

        c1, c2, c3 = st.columns(3)
        chave_data_emissao = get_real_col(['Data da ficha', 'Data_da_Pesagem'])
        chave_pesado_por = get_real_col(['Pesado Por', 'WEIGHED BY'])
        chave_local = get_real_col(['Local da pesagem'])
        data_emissao = c1.date_input(
            rotulo_form(chave_data_emissao, "Data de emissão:", c1),
            value=datetime.date.today(),
        )
        pesado_por = campo_com_preset(c2, rotulo_form(chave_pesado_por, "Pesado por:", c2), linha_existente.get(chave_pesado_por, ''), PRESETS["pesado_por"], f"{form_key}_pesado_por", categoria="pesado_por")
        local = campo_com_preset(c3, rotulo_form(chave_local, "Local da pesagem:", c3), linha_existente.get(chave_local, ''), PRESETS["local"], f"{form_key}_local", categoria="local")
        
        c4, c5, c6 = st.columns(3)
        key_data_pes = get_real_col(['Data_da_Pesagem', 'Data da ficha'])
        data_pesagem_existente = pd.to_datetime(
            linha_existente.get(key_data_pes), errors="coerce", dayfirst=True
        )
        val_pesagem = (
            data_pesagem_existente.date()
            if pd.notna(data_pesagem_existente)
            else datetime.date.today()
        )
        data_pesagem = c4.date_input(
            rotulo_form(get_real_col(['Data_da_Pesagem']), "Data da pesagem:", c4),
            value=val_pesagem,
        )
        
        lopa_inicial, config_lopa_inicial = separar_campos_lopa(linha_existente)
        lopa = campo_com_preset(c5, rotulo_form("lopa", "LOPA:", c5), lopa_inicial, PRESETS["lopa"], f"{form_key}_lopa")
        config_lopa = campo_com_preset(c6, rotulo_form("config_lopa", "Configuração LOPA:", c6), config_lopa_inicial, PRESETS["config_lopa"], f"{form_key}_config_lopa")
        
        c7, c8, c9 = st.columns(3)
        chave_vrbl = get_real_col(['VRBL', 'VRBL NUMBER'])
        chave_serial = get_real_col(['SERIAL', 'SERIAL NUMBER'])
        chave_line = get_real_col(['LINE', 'LINE NUMBER'])
        vrbl = c7.text_input(rotulo_form(chave_vrbl, "VRBL. NUMBER:", c7), value=safe_str(linha_existente.get(chave_vrbl)) or info_aero.get('vrbl', ''))
        serial = c8.text_input(rotulo_form(chave_serial, "SERIAL NUMBER:", c8), value=safe_str(linha_existente.get(chave_serial)) or info_aero.get('serial', ''))
        line = c9.text_input(rotulo_form(chave_line, "LINE NUMBER:", c9), value=safe_str(linha_existente.get(chave_line)) or info_aero.get('line', ''))

        chave_motivo = get_real_col(['Motivo', 'Razão'])
        razao = campo_com_preset(st, rotulo_form(chave_motivo, "Razão para emissão:"), linha_existente.get(chave_motivo, ''), PRESETS["motivo"], f"{form_key}_motivo", categoria="motivo")

    with aba2:
        def campos_formulario(chaves, rotulos):
            def render(container, indice):
                return container.number_input(
                    rotulo_form(chaves[indice], rotulos[indice], container),
                    value=safe_float(linha_existente.get(chaves[indice], 0.0)),
                    step=10.0,
                    key=f"{form_key}_celula_{chaves[indice]}",
                )
            return grade_celulas(render)

        st.markdown("**Pesagem 01**")
        chave_p1_nlh = get_real_col(['Peso nariz LH'])
        chave_p1_nrh = get_real_col(['Peso Nariz RH'])
        chave_p1_rhm = get_real_col(['Peso MLG RH 1 ', 'Peso MLG RH 1'])
        chave_p1_lhm = get_real_col(['Peso MLG LH 1', 'Peso MLG LH 1 '])
        chave_p1_rhm2 = get_real_col(['Peso MLG RH 2', 'Peso MLG RH 2 '])
        chave_p1_lhm2 = get_real_col(['Peso MLG LH2', 'Peso MLG LH 2'])
        p1_nlh, p1_nrh, p1_rhm, p1_rhm2, p1_lhm, p1_lhm2 = campos_formulario(
            [chave_p1_nlh, chave_p1_nrh, chave_p1_rhm, chave_p1_rhm2, chave_p1_lhm, chave_p1_lhm2],
            ["Nariz LH", "Nariz RH", "RH MLG 1", "RH MLG 2", "LH MLG 1", "LH MLG 2"],
        )

        st.markdown("**Pesagem 02**")
        chave_p2_nlh = get_real_col(['Peso nariz LH pesagem 2'])
        chave_p2_nrh = get_real_col(['Peso Nariz RH pesagem 2'])
        chave_p2_rhm = get_real_col(['Peso MLG RH 1  pesagem 2', 'Peso MLG RH 1 pesagem 2'])
        chave_p2_lhm = get_real_col(['Peso MLG LH 1 pesagem 2'])
        chave_p2_rhm2 = get_real_col(['Peso MLG RH 2 pesagem 2'])
        chave_p2_lhm2 = get_real_col(['Peso MLG LH2 pesagem 2'])
        p2_nlh, p2_nrh, p2_rhm, p2_rhm2, p2_lhm, p2_lhm2 = campos_formulario(
            [chave_p2_nlh, chave_p2_nrh, chave_p2_rhm, chave_p2_rhm2, chave_p2_lhm, chave_p2_lhm2],
            ["Nariz LH P2", "Nariz RH P2", "RH MLG 1 P2", "RH MLG 2 P2", "LH MLG 1 P2", "LH MLG 2 P2"],
        )
        
        st.markdown("**Canelas**")
        unidade_canela_key = f"{form_key}_unidade_canela"
        unidade_canela_anterior_key = f"{form_key}_unidade_canela_anterior"
        chave_canela_lh = f"{form_key}_canela_lh"
        chave_canela_rh = f"{form_key}_canela_rh"
        st.session_state.setdefault(unidade_canela_key, "in")
        st.session_state.setdefault(
            unidade_canela_anterior_key, st.session_state[unidade_canela_key]
        )
        fator_unidade = 25.4 if st.session_state[unidade_canela_key] == "mm" else 1.0
        st.session_state.setdefault(
            chave_canela_lh,
            safe_float(linha_existente.get(get_real_col(['Canela MLG LH']), 0.0)) * fator_unidade,
        )
        st.session_state.setdefault(
            chave_canela_rh,
            safe_float(linha_existente.get(get_real_col(['Canela MLG RH']), 0.0)) * fator_unidade,
        )
        unidade_canela = st.radio(
            "Unidade da canela",
            ["in", "mm"],
            horizontal=True,
            key=unidade_canela_key,
            on_change=converter_unidade_canela,
            args=(
                unidade_canela_key,
                unidade_canela_anterior_key,
                (chave_canela_lh, chave_canela_rh),
            ),
        )
        can_c1, can_c2 = st.columns(2)
        canela_lh_input = can_c1.number_input(
            rotulo_form("Canela MLG LH", f"Canela LH ({unidade_canela})", can_c1),
            step=0.1 if unidade_canela == "in" else 1.0,
            key=chave_canela_lh,
        )
        canela_rh_input = can_c2.number_input(
            rotulo_form("Canela MLG RH", f"Canela RH ({unidade_canela})", can_c2),
            step=0.1 if unidade_canela == "in" else 1.0,
            key=chave_canela_rh,
        )
        canela_lh = valor_canela_em_polegadas(canela_lh_input, unidade_canela)
        canela_rh = valor_canela_em_polegadas(canela_rh_input, unidade_canela)
        tail_val = 0.0

    with aba3:
        ded_ex = []
        for i in range(1, MAX_DEDUCTIONS + 1):
            desc = linha_existente.get(f'Deductions description {i}')
            if pd.notna(desc) and str(desc).strip():
                peso = safe_float(linha_existente.get(f'Deductions Weigth {i}'))
                arm = safe_float(linha_existente.get(f'Deductions arm {i}'))
                ded_ex.append({"Descrição": desc, "Peso (kgf)": peso, "Braço (in)": arm, "Momento (kgf.in)": peso * arm, "slot": i})
                
        if not ded_ex:
            ded_ex = [{"Descrição": "Fuel (Usable)", "Peso (kgf)": 0.0, "Braço (in)": 660.5, "Momento (kgf.in)": 0.0}]

        deducoes_editadas = renderizar_editor_itens(
            st,
            ded_ex,
            PRESETS["deductions"],
            f"{form_key}_deductions",
            campos_com_erro=campos_com_erro,
            prefixo_campos_erro="Deductions",
            categoria="deductions",
        )

    with aba4:
        momento_extra_flaps = momento_flaps_do_registro(linha_existente)
        if momento_extra_flaps is None:
            momento_extra_flaps = buscar_momento_flaps_banco(
                prefixo_selecionado, tipo_a
            )

        add_ex = []
        for i in range(1, MAX_ADDITIONS + 1):
            desc = linha_existente.get(f'Additions Description {i}')
            if pd.notna(desc) and str(desc).strip():
                peso = safe_float(linha_existente.get(f'Additions weigth {i}'))
                arm = safe_float(linha_existente.get(f'Additions arm {i}'))
                m_calc = peso * arm
                if "flap" in str(desc).casefold(): m_calc += momento_extra_flaps
                add_ex.append({"Descrição": desc, "Peso (kgf)": peso, "Braço (in)": arm, "Momento (kgf.in)": m_calc, "slot": i})
                
        if not add_ex:
            add_ex = [{"Descrição": "Flaps 0 - 40° (up when weighed)", "Peso (kgf)": 0.0, "Braço (in)": 0.0, "Momento (kgf.in)": momento_extra_flaps}]

        adicoes_editadas = renderizar_editor_itens(
            st,
            add_ex,
            PRESETS["additions"],
            f"{form_key}_additions",
            momento_flaps=momento_extra_flaps,
            campos_com_erro=campos_com_erro,
            prefixo_campos_erro="Additions",
            categoria="additions",
        )

    with aba5:
        angulo_salvo = normalizar_angulo_level_correction(
            linha_existente.get(get_real_col(['Graus correção do cg']), '')
        )
        opcoes_angulo = [""] + list(LEVEL_CORRECTION_VALUES)
        angulo_level_correction = st.selectbox(
            rotulo_form(
                get_real_col(['Graus correção do cg']),
                "Ângulo de inclinação",
            ),
            opcoes_angulo,
            index=opcoes_angulo.index(angulo_salvo),
            key=f"{form_key}_level_correction",
        )
        fator_level_correction = LEVEL_CORRECTION_VALUES.get(angulo_level_correction, 0.0)
        st.metric("Valor correspondente", f"{fator_level_correction:.1f}")

    with aba6:
        lh_val = (p1_lhm + p1_lhm2 + p2_lhm + p2_lhm2) / 2
        rh_val = (p1_rhm + p1_rhm2 + p2_rhm + p2_rhm2) / 2
        nose_val = ((p1_nlh + p2_nlh) / 2) + ((p1_nrh + p2_nrh) / 2)
        
        arm_a = info_aero.get('armnose', 93.0000)
        arm_b_lh = (canela_lh * 0.1304) + 706.17
        arm_b_rh = (canela_rh * 0.1304) + 706.17
        arm_c = 1298.9000

        m_lh = lh_val * arm_b_lh
        m_rh = rh_val * arm_b_rh
        m_nose = nose_val * arm_a
        m_tail = tail_val * arm_c

        tot_reg_w = lh_val + rh_val + nose_val + tail_val
        tot_reg_m = m_lh + m_rh + m_nose + m_tail
        tot_reg_arm = tot_reg_m / tot_reg_w if tot_reg_w > 0 else 0.0

        ded_w = deducoes_editadas["Peso (kgf)"].sum() if not deducoes_editadas.empty else 0.0
        ded_m = (deducoes_editadas["Peso (kgf)"] * deducoes_editadas["Braço (in)"]).sum() if not deducoes_editadas.empty else 0.0
        level_correction_factor, level_correction_moment, level_correction_side = (
            calcular_level_correction(angulo_level_correction, tot_reg_w)
        )
        if level_correction_side == "deductions":
            ded_m += level_correction_moment
        ded_arm = ded_m / ded_w if ded_w > 0 else 0.0

        add_w = adicoes_editadas["Peso (kgf)"].sum() if not adicoes_editadas.empty else 0.0
        add_m = 0.0
        for idx, row in adicoes_editadas.iterrows():
            m = row["Peso (kgf)"] * row["Braço (in)"]
            if "flap" in str(row["Descrição"]).casefold(): m += momento_extra_flaps
            add_m += m
        if level_correction_side == "additions":
            add_m += level_correction_moment
        add_arm = add_m / add_w if add_w > 0 else 0.0

        basic_w = tot_reg_w + add_w - ded_w
        basic_m = tot_reg_m + add_m - ded_m
        basic_arm = basic_m / basic_w if basic_w > 0 else 0.0
        cg_mac_val = ((basic_arm - 627.1) / 1.558) if basic_arm > 0 else 0.0

        mostrar_relatorio_pesagem({
            "Reaction / Item": ["LH", "RH", "NOSE", "TAIL", "TOTAL REGISTERED", "LEVEL CORRECTION", "DEDUCTIONS", "ADDITIONS", "AIRCRAFT BASIC WEIGHT"],
            "Weight (kgf)": [lh_val, rh_val, nose_val, tail_val, tot_reg_w, 0.0, ded_w, add_w, basic_w],
            "Arm (inch)": [arm_b_lh, arm_b_rh, arm_a, arm_c, tot_reg_arm, level_correction_factor, ded_arm, add_arm, basic_arm],
            "Moment (kgf x inch)": [m_lh, m_rh, m_nose, m_tail, tot_reg_m, level_correction_moment, ded_m, add_m, basic_m]
        }, basic_w, basic_arm, cg_mac_val)
        
        deducoes_excel = [
            {'desc': d["Descrição"], 'w': d["Peso (kgf)"], 'a': d["Braço (in)"], 'm': d["Peso (kgf)"] * d["Braço (in)"]}
            for d in deducoes_editadas.to_dict('records')
        ]
        adicoes_excel = [
            {'desc': a["Descrição"], 'w': a["Peso (kgf)"], 'a': a["Braço (in)"], 'm': a["Peso (kgf)"] * a["Braço (in)"] + (momento_extra_flaps if "flap" in str(a["Descrição"]).casefold() else 0)}
            for a in adicoes_editadas.to_dict('records')
        ]
        if level_correction_side == "deductions":
            deducoes_excel.append({
                'desc': LEVEL_CORRECTION_LABEL, 'w': 0.0,
                'a': level_correction_factor, 'm': level_correction_moment,
            })
        elif level_correction_side == "additions":
            adicoes_excel.append({
                'desc': LEVEL_CORRECTION_LABEL, 'w': 0.0,
                'a': level_correction_factor, 'm': level_correction_moment,
            })

        dados_excel = {
            "prefixo": prefixo_selecionado, "modelo": tipo_a, "pesado_por": pesado_por,
            "local": local, "data": str(data_pesagem), "config_lopa": config_lopa, "lopa": lopa,
            "razao": razao, "cg_mac": cg_mac_val,
            "arm_a": arm_a, "arm_b_lh": arm_b_lh, "arm_b_rh": arm_b_rh,
            "issue_date": str(data_emissao),
            "ultima_pesagem": ultima_pesagem,
            "ultima_revisao": ultima_revisao,
            "vrbl_number": vrbl, "serial_number": serial, "line_number": line,
            "numero_pesagem": str(int(id_pesagem_input)), "revisao": str(int(rev_input)),
            "reactions": {
                'LH': {'weight': lh_val, 'arm': arm_b_lh, 'moment': m_lh},
                'RH': {'weight': rh_val, 'arm': arm_b_rh, 'moment': m_rh},
                'NOSE': {'weight': nose_val, 'arm': arm_a, 'moment': m_nose},
                'TAIL': {'weight': tail_val, 'arm': arm_c, 'moment': m_tail},
                'TOTAL REGISTERED': {'weight': tot_reg_w, 'arm': tot_reg_arm, 'moment': tot_reg_m},
                'DEDUCTIONS': {'weight': ded_w, 'arm': ded_arm, 'moment': ded_m},
                'ADDITIONS': {'weight': add_w, 'arm': add_arm, 'moment': add_m},
                'AIRCRAFT BASIC WEIGHT': {'weight': basic_w, 'arm': basic_arm, 'moment': basic_m}
            },
            "deductions": deducoes_excel,
            "additions": adicoes_excel
        }
        dados_excel["assinatura_emissor"] = carregar_assinatura_usuario(
            st.session_state['usuario_id']
        )
        dados_excel["assinatura_aprovador"] = None
        
        st.session_state["dados_excel_temp"] = dados_excel

    # Mapeamento 100% seguro para evitar o Database Error
    novo_registro = linha_existente.copy()
    
    novo_registro[get_real_col(['Prefixo'])] = prefixo_selecionado
    if tipo_a:
        novo_registro[get_real_col(['Tipo_aeronave'])] = tipo_a
    novo_registro[get_real_col(['Data da ficha', 'Data_da_Pesagem'])] = str(data_emissao)
    novo_registro[get_real_col(['Pesado Por', 'WEIGHED BY'])] = pesado_por
    novo_registro[get_real_col(['Local da pesagem'])] = local
    novo_registro[get_real_col(['Data_da_Pesagem'])] = str(data_pesagem)
    novo_registro[get_real_col([' LOPA', 'LOPA'])] = lopa
    novo_registro[get_real_col(['Configuração LOPA ', 'Configuração LOPA'])] = config_lopa
    novo_registro[get_real_col(['Motivo', 'Razão'])] = razao
    novo_registro[get_real_col(['VRBL', 'VRBL NUMBER'])] = vrbl
    novo_registro[get_real_col(['SERIAL', 'SERIAL NUMBER'])] = serial
    novo_registro[get_real_col(['LINE', 'LINE NUMBER'])] = line
    
    novo_registro[get_real_col(['Peso nariz LH'])] = p1_nlh
    novo_registro[get_real_col(['Peso Nariz RH'])] = p1_nrh
    novo_registro[get_real_col(['Peso MLG RH 1 ', 'Peso MLG RH 1'])] = p1_rhm
    novo_registro[get_real_col(['Peso MLG LH 1', 'Peso MLG LH 1 '])] = p1_lhm
    novo_registro[get_real_col(['Peso MLG RH 2', 'Peso MLG RH 2 '])] = p1_rhm2
    novo_registro[get_real_col(['Peso MLG LH2', 'Peso MLG LH 2'])] = p1_lhm2
    
    novo_registro[get_real_col(['Peso nariz LH pesagem 2'])] = p2_nlh
    novo_registro[get_real_col(['Peso Nariz RH pesagem 2'])] = p2_nrh
    novo_registro[get_real_col(['Peso MLG RH 1  pesagem 2', 'Peso MLG RH 1 pesagem 2'])] = p2_rhm
    novo_registro[get_real_col(['Peso MLG LH 1 pesagem 2'])] = p2_lhm
    novo_registro[get_real_col(['Peso MLG RH 2 pesagem 2'])] = p2_rhm2
    novo_registro[get_real_col(['Peso MLG LH2 pesagem 2'])] = p2_lhm2
    
    novo_registro[CAMPO_ULTIMA_PESAGEM_INFORMADA] = pesagem_informada or None
    novo_registro[CAMPO_ULTIMA_REVISAO_INFORMADA] = (
        (revisao_informada or "0") if pesagem_informada else None
    )
    novo_registro[get_real_col(['Canela MLG LH'])] = canela_lh
    novo_registro[get_real_col(['Canela MLG RH'])] = canela_rh
    novo_registro[get_real_col(['Tail'])] = tail_val
    novo_registro[get_real_col(['Adction LAP'])] = (
        momento_extra_flaps
        if any("flap" in str(descricao).casefold() for descricao in adicoes_editadas["Descrição"])
        else 0.0
    )
    novo_registro[get_real_col(['Graus correção do cg'])] = angulo_level_correction or None
    novo_registro[get_real_col(['Braço correspondente '])] = (
        LEVEL_CORRECTION_VALUES[angulo_level_correction]
        if angulo_level_correction else None
    )
    
    # Todos os itens digitados são gravados. O filtro antigo por colunas do
    # histórico descartava, por exemplo, o 6º item e os campos VRBL/SERIAL/LINE.
    for i in range(1, MAX_DEDUCTIONS + 1):
        novo_registro[f'Deductions description {i}'] = None
        novo_registro[f'Deductions Weigth {i}'] = None
        novo_registro[f'Deductions arm {i}'] = None
    for i, row in enumerate(deducoes_editadas.to_dict('records')[:MAX_DEDUCTIONS], start=1):
        novo_registro[f'Deductions description {i}'] = row.get('Descrição')
        novo_registro[f'Deductions Weigth {i}'] = row.get('Peso (kgf)')
        novo_registro[f'Deductions arm {i}'] = row.get('Braço (in)')

    for i in range(1, MAX_ADDITIONS + 1):
        novo_registro[f'Additions Description {i}'] = None
        novo_registro[f'Additions weigth {i}'] = None
        novo_registro[f'Additions arm {i}'] = None
    for i, row in enumerate(adicoes_editadas.to_dict('records')[:MAX_ADDITIONS], start=1):
        novo_registro[f'Additions Description {i}'] = row.get('Descrição')
        novo_registro[f'Additions weigth {i}'] = row.get('Peso (kgf)')
        novo_registro[f'Additions arm {i}'] = row.get('Braço (in)')

    # Células vazias ou iguais às da pesagem anterior são confirmadas na
    # própria etapa, no botão "Próxima"; a confirmação vale para esses valores.
    aviso_celulas = (
        aviso_celulas_ficha(novo_registro, celulas_base) if checar_celulas else None
    )
    chave_celulas_ok = f"{form_key}_celulas_ok"
    valores_celulas = tuple(
        safe_float(novo_registro.get(chave)) for chave in chaves_celulas_ficha()
    )
    celulas_pendentes = bool(aviso_celulas) and (
        st.session_state.get(chave_celulas_ok) != valores_celulas
    )

    st.divider()
    if passo_atual == "Células de Carga" and celulas_pendentes:
        st.warning(
            f"As Células de Carga estão {aviso_celulas}. Confira os valores "
            "antes de seguir: a Próxima tela confirma que estão certas."
        )
    elif indice_passo == len(PASSOS_FICHA) - 1 and celulas_pendentes:
        st.warning(
            f"As Células de Carga estão {aviso_celulas}. Volte à etapa "
            "Células de Carga e confirme antes de gerar a ficha."
        )
    # Na última etapa o "Voltar" fica na linha do botão de gerar (tela).
    st.session_state["_voltar_ficha"] = (
        chave_passo, PASSOS_FICHA[indice_passo - 1], f"{form_key}_voltar"
    )
    if indice_passo < len(PASSOS_FICHA) - 1:
        col_voltar, _, col_proxima = st.columns([1, 1, 1])
        if indice_passo > 0:
            botao_voltar_ficha(col_voltar)
        proximo = PASSOS_FICHA[indice_passo + 1]
        confirmar = passo_atual == "Células de Carga" and celulas_pendentes
        col_proxima.button(
            "Próxima tela",
            icon=":material/arrow_forward:",
            icon_position="right",
            type="primary",
            key=f"{form_key}_proxima",
            on_click=ir_para_passo,
            args=(chave_passo, proximo),
            kwargs=(
                {"confirmar": (chave_celulas_ok, valores_celulas)} if confirmar else {}
            ),
            use_container_width=True,
        )
    rolar_para_o_topo(passo_atual)

    return (
        novo_registro, id_pesagem_input, rev_input,
        indice_passo == len(PASSOS_FICHA) - 1, celulas_pendentes,
    )


PASSOS_FICHA = [
    "Dados da Ficha", "Células de Carga", "Additions", "Deductions",
    "Level Correction", "Preview",
]


def botao_voltar_ficha(container):
    chave_passo, anterior, chave = st.session_state["_voltar_ficha"]
    container.button(
        f"Voltar: {anterior}",
        icon=":material/arrow_back:",
        key=chave,
        on_click=ir_para_passo,
        args=(chave_passo, anterior),
        use_container_width=True,
    )


def ir_para_passo(chave_passo, passo, confirmar=None):
    if confirmar:
        st.session_state[confirmar[0]] = confirmar[1]
    st.session_state[chave_passo] = passo


def rolar_para_o_topo(passo):
    """Ao trocar de etapa, volta a tela para o topo (senão a pessoa fica no
    rodapé da etapa nova)."""
    if st.session_state.get("_ultimo_passo_ficha") == passo:
        return
    st.session_state["_ultimo_passo_ficha"] = passo
    components.html(
        "<script>"
        "const main = window.parent.document.querySelector("
        "'[data-testid=\"stMain\"], section.main');"
        "if (main) main.scrollTo({top: 0, behavior: 'smooth'});"
        "window.parent.scrollTo({top: 0, behavior: 'smooth'});"
        f"/* {passo} */</script>",
        height=0,
    )

OPCAO_NUMERACAO_MANUAL = "Digitar Pesagem e Revisão manualmente"


def chaves_celulas_ficha():
    return [
        get_real_col(nomes)
        for nomes in (
            ['Peso nariz LH'], ['Peso Nariz RH'],
            ['Peso MLG RH 1 ', 'Peso MLG RH 1'], ['Peso MLG LH 1', 'Peso MLG LH 1 '],
            ['Peso MLG RH 2', 'Peso MLG RH 2 '], ['Peso MLG LH2', 'Peso MLG LH 2'],
        )
    ]


def aviso_celulas_ficha(registro, ficha_base=None):
    """Diz se as Células de Carga estão vazias ou, numa nova pesagem, ainda
    iguais às copiadas da pesagem anterior. Devolve None quando estão em ordem.
    """
    chaves_celulas = chaves_celulas_ficha()
    celulas = [safe_float(registro.get(chave)) for chave in chaves_celulas]
    if sum(celulas) <= 0:
        return "vazias"
    if ficha_base and celulas == [
        safe_float(ficha_base.get(chave)) for chave in chaves_celulas
    ]:
        return "iguais às da pesagem anterior"
    return None


@st.fragment
@protegido
def tela_nova_ficha():
    st.session_state["historico_colunas"] = ()

    if not exigir_nivel(NIVEL_OPERACAO, NIVEL_ADMIN, NIVEL_ESPECIALISTA):
        return
    cabecalho("Nova ficha", "Escolha a aeronave e siga as etapas com o botão Próxima. Na última (Preview) você confere o resultado e gera a ficha.")
    aviso_somente_visualizacao()
    ficha_salva = st.session_state.get("nova_ficha_salva")
    if ficha_salva:
        # Evita que um segundo clique em "Salvar" grave outra revisão igual.
        st.success(
            f"Ficha {ficha_salva[0]}, Pesagem {ficha_salva[1]}, "
            f"Revisão {ficha_salva[2]} salva e enviada para aprovação."
        )
        if st.button("Emitir outra ficha", type="primary"):
            st.session_state.pop("nova_ficha_salva", None)
            st.rerun()
        return

    usuario_atual = st.session_state['usuario_id']
    minhas_devolvidas = listar_fichas_devolvidas_usuario(usuario_atual)
    opcoes_correcao = {
        (
            f"{ficha['chave'][0]} | Pesagem {ficha['chave'][1]} | "
            f"Revisão {ficha['chave'][2]}"
        ): ficha
        for ficha in minhas_devolvidas
    }
    correcao_selecionada = None
    if opcoes_correcao:
        st.info("Há fichas devolvidas para correção. Selecione uma para revisar e reenviar.")
        rotulo_correcao = st.selectbox(
            "Ficha pendente de correção",
            ["Nova ficha"] + list(opcoes_correcao),
            key="ficha_devolvida_emissao",
        )
        correcao_selecionada = opcoes_correcao.get(rotulo_correcao)

    if correcao_selecionada:
        prefixo = correcao_selecionada["chave"][0]
    else:
        chaves_validas = [
            str(k)
            for k in dict_tipos_aeronave.keys()
            if str(k) not in ('nan', 'None', '')
        ]
        prefixo = st.selectbox(
            "Aeronave",
            [""] + sorted(chaves_validas),
            key="nova_ficha_aeronave",
        )

    if prefixo:
        st.divider()
        chave_ficha_anterior = f"ficha_anterior_{prefixo}"
        ficha_anterior = (
            not correcao_selecionada
            and st.session_state.get(chave_ficha_anterior, False)
        )
        df_aero = carregar_fichas_prefixo(prefixo)
        st.session_state["historico_colunas"] = tuple(
            df_aero.columns if not df_aero.empty else df_historico.columns
        )
        if not df_aero.empty:
            df_aero['Pesagem_num'] = pd.to_numeric(
                df_aero['Pesagem'], errors='coerce'
            ).fillna(0)
            df_aero['Revisao_num'] = pd.to_numeric(
                df_aero['Revisao'], errors='coerce'
            ).fillna(0)

        if correcao_selecionada:
            ficha_correcao = df_aero.loc[
                (df_aero["Pesagem"].map(normalizar_chave_ficha)
                 == correcao_selecionada["chave"][1])
                & (df_aero["Revisao"].map(normalizar_chave_ficha)
                   == correcao_selecionada["chave"][2])
            ]
            if ficha_correcao.empty:
                st.error(
                    "A ficha devolvida não foi encontrada no histórico desta aeronave. "
                    "Atualize a lista e tente novamente."
                )
                return
            linha_base = ficha_correcao.iloc[0].to_dict()
            p_ant = int(correcao_selecionada["pesagem_num"])
            r_ant = int(correcao_selecionada["revisao_num"])
            p_nova = p_ant
            r_nova = r_ant
            campos_correcao = {
                campo["campo"] for campo in correcao_selecionada["erros"]
            }
            st.markdown("**Campos que precisam de correção:**")
            for erro in correcao_selecionada["erros"]:
                st.markdown(f":red[🔴 {erro['descricao']}]")
        elif not df_aero.empty and ficha_anterior:
            df_aero = df_aero.sort_values(by=['Pesagem_num', 'Revisao_num'])
            st.info(
                "📥 **Cadastrar ficha anterior.** Use para incluir uma ficha "
                "antiga que não está no sistema. Digite a Pesagem e a Revisão "
                "dela e preencha as abas com os dados da ficha em papel. Ela "
                "não muda a última ficha da aeronave."
            )
            if st.button("Voltar para nova ficha", key=f"cancelar_anterior_{prefixo}"):
                st.session_state.pop(chave_ficha_anterior, None)
                st.rerun()
            u_pes = int(df_aero.iloc[-1]['Pesagem_num'])
            u_rev = int(df_aero.iloc[-1]['Revisao_num'])
            tipo_acao = OPCAO_NUMERACAO_MANUAL
            linha_base = {}
            p_nova, r_nova = (u_pes, u_rev - 1) if u_rev > 0 else (max(u_pes - 1, 1), 0)
            p_ant = r_ant = 0
            campos_correcao = set()
        elif not df_aero.empty:
            df_aero = df_aero.sort_values(by=['Pesagem_num', 'Revisao_num'])

            ultima_linha = df_aero.iloc[-1]
            linha_base = ultima_linha.to_dict()
            linha_base.pop('Pesagem_num', None)
            linha_base.pop('Revisao_num', None)
            
            u_pes = int(ultima_linha['Pesagem_num'])
            u_rev = int(ultima_linha['Revisao_num'])
            
            st.info(f"📋 **Base Encontrada:** A última ficha dessa aeronave foi a **Pesagem {u_pes} | Revisão {u_rev}**. Os dados foram copiados para agilizar o preenchimento.")
            tipo_acao = st.radio(
                "Escolha a ação desejada:",
                [
                    f"Nova Revisão (Pesagem {u_pes}, Revisão {u_rev + 1})",
                    f"Nova Pesagem (Pesagem {u_pes + 1}, Revisão 0)",
                    OPCAO_NUMERACAO_MANUAL,
                ],
                horizontal=True,
                key=f"tipo_acao_{prefixo}",
            )
            if st.button(
                "Cadastrar ficha anterior que falta no sistema",
                icon=":material/history:",
                type="tertiary",
                key=f"abrir_anterior_{prefixo}",
                help="Inclui uma ficha antiga (por exemplo, a que veio antes "
                "da última) sem copiar os dados da última ficha.",
            ):
                st.session_state[chave_ficha_anterior] = True
                st.rerun()

            if tipo_acao.startswith("Nova Pesagem"):
                p_nova = u_pes + 1
                r_nova = 0
            else:
                p_nova = u_pes
                r_nova = u_rev + 1
                
            p_ant = u_pes
            r_ant = u_rev
            campos_correcao = set()
        else:
            st.info("🆕 **Nenhuma ficha anterior encontrada.** Iniciando a primeira Pesagem (1), Revisão (0) com a ficha em branco.")
            tipo_acao = st.radio(
                "Escolha a ação desejada:",
                ["Primeira ficha (Pesagem 1, Revisão 0)", OPCAO_NUMERACAO_MANUAL],
                horizontal=True,
                key=f"tipo_acao_{prefixo}",
            )
            linha_base = {}
            p_nova = 1
            r_nova = 0
            p_ant = 0
            r_ant = 0
            campos_correcao = set()

        if not correcao_selecionada and p_ant > 0:
            campos_pendentes = carregar_campos_com_erro(prefixo, p_ant, r_ant)
            if campos_pendentes:
                emissor = buscar_fluxo_ficha(prefixo, p_ant, r_ant)['gerador_nome']
                st.warning(
                    f"A revisão anterior foi devolvida ao emissor ({emissor}). "
                    "Campos a corrigir: "
                    + "; ".join(campo['descricao'] for campo in campos_pendentes)
                )

        ficha_ja_existe = False
        if not correcao_selecionada and tipo_acao == OPCAO_NUMERACAO_MANUAL:
            col_pesagem, col_revisao = st.columns(2)
            p_nova = col_pesagem.number_input(
                "Número da Pesagem",
                min_value=1,
                value=max(1, int(p_nova)),
                step=1,
                key=f"pesagem_manual_{prefixo}_{int(ficha_anterior)}",
            )
            r_nova = col_revisao.number_input(
                "Número da Revisão",
                min_value=0,
                value=max(0, int(r_nova)),
                step=1,
                key=f"revisao_manual_{prefixo}_{int(ficha_anterior)}",
            )
            # Digitar um número que já existe gravaria por cima de uma ficha
            # emitida; o aviso aparece antes de a pessoa preencher as abas.
            ficha_ja_existe = not df_aero.empty and (
                (df_aero['Pesagem_num'] == int(p_nova))
                & (df_aero['Revisao_num'] == int(r_nova))
            ).any()
            if ficha_ja_existe:
                st.error(
                    f"Já existe uma ficha para {prefixo}, Pesagem {int(p_nova)}, "
                    f"Revisão {int(r_nova)}. Escolha outro número. Para alterar "
                    "essa ficha, use a tela Editar ficha."
                )
            
        (
            novo_registro, p_final, r_final, na_ultima_etapa, celulas_pendentes,
        ) = formulario_pesagem(
            prefixo,
            p_nova,
            r_nova,
            p_ant,
            r_ant,
            linha_existente=linha_base,
            campos_com_erro=campos_correcao,
            modo="anterior" if ficha_anterior else "nova",
            checar_celulas=not correcao_selecionada,
            celulas_base=linha_base if int(p_nova) != int(p_ant) else None,
        )
        # Gerar e exportar ficam na última etapa (Preview), depois de a pessoa
        # passar por todas.
        if not na_ultima_etapa:
            return

        col_voltar, col_btn1, col_btn2 = st.columns([1, 1.3, 1.7])
        botao_voltar_ficha(col_voltar)
        with col_btn1:
            label_salvar = (
                "Emitir ficha e enviar para aprovação"
                if correcao_selecionada
                else "Cadastrar ficha anterior"
                if ficha_anterior
                else "Gerar ficha e enviar para aprovação"
            )
            if st.button(
                label_salvar,
                type="primary",
                use_container_width=True,
                disabled=(
                    bool(ficha_ja_existe) or celulas_pendentes
                    or not pode_emitir()
                ),
            ):
                ficha_existente = False
                if not df_aero.empty:
                    ficha_existente = (
                        (df_aero['Pesagem_num'] == int(p_final))
                        & (df_aero['Revisao_num'] == int(r_final))
                    ).any()

                novo_registro['Pesagem'] = str(int(p_final))
                novo_registro['Revisao'] = str(int(r_final))
                if correcao_selecionada:
                    alteracoes = atualizar_ficha(prefixo, p_final, r_final, novo_registro)
                    if alteracoes is None:
                        st.error(
                            "Não foi possível localizar a ficha original para "
                            "salvar a correção. Atualize a página e tente novamente."
                        )
                    else:
                        registrar_ficha(
                            prefixo,
                            p_final,
                            r_final,
                            st.session_state['usuario_id'],
                            acao="Corrigida",
                            alteracoes=alteracoes,
                        )
                        limpar_campos_com_erro(prefixo, p_ant, r_ant)
                        st.session_state["nova_ficha_salva"] = (
                            prefixo, int(p_final), int(r_final)
                        )
                        st.rerun()
                elif ficha_existente:
                    st.error(f"Já existe uma ficha para {prefixo}, Pesagem {int(p_final)}, Revisão {int(r_final)}.")
                elif not inserir_ficha(novo_registro):
                    st.error(f"Já existe uma ficha para {prefixo}, Pesagem {int(p_final)}, Revisão {int(r_final)}.")
                else:
                    registrar_ficha(
                        prefixo, p_final, r_final, st.session_state['usuario_id'],
                        acao="Ficha anterior cadastrada" if ficha_anterior else "Gerada",
                    )
                    st.session_state.pop(chave_ficha_anterior, None)
                    st.session_state["nova_ficha_salva"] = (
                        prefixo, int(p_final), int(r_final)
                    )
                    st.rerun()
        with col_btn2:
            botoes_exportacao(
                st, lambda dados=st.session_state["dados_excel_temp"]: dados, prefixo,
                p_final, r_final, f"nova_{prefixo}",
            )

@st.fragment
@protegido
def tela_edicao():
    st.session_state["historico_colunas"] = ()
    cabecalho("Editar ficha", "Altere uma revisão já emitida. Ela volta para aprovação.")
    if not exigir_nivel(NIVEL_OPERACAO, NIVEL_ADMIN, NIVEL_ESPECIALISTA):
        return
    aviso_somente_visualizacao()

    prefixos = [""] + carregar_prefixos_fichas()
    prefixo = st.selectbox("Aeronave", prefixos)
    
    if prefixo:
        df_filtrado = carregar_fichas_prefixo(prefixo)
        st.session_state["historico_colunas"] = tuple(df_filtrado.columns)
        pesagem = st.selectbox("Pesagem a editar", [""] + sorted(df_filtrado['Pesagem'].dropna().unique().tolist()))
        
        if pesagem:
            revisao = st.selectbox("Revisão a editar", [""] + sorted(df_filtrado[df_filtrado['Pesagem'] == pesagem]['Revisao'].dropna().unique().tolist()))
            
            if revisao:
                st.divider()
                linha_atual = df_filtrado[(df_filtrado['Pesagem'] == pesagem) & (df_filtrado['Revisao'] == revisao)].iloc[0].to_dict()
                linha_atual.pop('Pesagem_num', None)
                linha_atual.pop('Revisao_num', None)

                u_pes = int(safe_float(linha_atual.get('Pesagem', 1)))
                u_rev = int(safe_float(linha_atual.get('Revisao', 0)))
                
                p_nova = u_pes
                r_nova = u_rev
                # A anterior é a ficha logo antes desta no sistema (não ela mesma).
                anteriores = sorted(
                    (safe_float(p), safe_float(r))
                    for p, r in zip(df_filtrado['Pesagem'], df_filtrado['Revisao'])
                    if (safe_float(p), safe_float(r)) < (u_pes, u_rev)
                )
                p_ant, r_ant = (
                    (int(anteriores[-1][0]), int(anteriores[-1][1]))
                    if anteriores else (0, 0)
                )
                
                novo_registro, p_final, r_final, na_ultima_etapa, _ = formulario_pesagem(
                    prefixo, p_nova, r_nova, p_ant, r_ant,
                    linha_existente=linha_atual, modo="edicao",
                )
                if not na_ultima_etapa:
                    return

                col_voltar, col_btn1, col_btn2 = st.columns([1, 1.3, 1.7])
                botao_voltar_ficha(col_voltar)
                with col_btn1:
                    if st.button(
                        "Enviar alterações para aprovação", type="primary",
                        use_container_width=True, disabled=not pode_emitir(),
                    ):
                        fluxo_origem = buscar_fluxo_ficha(prefixo, pesagem, revisao)
                        novo_registro['Pesagem'] = str(int(p_final))
                        novo_registro['Revisao'] = str(int(r_final))
                        alteracoes = atualizar_ficha(
                            prefixo, pesagem, revisao, novo_registro
                        )
                        if alteracoes is None:
                            st.error(
                                "Não foi possível localizar a ficha selecionada "
                                "para salvar as alterações. Atualize a página e tente novamente."
                            )
                            return
                        registrar_ficha(
                            prefixo, p_final, r_final, st.session_state['usuario_id'],
                            acao="Alterada",
                            alteracoes=alteracoes,
                        )
                        if fluxo_origem['gerador_usuario'] == st.session_state['usuario_id']:
                            limpar_campos_com_erro(prefixo, pesagem, revisao)
                        st.success(
                            "Alterações salvas na mesma revisão e enviadas para aprovação."
                        )
                with col_btn2:
                    botoes_exportacao(
                        st, lambda dados=st.session_state["dados_excel_temp"]: dados, prefixo,
                        p_final, r_final, f"edicao_{prefixo}",
                    )

@protegido
def tela_criar_login():
    if not exigir_nivel(NIVEL_ADMIN):
        return

    cabecalho("Usuários", "Crie acessos para a equipe.")
    with st.form("form_criar_login"):
        nome = st.text_input("Nome completo")
        usuario = st.text_input("Usuário")
        senha = st.text_input("Senha", type="password")
        confirmar_senha = st.text_input("Confirmar senha", type="password")
        nivel = st.selectbox(
            "Perfil", list(PERFIS), format_func=PERFIS.get,
            help="Consulta só visualiza. Gera e aprova emite, edita e aprova fichas. "
                 "Especialista GOL vê as telas de emissão e aprovação, mas não gera nem aprova. "
                 "Administrador também gerencia usuários e assinaturas.",
        )
        enviar = st.form_submit_button("Criar usuário", type="primary")

    if enviar:
        if not nome.strip() or not usuario.strip():
            st.error("Informe o nome e o usuário.")
        elif len(senha) < 1:
            st.error("A senha deve ter pelo menos 1 caracteres.")
        elif senha != confirmar_senha:
            st.error("As senhas não coincidem.")
        else:
            try:
                criar_usuario(nome, usuario, senha, nivel)
                st.success(f"Login criado para {nome.strip()}.")
            except IntegrityError:
                st.error("Esse nome de usuário já está cadastrado.")

    st.divider()
    st.subheader("Alterar perfil")
    usuarios = listar_usuarios_assinaturas()
    if usuarios:
        opcoes = {u['usuario']: f"{u['nome']} ({u['usuario']})" for u in usuarios}
        perfis_atuais = {u['usuario']: u['nivel_acesso'] for u in usuarios}
        c1, c2, c3 = st.columns([2, 1.3, 1])
        escolhido = c1.selectbox("Usuário", list(opcoes), format_func=opcoes.get, key="perfil_usuario")
        novo_nivel = c2.selectbox(
            "Novo perfil", list(PERFIS), format_func=PERFIS.get,
            index=list(PERFIS).index(perfis_atuais.get(escolhido, NIVEL_CONSULTA))
            if perfis_atuais.get(escolhido) in PERFIS else 0,
            key=f"perfil_novo_{escolhido}",
        )
        c3.write("")
        c3.write("")
        if c3.button("Salvar perfil", use_container_width=True):
            if escolhido == st.session_state['usuario_id'] and novo_nivel != NIVEL_ADMIN:
                st.error("Você não pode remover o seu próprio acesso de Administrador.")
            else:
                with conectar_banco() as conn:
                    conn.execute(
                        'UPDATE usuarios SET nivel_acesso = ? WHERE usuario = ?',
                        (novo_nivel, escolhido),
                    )
                st.success(f"Perfil de {opcoes[escolhido]} alterado para {PERFIS[novo_nivel]}.")

        st.divider()
        st.subheader("Redefinir senha")
        st.caption("As senhas ficam criptografadas e não podem ser vistas; defina uma nova e informe à pessoa.")
        with st.form("form_redefinir_senha", clear_on_submit=True):
            alvo = st.selectbox("Usuário", list(opcoes), format_func=opcoes.get)
            c1, c2 = st.columns(2)
            nova = c1.text_input("Nova senha", type="password")
            confirmar = c2.text_input("Confirmar nova senha", type="password")
            trocar = st.form_submit_button("Redefinir senha", type="primary")
        if trocar:
            if not nova:
                st.error("Informe a nova senha.")
            elif nova != confirmar:
                st.error("As senhas não coincidem.")
            else:
                redefinir_senha(alvo, nova)
                st.success(f"Senha de {opcoes[alvo]} redefinida.")


@protegido
def tela_cadastrar_aeronave():
    if not exigir_nivel(NIVEL_OPERACAO, NIVEL_ADMIN):
        return

    cabecalho("Aeronaves", "Cadastre aeronaves que ainda não estão na frota do sistema.")
    st.caption(
        "A aeronave cadastrada aparece na lista de Nova Ficha. O modelo define "
        "o momento do flap sugerido na primeira ficha, a partir das fichas de "
        "outras aeronaves do mesmo modelo."
    )

    modelos = sorted({
        dados['modelo'] for dados in dict_tipos_aeronave.values()
        if dados.get('modelo')
    })
    modelo = st.selectbox(
        "Modelo da aeronave",
        [""] + modelos,
        key="cadastro_aeronave_modelo",
    )
    armnose_sugerido = armnose_padrao_modelo(modelo) if modelo else 93.0

    with st.form("form_cadastrar_aeronave", clear_on_submit=True):
        c1, c2 = st.columns(2)
        prefixo = c1.text_input("Prefixo", placeholder="PR-XXX")
        serial = c2.text_input("Serial number (MSN)")
        c3, c4, c5 = st.columns(3)
        vrbl = c3.text_input("VRBL number")
        line = c4.text_input("Line number")
        armnose = c5.number_input(
            "Arm nose (in)",
            value=armnose_sugerido,
            format="%.4f",
            help="Sugerido a partir das outras aeronaves do mesmo modelo.",
            key=f"cadastro_aeronave_armnose_{modelo}",
        )
        enviar = st.form_submit_button("Cadastrar aeronave", type="primary")

    if enviar:
        prefixo = normalizar_prefixo(prefixo)
        if not modelo:
            st.error("Selecione o modelo da aeronave.")
        elif not prefixo:
            st.error("Informe o prefixo.")
        elif any(normalizar_prefixo(k) == prefixo for k in dict_tipos_aeronave):
            st.error(f"A aeronave {prefixo} já está cadastrada.")
        else:
            try:
                cadastrar_aeronave(
                    prefixo,
                    modelo,
                    safe_identifier(serial),
                    safe_identifier(vrbl),
                    safe_identifier(line),
                    armnose,
                    st.session_state['usuario_id'],
                )
            except IntegrityError:
                st.error(f"A aeronave {prefixo} já está cadastrada.")
            else:
                st.session_state['aeronave_cadastrada'] = prefixo
                st.rerun()

    if st.session_state.get('aeronave_cadastrada'):
        st.success(
            f"Aeronave {st.session_state.pop('aeronave_cadastrada')} cadastrada."
        )

    cadastradas = carregar_aeronaves_cadastradas()
    st.divider()
    st.subheader("Aeronaves cadastradas pelo app")
    if not cadastradas:
        st.info("Nenhuma aeronave cadastrada pelo app ainda.")
        return

    st.dataframe(
        pd.DataFrame([
            {
                "Prefixo": prefixo,
                "Modelo": dados['modelo'],
                "Serial": dados['serial'],
                "VRBL": dados['vrbl'],
                "Line": dados['line'],
                "Arm nose": dados['armnose'],
                "Cadastrado por": dados['cadastrado_por'],
            }
            for prefixo, dados in cadastradas.items()
        ]),
        hide_index=True,
        use_container_width=True,
    )

    prefixos_com_ficha = set(carregar_prefixos_fichas())
    removiveis = [p for p in cadastradas if p not in prefixos_com_ficha]
    if removiveis:
        with st.expander("Remover aeronave cadastrada por engano"):
            st.caption("Só é possível remover aeronaves que ainda não têm ficha.")
            prefixo_remover = st.selectbox("Aeronave", removiveis)
            if st.button("Remover aeronave"):
                remover_aeronave_cadastrada(prefixo_remover)
                st.rerun()


@protegido
def tela_assinaturas():
    if not exigir_nivel(NIVEL_ADMIN):
        return

    cabecalho("Assinaturas", "Assinaturas usadas no Excel da ficha.")
    usuarios = listar_usuarios_assinaturas()
    if not usuarios:
        st.info("Não há usuários cadastrados.")
        return

    st.dataframe(
        pd.DataFrame([
            {
                "Nome": usuario['nome'],
                "Usuário": usuario['usuario'],
                "Perfil": PERFIS.get(usuario['nivel_acesso'], "—"),
                "Assinatura": "Cadastrada" if usuario['atualizado_em'] else "Pendente",
            }
            for usuario in usuarios
        ]),
        use_container_width=True,
        hide_index=True,
    )

    opcoes = {
        f"{usuario['nome']} ({usuario['usuario']})": usuario['usuario']
        for usuario in usuarios
    }
    usuario_selecionado = st.selectbox(
        "Selecione o usuário para cadastrar ou atualizar a assinatura",
        list(opcoes),
        key="usuario_assinatura",
    )
    usuario = opcoes[usuario_selecionado]
    assinatura_atual = carregar_assinatura_usuario(usuario)
    if assinatura_atual:
        st.image(assinatura_atual, caption="Assinatura cadastrada", width=300)

    arquivo = st.file_uploader(
        "Enviar assinatura (PNG ou JPEG, até 5 MB)",
        type=["png", "jpg", "jpeg"],
        key=f"arquivo_assinatura_{usuario}",
    )
    coluna_salvar, coluna_remover = st.columns(2)
    if coluna_salvar.button(
        "Salvar assinatura",
        type="primary",
        disabled=arquivo is None,
        key=f"salvar_assinatura_{usuario}",
        use_container_width=True,
    ):
        try:
            imagem_png = normalizar_imagem_assinatura(arquivo.getvalue())
        except ValueError as erro:
            st.error(str(erro))
        else:
            salvar_assinatura_usuario(usuario, imagem_png)
            st.success(f"Assinatura de {usuario_selecionado} salva.")
            st.rerun()

    if assinatura_atual and coluna_remover.button(
        "Remover assinatura",
        key=f"remover_assinatura_{usuario}",
        use_container_width=True,
    ):
        remover_assinatura_usuario(usuario)
        st.success(f"Assinatura de {usuario_selecionado} removida.")
        st.rerun()


@st.fragment
@protegido
def tela_aprovar_fichas():
    if not exigir_nivel(NIVEL_OPERACAO, NIVEL_ADMIN, NIVEL_ESPECIALISTA):
        return

    cabecalho("Aprovação", "Revise, devolva para correção ou aprove as fichas emitidas.")
    aviso_somente_visualizacao()
    pendentes_aprovacao, pendentes_correcao = listar_fichas_pendentes()
    aba_aprovacao, aba_correcao = st.tabs(
        ["Pendentes de aprovação", "Pendentes de correção"]
    )

    with aba_aprovacao:
        recem_aprovada = st.session_state.get("ficha_recem_aprovada")
        linha_aprovada = carregar_linha_ficha(*recem_aprovada) if recem_aprovada else None
        if linha_aprovada is not None:
            prefixo, pesagem, revisao = recem_aprovada
            st.success(
                f"Ficha já aprovada: {prefixo} | Pesagem {pesagem} | Revisão {revisao}. "
                "Baixe o Excel ou o PDF abaixo."
            )
            if st.button("Voltar às fichas pendentes", key="voltar_pendentes"):
                st.session_state.pop("ficha_recem_aprovada", None)
                st.rerun()
            renderizar_ficha_visualizacao(prefixo, pesagem, revisao, linha_aprovada)
        elif not pendentes_aprovacao:
            st.success("Não há fichas pendentes de aprovação.")
        else:
            opcoes_aprovacao = {
                (
                    f"{ficha['chave'][0]} | Pesagem {ficha['chave'][1]} | "
                    f"Revisão {ficha['chave'][2]}"
                ): ficha
                for ficha in pendentes_aprovacao
            }
            selecionada = st.selectbox(
                "Selecione a ficha para revisar",
                list(opcoes_aprovacao),
                key="ficha_pendente_aprovacao",
            )
            ficha = opcoes_aprovacao[selecionada]
            prefixo, pesagem, revisao = ficha["chave"]
            linha = carregar_linha_ficha(prefixo, pesagem, revisao)
            if linha is None:
                st.error("A ficha selecionada não foi encontrada no banco de dados.")
            else:
                renderizar_ficha_visualizacao(
                    prefixo,
                    pesagem,
                    revisao,
                    linha,
                    modo_aprovacao=True,
                )

    with aba_correcao:
        if not pendentes_correcao:
            st.success("Não há fichas aguardando correção do emissor.")
        else:
            opcoes_correcao = {
                (
                    f"{ficha['chave'][0]} | Pesagem {ficha['chave'][1]} | "
                    f"Revisão {ficha['chave'][2]}"
                ): ficha
                for ficha in pendentes_correcao
            }
            selecionada = st.selectbox(
                "Selecione a ficha devolvida",
                list(opcoes_correcao),
                key="ficha_pendente_correcao",
            )
            ficha = opcoes_correcao[selecionada]
            prefixo, pesagem, revisao = ficha["chave"]
            st.warning(
                f"Devolvida para correção por {ficha['fluxo']['gerador_nome']}. "
                "O emissor pode abrir “Nova Ficha”, selecionar esta ficha e "
                "reenviá-la para aprovação."
            )
            st.markdown("**Campos que precisam de correção:**")
            for erro in ficha["erros"]:
                st.markdown(f":red[🔴 {erro['descricao']}]")
            linha = carregar_linha_ficha(prefixo, pesagem, revisao)
            if linha is None:
                st.error("A ficha selecionada não foi encontrada no banco de dados.")
            else:
                renderizar_ficha_visualizacao(
                    prefixo,
                    pesagem,
                    revisao,
                    linha,
                    destacar_erros=True,
                )

GA_FALTANDO = "Faltando quem gerou ou aprovou"
GA_TODAS = "Todas as fichas"
# Rótulo da célula vazia na tabela (o Streamlit mostraria "None").
GA_VAZIO = "—"


def ga_login(valor):
    return valor if isinstance(valor, str) and valor else None


def ga_novo_login(novos, coluna, antes, usuario_por_rotulo):
    if coluna not in novos:
        return antes
    rotulo = novos[coluna]
    if not isinstance(rotulo, str) or rotulo in ("", GA_VAZIO):
        return None
    return usuario_por_rotulo.get(rotulo, antes)


def ga_guardar_edicoes():
    # Passa o que foi digitado na tabela para a lista de alterações pendentes,
    # para não se perder quando o filtro muda.
    estado = st.session_state.get(f"ga_editor_{st.session_state.get('ga_versao', 0)}")
    exibidas = st.session_state.get("ga_chaves_exibidas", [])
    pendentes = st.session_state.setdefault("ga_pendentes", {})
    if estado:
        for posicao, valores in estado.get("edited_rows", {}).items():
            if int(posicao) < len(exibidas):
                pendentes.setdefault(exibidas[int(posicao)], {}).update(valores)
    st.session_state["ga_versao"] = st.session_state.get("ga_versao", 0) + 1


def ga_descartar():
    st.session_state["ga_pendentes"] = {}
    st.session_state["ga_versao"] = st.session_state.get("ga_versao", 0) + 1


@protegido
def tela_gerou_aprovou():
    if not pode_preencher_gerou_aprovou():
        st.error("Esta tela é só do usuário Engenharia.")
        return

    cabecalho(
        "Gerou e aprovou",
        "Informe quem gerou e quem aprovou as fichas que vieram do banco "
        "antigo sem esse registro.",
    )
    mensagem = st.session_state.pop("ga_mensagem", None)
    if mensagem:
        st.success(mensagem)

    df = listar_fichas_gerou_aprovou()
    usuarios = listar_usuarios_assinaturas()

    # Rótulo de cada usuário na lista: o nome; se dois tiverem o mesmo nome,
    # o login vai junto.
    nomes = [safe_str(u["nome"]) or u["usuario"] for u in usuarios]
    rotulo_por_usuario = {}
    for usuario, nome in zip((u["usuario"] for u in usuarios), nomes):
        rotulo_por_usuario[usuario] = (
            f"{nome} ({usuario})" if nomes.count(nome) > 1 else nome
        )
    # Logins gravados que não existem mais em Usuários continuam aparecendo.
    for valor in pd.concat([df["gerado_por"], df["aprovado_por"]]).dropna().unique():
        rotulo_por_usuario.setdefault(valor, f"{valor} (sem cadastro)")
    usuario_por_rotulo = {rotulo: usuario for usuario, rotulo in rotulo_por_usuario.items()}
    opcoes = sorted(usuario_por_rotulo, key=str.casefold)

    sem_gerou = int(df["gerado_por"].isna().sum())
    sem_aprovou = int(df["aprovado_por"].isna().sum())
    m1, m2, m3 = st.columns(3)
    m1.metric("Fichas", len(df))
    m2.metric("Sem quem gerou", sem_gerou)
    m3.metric("Sem quem aprovou", sem_aprovou)
    fora_aprovacao = carregar_fichas_fora_aprovacao()
    if fora_aprovacao:
        st.caption(
            f"{len(fora_aprovacao)} fichas que aguardavam aprovação em 21/10/2026 "
            "não aparecem na tela Aprovação: são resolvidas só aqui."
        )

    f1, f2 = st.columns([1, 1.4])
    mostrar = f1.radio(
        "Mostrar", [GA_FALTANDO, GA_TODAS], key="ga_mostrar",
        horizontal=True, on_change=ga_guardar_edicoes,
    )
    aeronaves = f2.multiselect(
        "Aeronaves", sorted(df["Aeronave"].unique()), key="ga_aeronaves",
        placeholder="Todas", on_change=ga_guardar_edicoes,
    )
    filtrado = df
    if mostrar == GA_FALTANDO:
        filtrado = filtrado.loc[
            filtrado["gerado_por"].isna() | filtrado["aprovado_por"].isna()
        ]
    if aeronaves:
        filtrado = filtrado.loc[filtrado["Aeronave"].isin(aeronaves)]
    filtrado = filtrado.reset_index(drop=True)

    pendentes = st.session_state.setdefault("ga_pendentes", {})

    with st.expander("Preencher várias fichas de uma vez"):
        st.caption(
            "Coloca a pessoa escolhida só nos campos vazios das fichas listadas "
            "abaixo. Nada é gravado até você clicar em Salvar."
        )
        l1, l2, l3 = st.columns([1, 1, 0.8], vertical_alignment="bottom")
        lote_gerou = l1.selectbox(
            "Gerou", opcoes, index=None, key="ga_lote_gerou", placeholder="Não alterar"
        )
        lote_aprovou = l2.selectbox(
            "Aprovou", opcoes, index=None, key="ga_lote_aprovou", placeholder="Não alterar"
        )
        if l3.button(
            f"Preencher {len(filtrado)} fichas", key="ga_lote",
            disabled=filtrado.empty or not (lote_gerou or lote_aprovou),
            use_container_width=True,
        ):
            ga_guardar_edicoes()
            for _, linha in filtrado.iterrows():
                atual = {
                    "Gerou": rotulo_por_usuario.get(linha["gerado_por"]),
                    "Aprovou": rotulo_por_usuario.get(linha["aprovado_por"]),
                    **pendentes.get(linha["chave"], {}),
                }
                novos = {}
                if lote_gerou and atual["Gerou"] in (None, "", GA_VAZIO):
                    novos["Gerou"] = lote_gerou
                if lote_aprovou and atual["Aprovou"] in (None, "", GA_VAZIO):
                    novos["Aprovou"] = lote_aprovou
                if novos:
                    pendentes.setdefault(linha["chave"], {}).update(novos)
            st.rerun()

    # Campos vazios aparecem como "—" (e não como "None") na tabela.
    tabela = pd.DataFrame({
        "Aeronave": filtrado["Aeronave"],
        "Pesagem": filtrado["Pesagem"],
        "Revisão": filtrado["Revisão"],
        "Data da pesagem": filtrado["Data da pesagem"].map(
            lambda data: data.strftime("%d/%m/%Y") if pd.notna(data) else ""
        ),
        "Gerou": filtrado["gerado_por"].map(rotulo_por_usuario).fillna(GA_VAZIO),
        "Aprovou": filtrado["aprovado_por"].map(rotulo_por_usuario).fillna(GA_VAZIO),
    })
    for posicao, chave in enumerate(filtrado["chave"]):
        for coluna, valor in pendentes.get(chave, {}).items():
            tabela.at[posicao, coluna] = valor
    st.session_state["ga_chaves_exibidas"] = list(filtrado["chave"])

    if filtrado.empty:
        st.success("Nenhuma ficha faltando quem gerou ou aprovou nesta seleção.")
        editada = tabela
    else:
        editada = st.data_editor(
            tabela,
            key=f"ga_editor_{st.session_state.get('ga_versao', 0)}",
            hide_index=True,
            use_container_width=True,
            height=min(38 + 35 * max(len(tabela), 1), 560),
            disabled=["Aeronave", "Pesagem", "Revisão", "Data da pesagem"],
            column_config={
                "Gerou": st.column_config.SelectboxColumn(options=[GA_VAZIO, *opcoes], width="medium"),
                "Aprovou": st.column_config.SelectboxColumn(options=[GA_VAZIO, *opcoes], width="medium"),
            },
        )

    # Junta o que está na tabela agora com o que ficou guardado de outros filtros.
    valores = dict(pendentes)
    for posicao, chave in enumerate(filtrado["chave"]):
        valores[chave] = {
            "Gerou": editada.at[posicao, "Gerou"],
            "Aprovou": editada.at[posicao, "Aprovou"],
        }
    originais = df.set_index("chave")
    mudancas = {}
    for chave, novos in valores.items():
        if chave not in originais.index:
            continue
        # O pandas pode trocar None por NaN; tudo que não é login vira None.
        antes_gerou = ga_login(originais.at[chave, "gerado_por"])
        antes_aprovou = ga_login(originais.at[chave, "aprovado_por"])
        novo_gerou = ga_novo_login(novos, "Gerou", antes_gerou, usuario_por_rotulo)
        novo_aprovou = ga_novo_login(novos, "Aprovou", antes_aprovou, usuario_por_rotulo)
        alteracoes = [
            {
                "campo": campo,
                "antes": rotulo_por_usuario.get(antes, "—") if antes else "—",
                "depois": rotulo_por_usuario.get(depois, "—") if depois else "—",
            }
            for campo, antes, depois in (
                ("Gerada por", antes_gerou, novo_gerou),
                ("Aprovada por", antes_aprovou, novo_aprovou),
            )
            if antes != depois
        ]
        if alteracoes:
            mudancas[chave] = {
                "gerado_por": novo_gerou,
                "aprovado_por": novo_aprovou,
                "alteracoes": alteracoes,
            }

    b1, b2, b3 = st.columns([1, 1, 2], vertical_alignment="center")
    if b1.button(
        f"Salvar {len(mudancas)} fichas" if mudancas else "Salvar",
        key="ga_salvar", type="primary", disabled=not mudancas,
        use_container_width=True,
    ):
        definir_gerou_aprovou(mudancas, st.session_state["usuario_id"])
        ga_descartar()
        st.session_state["ga_mensagem"] = f"{len(mudancas)} fichas atualizadas."
        st.rerun()
    b2.button(
        "Descartar alterações", key="ga_descartar", disabled=not mudancas,
        on_click=ga_descartar, use_container_width=True,
    )
    if mudancas:
        b3.caption(f"{len(mudancas)} fichas com alterações ainda não salvas.")

def excluir_ficha(prefixo, pesagem, revisao):
    garantir_pode_emitir()
    chave = (
        safe_str(prefixo),
        normalizar_chave_ficha(pesagem),
        normalizar_chave_ficha(revisao),
    )
    with conectar_banco() as conn:
        bloquear_prefixo(conn, prefixo)
        quantidade = 0
        for id_ficha in buscar_id_ficha(conn, prefixo, pesagem, revisao):
            quantidade += conn.execute(
                'DELETE FROM pesagens WHERE id = ?', (id_ficha,)
            ).rowcount
        if quantidade:
            conn.execute(
                '''DELETE FROM fluxo_fichas
                   WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
                chave,
            )
            conn.execute(
                '''DELETE FROM campos_com_erro
                   WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
                chave,
            )
            conn.execute(
                '''DELETE FROM historico_fichas
                   WHERE prefixo = ? AND pesagem = ? AND revisao = ?''',
                chave,
            )
    if quantidade:
        invalidar_cache_fichas()
    return quantidade


@protegido
def tela_excluir_ficha():
    if not exigir_nivel(NIVEL_OPERACAO, NIVEL_ADMIN):
        return

    cabecalho("Excluir ficha", "Remoção permanente de uma revisão.")
    colunas_chave = ['Prefixo', 'Pesagem', 'Revisao']
    if df_historico.empty or not all(coluna in df_historico.columns for coluna in colunas_chave):
        st.info("Não há fichas disponíveis para exclusão.")
        return

    prefixos = sorted(df_historico['Prefixo'].dropna().astype(str).unique().tolist())
    prefixo = st.selectbox("Aeronave", prefixos, key="excluir_prefixo")
    fichas = df_historico.loc[
        df_historico['Prefixo'].astype(str) == prefixo,
        ['Pesagem', 'Revisao'],
    ].dropna().astype(str).drop_duplicates()
    opcoes = {
        f"Pesagem {linha.Pesagem} | Revisão {linha.Revisao}": (linha.Pesagem, linha.Revisao)
        for linha in fichas.itertuples(index=False)
    }
    if not opcoes:
        st.info("Não há fichas para esta aeronave.")
        return

    selecionada = st.selectbox("Ficha", list(opcoes), key="excluir_ficha")
    pesagem, revisao = opcoes[selecionada]
    fluxo = buscar_fluxo_ficha(prefixo, pesagem, revisao)
    st.write(f"Gerada por: {fluxo['gerador_nome']}")
    st.write(f"Aprovada por: {fluxo['aprovador_nome']}")

    with st.form("form_excluir_ficha"):
        st.warning("A exclusão é permanente e também remove os dados de aprovação desta ficha.")
        confirmar = st.checkbox(
            f"Confirmo excluir {prefixo}, Pesagem {pesagem}, Revisão {revisao}."
        )
        enviar = st.form_submit_button("Excluir ficha", type="primary")

    if enviar:
        if not confirmar:
            st.error("Marque a confirmação para excluir a ficha.")
        else:
            removidas = excluir_ficha(prefixo, pesagem, revisao)
            if removidas:
                st.success(f"Ficha excluída. Registros removidos: {removidas}.")
                st.rerun()
            else:
                st.error("A ficha não foi encontrada no banco de dados.")

# 5. ROTEAMENTO E BARRA LATERAL
def logo_html():
    for caminho in ("assets/logo_gol.svg", "assets/logo_gol.png"):
        if os.path.exists(caminho):
            import base64
            tipo = "svg+xml" if caminho.endswith(".svg") else "png"
            with open(caminho, "rb") as arquivo:
                dados = base64.b64encode(arquivo.read()).decode()
            return (
                f'<img src="data:image/{tipo};base64,{dados}" '
                'style="height:34px;width:auto;flex-shrink:0">'
            )
    return '<div class="marca-logo">W&B</div>'


MARCA_HTML = (
    f'<div class="marca">{logo_html()}<div>'
    '<div class="marca-nome">Pesagem e Balanceamento</div>'
    '<div class="marca-sub">Engenharia GOL</div></div></div>'
)

PAGINAS = {
    'consulta': ("Consultar fichas", ":material/dashboard:", tela_consulta, {1, 2, 3, 4, 5}),
    'graficos': ("Gráficos da frota", ":material/bar_chart:", tela_graficos, {1, 2, 3, 5}),
    'nova_ficha': ("Nova ficha", ":material/add_circle:", tela_nova_ficha, {1, 3, 5}),
    'edicao': ("Editar ficha", ":material/edit_note:", tela_edicao, {1, 3, 5}),
    'aprovar': ("Aprovação", ":material/task_alt:", tela_aprovar_fichas, {1, 3, 5}),
    'cadastrar_aeronave': ("Aeronaves", ":material/flight:", tela_cadastrar_aeronave, {1, 3}),
    'criar_login': ("Usuários", ":material/group:", tela_criar_login, {3}),
    'assinaturas': ("Assinaturas", ":material/draw:", tela_assinaturas, {3}),
    'excluir': ("Excluir ficha", ":material/delete:", tela_excluir_ficha, {1, 3}),
    # Além do perfil, só aparece para o usuário Engenharia (pode_preencher_gerou_aprovou).
    'gerou_aprovou': ("Gerou e aprovou", ":material/how_to_reg:", tela_gerou_aprovou, {1, 2, 3, 4, 5}),
}
GRUPOS_MENU = [
    ("Consulta", ['consulta', 'graficos']),
    ("Emissão", ['nova_ficha', 'edicao', 'aprovar', 'excluir']),
    ("Administração", ['cadastrar_aeronave', 'criar_login', 'assinaturas', 'gerou_aprovou']),
]


def pagina_permitida(pagina, nivel):
    if pagina == 'gerou_aprovou' and not pode_preencher_gerou_aprovou():
        return False
    return pagina in PAGINAS and nivel in PAGINAS[pagina][3]


def ir_para(pagina):
    st.session_state['pagina_atual'] = pagina


if st.session_state['usuario_logado'] and st.session_state.get('boas_vindas'):
    # Tela curta de boas-vindas enquanto os dados das telas carregam.
    primeiro_nome = safe_str(st.session_state.get('nome_usuario')).split(" ")[0]
    st.markdown(
        f"""
<style>
[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] {{ display: none; }}
@keyframes gol-girar {{ to {{ transform: rotate(360deg); }} }}
@keyframes gol-surgir {{ from {{ opacity: 0; transform: translateY(6px); }} to {{ opacity: 1; transform: none; }} }}
.boas-vindas {{ height: 70vh; display: flex; flex-direction: column; align-items: center;
  justify-content: center; gap: 18px; animation: gol-surgir .5s ease-out; }}
.boas-vindas h2 {{ margin: 0; font-weight: 600; color: #1F2430; }}
.boas-vindas p {{ margin: 0; color: #6B7280; font-size: .95rem; }}
.gol-loader {{ width: 28px; height: 28px; border-radius: 50%; border: 3px solid #FFE1CC;
  border-top-color: #FF6A13; animation: gol-girar .8s linear infinite; }}
</style>
<div class="boas-vindas">
  {MARCA_HTML}
  <h2>Bem-vindo, {primeiro_nome}</h2>
  <div class="gol-loader"></div>
  <p>Carregando as fichas…</p>
</div>
""",
        unsafe_allow_html=True,
    )
    inicio = time.monotonic()
    listar_fichas_pendentes()
    carregar_prefixos_fichas()
    restante = 1.2 - (time.monotonic() - inicio)
    if restante > 0:
        time.sleep(restante)
    st.session_state.pop('boas_vindas', None)
    st.rerun()
elif not st.session_state['usuario_logado']:
    _, centro, _ = st.columns([1, 1.2, 1])
    with centro:
        st.markdown("<div style='height:12vh'></div>", unsafe_allow_html=True)
        with st.container(border=True):
            st.markdown(MARCA_HTML, unsafe_allow_html=True)
            st.caption("Entre com seu usuário para acessar as fichas.")
            with st.form("login_form", border=False):
                usuario = st.text_input("Usuário")
                senha = st.text_input("Senha", type="password")
                entrar = st.form_submit_button("Acessar", type="primary", use_container_width=True)
            if entrar:
                registro_usuario = autenticar_usuario(usuario, senha)
                if registro_usuario:
                    st.session_state['usuario_logado'] = True
                    st.session_state['usuario_id'] = registro_usuario['usuario']
                    st.session_state['nivel_acesso'] = registro_usuario['nivel_acesso']
                    st.session_state['nome_usuario'] = registro_usuario['nome']
                    st.session_state['pagina_atual'] = 'consulta'
                    st.session_state['boas_vindas'] = True
                    st.rerun()
                else:
                    st.error("Usuário ou senha inválidos.")
    # Enquanto a pessoa digita a senha, deixa os dados das telas prontos.
    listar_fichas_pendentes()
    carregar_prefixos_fichas()
else:
    nivel = st.session_state['nivel_acesso']
    if not pagina_permitida(st.session_state['pagina_atual'], nivel):
        st.session_state['pagina_atual'] = 'consulta'

    with st.sidebar:
        st.markdown(MARCA_HTML, unsafe_allow_html=True)
        for grupo, paginas in GRUPOS_MENU:
            visiveis = [p for p in paginas if pagina_permitida(p, nivel)]
            if not visiveis:
                continue
            st.markdown(f'<div class="grupo-menu">{grupo}</div>', unsafe_allow_html=True)
            for pagina in visiveis:
                rotulo, icone, _, _ = PAGINAS[pagina]
                st.button(
                    rotulo,
                    key=f"menu_{pagina}",
                    use_container_width=True,
                    type="primary" if st.session_state['pagina_atual'] == pagina else "secondary",
                    on_click=ir_para,
                    args=(pagina,),
                )
        st.divider()
        nome = safe_str(st.session_state["nome_usuario"])
        iniciais = "".join(parte[0] for parte in nome.split()[:2]).upper() or "?"
        st.markdown(
            f'<div class="usuario"><div class="usuario-avatar">{iniciais}</div>'
            f'<div><b>{nome}</b><br><span>{PERFIS.get(nivel, "—")}</span></div></div>',
            unsafe_allow_html=True,
        )
        if st.button("Sair", key="menu_sair", use_container_width=True):
            st.session_state['usuario_logado'] = False
            st.session_state['nivel_acesso'] = 0
            st.session_state['usuario_id'] = ""
            st.session_state['nome_usuario'] = ""
            st.session_state['pagina_atual'] = 'consulta'
            st.rerun()

    PAGINAS[st.session_state['pagina_atual']][2]()
